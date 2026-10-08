"""SEMP - System Equipment Management Program.
Run: python SEMP.py [path/to/project.semp]
Sections: imports, store, theme, icons, Excel, widgets, dialogs, main window.
"""
from __future__ import annotations

import sys
from pathlib import Path


def _relaunch_windowed():
    """Use pythonw when this file itself was opened on Windows."""
    if sys.platform != "win32" or __name__ != "__main__":
        return False
    try:
        script = Path(__file__).resolve()
        direct = Path(sys.argv[0]).resolve() == script
        local_pythonw = script.parent / ".venv" / "Scripts" / "pythonw.exe"
        pythonw = local_pythonw if local_pythonw.is_file() else Path(sys.executable).with_name("pythonw.exe")
        already_windowed = Path(sys.executable).resolve() == pythonw.resolve()
    except OSError:
        return False
    if not direct or already_windowed or not pythonw.is_file():
        return False
    import os
    import subprocess
    environment = os.environ.copy()
    environment["SEMP_WINDOWED"] = "1"
    subprocess.Popen([str(pythonw), str(script), *sys.argv[1:]],
                     cwd=str(script.parent), env=environment)
    return True


def _startup_error(exc_type, exc_value, exc_tb):
    """Keep startup failures visible when launched using pythonw on Windows."""
    import traceback
    import tempfile
    from pathlib import Path
    details = ''.join(traceback.format_exception(exc_type, exc_value, exc_tb))
    log = Path(tempfile.gettempdir()) / 'semp-startup-error.log'
    try:
        log.write_text(details, encoding='utf-8')
    except OSError:
        pass
    if sys.stderr is not None:
        sys.__excepthook__(exc_type, exc_value, exc_tb)
    elif sys.platform == 'win32':
        import ctypes
        ctypes.windll.user32.MessageBoxW(
            0, f'SEMP를 시작하지 못했습니다.\n먼저 install.cmd를 실행하세요.\n\n진단 파일: {log}\n\n{details[-900:]}',
            'SEMP 실행 오류', 16)


if _relaunch_windowed():
    raise SystemExit(0)

if __name__ == '__main__':
    sys.excepthook = _startup_error

from contextlib import contextmanager
from datetime import datetime, date, timedelta
import json
import re
import sqlite3
import uuid
from PySide6.QtCore import QByteArray, Qt
from PySide6.QtGui import QIcon, QPainter, QPixmap
from PySide6.QtSvg import QSvgRenderer
from dataclasses import dataclass
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PySide6.QtCore import Qt, Signal, QMimeData, QSize, QRect, QPoint
from PySide6.QtGui import QDrag, QTextCharFormat, QFont, QColor, QKeySequence, QTextDocument, QPainter, QPalette
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton, QTreeWidget,
    QListWidget, QTextEdit, QToolButton, QDialog, QGridLayout, QSpinBox, QComboBox, QLineEdit, QCompleter, QSizePolicy, QLayout, QButtonGroup, QInputDialog, QCalendarWidget,
)
from PySide6.QtCore import Qt, QDate, QMimeData, QPoint, QSize
from PySide6.QtGui import QDrag, QColor
from PySide6.QtWidgets import (
    QDialog, QVBoxLayout, QFormLayout, QLineEdit, QComboBox, QTextEdit,
    QDialogButtonBox, QMessageBox, QDateEdit, QLabel, QTableWidget,
    QTableWidgetItem, QHeaderView, QFileDialog, QWidget, QHBoxLayout, QTreeWidget, QTreeWidgetItem, QCalendarWidget, QInputDialog, QListWidget, QSplitter, QFrame,
)
from datetime import datetime
import html
import shutil
from PySide6.QtCore import Qt, QSize, QLockFile, QDate, QSettings, QTimer
from PySide6.QtGui import QAction, QKeySequence, QFont, QColor
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QFrame, QVBoxLayout, QHBoxLayout,
    QSplitter, QStackedWidget, QTreeWidget, QTreeWidgetItem, QListWidgetItem,
    QLineEdit, QComboBox, QCheckBox, QTabWidget, QTextBrowser, QFormLayout,
    QFileDialog, QInputDialog, QMessageBox, QMenu, QDialog, QAbstractItemView,
    QStyle, QListWidget, QTextEdit, QSizePolicy, QTabBar, QScrollArea, QStyledItemDelegate, QStyleOptionComboBox,
)

__version__ = "0.1.0"

# ============================================================================
# STORE
# ============================================================================

KINDS = {"group": "그룹", "device": "장치", "unit": "개별 장비"}
CATEGORIES = ("사용", "위치", "출고", "입고", "이슈", "용도변경", "기타")


def tags(text: str) -> str:
    return " ".join(dict.fromkeys("#" + t for t in re.split(r"[\s,#]+", text.strip()) if t))


def serial_range(first: str, last: str, limit: int = 10000) -> list[str]:
    first, last = first.strip(), last.strip()
    if not first:
        if last:
            raise ValueError("일련번호 끝만 입력할 수 없습니다.")
        return []
    if not last or first == last:
        return [first]
    a, b = re.fullmatch(r"(.*?)([0-9]+)", first), re.fullmatch(r"(.*?)([0-9]+)", last)
    if not a or not b or a[1] != b[1]:
        raise ValueError("범위는 같은 접두어와 끝의 숫자로 입력하세요. 예: PSU-001 ~ PSU-010")
    lo, hi = int(a[2]), int(b[2])
    padded = a[2].startswith("0") or b[2].startswith("0")
    if padded and len(a[2]) != len(b[2]):
        raise ValueError("앞에 0을 붙인 일련번호는 처음과 끝의 자릿수를 맞춰주세요.")
    if hi < lo:
        raise ValueError("일련번호 끝이 처음보다 작습니다.")
    if hi - lo + 1 > limit:
        raise ValueError(f"한 번에 최대 {limit:,}개까지 등록할 수 있습니다.")
    width = len(a[2]) if padded else 0
    return [a[1] + str(i).zfill(width) for i in range(lo, hi + 1)]


class Store:
    def __init__(self, project_file: str | Path):
        self.path = Path(project_file).resolve()
        manifest = json.loads(self.path.read_text(encoding="utf-8"))
        if manifest.get("format") != "SEMP" or manifest.get("version") != 1:
            raise ValueError("지원하지 않는 SEMP 프로젝트 형식입니다.")
        self.name = manifest["name"]
        self.root = self.path.parent
        db_path = self.root / "data" / "semp.sqlite3"
        if not db_path.is_file():
            raise ValueError("프로젝트의 data/semp.sqlite3 파일이 없습니다. 전체 폴더를 옮겨주세요.")
        self.db = sqlite3.connect(db_path, isolation_level=None)
        self.db.row_factory = sqlite3.Row
        self.db.execute("PRAGMA foreign_keys = ON")
        self.db.execute("PRAGMA journal_mode = DELETE")
        if self.db.execute("PRAGMA user_version").fetchone()[0] != 1:
            self.db.close()
            raise ValueError("지원하지 않는 데이터베이스 버전입니다.")

        entry_columns = {r[1] for r in self.db.execute("PRAGMA table_info(entries)")}
        if "end_day" not in entry_columns:
            self.db.execute("ALTER TABLE entries ADD COLUMN end_day TEXT NOT NULL DEFAULT ''")
        if "purpose_after" not in entry_columns:
            self.db.execute("ALTER TABLE entries ADD COLUMN purpose_after TEXT")
        if "location" not in {r[1] for r in self.db.execute("PRAGMA table_info(nodes)")}:
            self.db.execute("ALTER TABLE nodes ADD COLUMN location TEXT NOT NULL DEFAULT ''")
        self.db.executescript("""
            CREATE TABLE IF NOT EXISTS project_metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS categories (name TEXT PRIMARY KEY);
            CREATE TABLE IF NOT EXISTS recent_inputs (
                kind TEXT NOT NULL, value TEXT NOT NULL, used_at TEXT NOT NULL,
                PRIMARY KEY(kind,value));
            CREATE TABLE IF NOT EXISTS collections (id TEXT PRIMARY KEY, name TEXT NOT NULL UNIQUE);
            CREATE TABLE IF NOT EXISTS collection_nodes (
                collection_id TEXT REFERENCES collections(id) ON DELETE CASCADE,
                node_id TEXT REFERENCES nodes(id), PRIMARY KEY(collection_id,node_id));
        """)
        self.db.execute("INSERT OR IGNORE INTO project_metadata VALUES('identity',?)", (uuid.uuid4().hex,))
        self.project_identity = self.db.execute("SELECT value FROM project_metadata WHERE key='identity'").fetchone()[0]

    @classmethod
    def create(cls, parent: str | Path, name: str, default_groups=True):
        name = name.strip()
        if not name or re.search(r'[<>:"/\\|?*\x00-\x1f]', name) or name.endswith((".", " ")):
            raise ValueError("폴더 이름에 사용할 수 있는 프로젝트 이름을 입력하세요.")
        if name.upper().split(".")[0] in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}:
            raise ValueError("Windows 예약 이름은 사용할 수 없습니다.")
        root = Path(parent) / name
        root.mkdir(exist_ok=False)
        (root / "data").mkdir()
        (root / "backups").mkdir()
        db = sqlite3.connect(root / "data" / "semp.sqlite3")
        db.executescript("""
            PRAGMA user_version = 1;
            CREATE TABLE nodes (
                id TEXT PRIMARY KEY, parent TEXT REFERENCES nodes(id),
                kind TEXT NOT NULL CHECK(kind IN ('group','device','unit')),
                name TEXT NOT NULL, serial TEXT NOT NULL DEFAULT '',
                purpose TEXT NOT NULL DEFAULT '', description TEXT NOT NULL DEFAULT '',
                tags TEXT NOT NULL DEFAULT '', deleted TEXT,
                created TEXT NOT NULL DEFAULT (strftime('%Y-%m-%dT%H:%M:%f','now'))
            );
            CREATE INDEX node_parent ON nodes(parent);
            CREATE UNIQUE INDEX unit_serial ON nodes(parent,serial)
                WHERE kind='unit' AND deleted IS NULL;
            CREATE UNIQUE INDEX container_name ON nodes(ifnull(parent,''),kind,name)
                WHERE kind!='unit' AND deleted IS NULL;
            CREATE TABLE entries (
                id TEXT PRIMARY KEY, event INTEGER NOT NULL DEFAULT 0,
                category TEXT NOT NULL, day TEXT NOT NULL, title TEXT NOT NULL,
                html TEXT NOT NULL DEFAULT '', plain TEXT NOT NULL DEFAULT '',
                tags TEXT NOT NULL DEFAULT '', location TEXT NOT NULL DEFAULT '',
                end_day TEXT NOT NULL DEFAULT '', purpose_after TEXT,
                deleted TEXT, created TEXT NOT NULL DEFAULT (strftime('%Y-%m-%dT%H:%M:%f','now'))
            );
            CREATE TABLE entry_nodes (
                entry_id TEXT NOT NULL REFERENCES entries(id),
                node_id TEXT NOT NULL REFERENCES nodes(id),
                PRIMARY KEY(entry_id,node_id)
            );
            CREATE INDEX entry_target ON entry_nodes(node_id);
            CREATE TABLE purposes (name TEXT PRIMARY KEY);
            CREATE TABLE trash (id TEXT PRIMARY KEY, label TEXT NOT NULL, created TEXT NOT NULL);
        """)
        db.close()
        path = root / "project.semp"
        path.write_text(json.dumps({"format": "SEMP", "version": 1, "name": name}, ensure_ascii=False, indent=2), encoding="utf-8")
        store = cls(path)
        if default_groups:
            with store.atomic():
                for group in ("탑재장치", "점검장비", "케이블", "기타"):
                    store.add_node("group", group)
        return store

    @contextmanager
    def atomic(self):
        token = "s" + uuid.uuid4().hex
        self.db.execute(f"SAVEPOINT {token}")
        try:
            yield
        except Exception:
            self.db.execute(f"ROLLBACK TO {token}")
            self.db.execute(f"RELEASE {token}")
            raise
        else:
            self.db.execute(f"RELEASE {token}")

    def categories(self):
        extra = list(dict.fromkeys([r[0] for r in self.db.execute("SELECT name FROM categories ORDER BY rowid")] + [r[0] for r in self.db.execute("SELECT DISTINCT category FROM entries")]))
        return list(CATEGORIES) + [name for name in extra if name not in CATEGORIES]

    def collections(self):
        return [dict(r) for r in self.db.execute("SELECT * FROM collections ORDER BY rowid")]

    def collection_targets(self, ident):
        return [r[0] for r in self.db.execute("SELECT node_id FROM collection_nodes JOIN nodes ON nodes.id=node_id WHERE collection_id=? AND deleted IS NULL", (ident,))]

    def save_collection(self, name, node_ids, ident=None):
        if not name.strip(): raise ValueError("장비 모음 이름을 입력하세요.")
        node_ids = list(dict.fromkeys(node_ids))
        if not node_ids: raise ValueError("모음에 넣을 장비를 선택하세요.")
        for node_id in node_ids: self.node(node_id)
        ident = ident or uuid.uuid4().hex
        with self.atomic():
            self.db.execute("INSERT INTO collections VALUES(?,?) ON CONFLICT(id) DO UPDATE SET name=excluded.name", (ident, name.strip()))
            self.db.execute("DELETE FROM collection_nodes WHERE collection_id=? AND node_id IN (SELECT id FROM nodes WHERE deleted IS NULL)", (ident,))
            self.db.executemany("INSERT INTO collection_nodes VALUES(?,?)", [(ident, i) for i in node_ids])
        return ident

    def delete_collection(self, ident):
        with self.atomic():
            self.db.execute("DELETE FROM collections WHERE id=?", (ident,))

    def nodes(self, include_deleted=False):
        return [self._node_record(r) for r in self.db.execute(
            "SELECT n.*, p.name AS device_name FROM nodes n LEFT JOIN nodes p ON p.id=n.parent"
            + ("" if include_deleted else " WHERE n.deleted IS NULL") + " ORDER BY n.created,n.rowid")]

    @staticmethod
    def _node_record(row):
        node = dict(row)
        device_name = node.pop("device_name", None)
        if node["kind"] == "unit" and device_name is not None:
            node["stored_name"] = node["name"]
            node["name"] = device_name
        return node

    def node(self, node_id):
        row = self.db.execute("SELECT n.*, p.name AS device_name FROM nodes n LEFT JOIN nodes p ON p.id=n.parent WHERE n.id=? AND n.deleted IS NULL", (node_id,)).fetchone()
        if not row:
            raise ValueError("장비를 찾을 수 없습니다.")
        return self._node_record(row)

    def validate_parent(self, kind, parent):
        p = self.node(parent) if parent else None
        if kind == "unit" and (not p or p["kind"] != "device"):
            raise ValueError("일련번호는 장치 아래에 등록하세요.")
        if kind != "unit" and p and p["kind"] != "group":
            raise ValueError("그룹과 장치는 그룹 아래 또는 프로젝트 최상위에 등록하세요.")

    def add_node(self, kind, name, parent=None, serial="", purpose="", description="", tag_text="", purpose_options=()):
        if kind not in KINDS or not name.strip():
            raise ValueError("이름을 입력하세요.")
        self.validate_parent(kind, parent)
        if kind == "unit" and not serial.strip():
            raise ValueError("개별 장비의 일련번호를 입력하세요.")
        ident = uuid.uuid4().hex
        with self.atomic():
            self.db.execute("INSERT INTO nodes(id,parent,kind,name,serial,purpose,description,tags) VALUES(?,?,?,?,?,?,?,?)",
                            (ident, parent, kind, name.strip(), serial.strip(), purpose.strip(), description, tags(tag_text)))
            self.db.executemany("INSERT OR IGNORE INTO purposes VALUES(?)", [(value.strip(),) for value in [*purpose_options, purpose] if value.strip()])
            self.remember_inputs(tag_text=tag_text)
        return ident

    def update_node(self, ident, name, serial, purpose, description, tag_text, purpose_options=()):
        node = self.node(ident)
        if not name.strip() or (node["kind"] == "unit" and not serial.strip()):
            raise ValueError("이름과 개별 장비의 일련번호를 입력하세요.")
        with self.atomic():
            self.db.execute("UPDATE nodes SET name=?,serial=?,purpose=?,description=?,tags=? WHERE id=?",
                            (name.strip(), serial.strip(), purpose.strip(), description, tags(tag_text), ident))
            self.db.executemany("INSERT OR IGNORE INTO purposes VALUES(?)", [(value.strip(),) for value in [*purpose_options, purpose] if value.strip()])
            self.remember_inputs(tag_text=tag_text)

    def bulk_update_purpose(self, ids, purpose, purpose_options=()):
        targets = [ident for ident in self.descendants(ids) if self.node(ident)["kind"] == "unit"]
        if not targets:
            raise ValueError("선택한 항목 아래에 용도를 변경할 개별 장비가 없습니다.")
        purpose = purpose.strip()
        with self.atomic():
            self.db.executemany("UPDATE nodes SET purpose=? WHERE id=?", [(purpose, ident) for ident in targets])
            self.db.executemany("INSERT OR IGNORE INTO purposes VALUES(?)",
                                [(value.strip(),) for value in [*purpose_options, purpose] if value.strip()])
        return len(targets)

    def update_tags(self, ident, tag_text):
        self.node(ident)
        with self.atomic():
            self.db.execute("UPDATE nodes SET tags=? WHERE id=?", (tags(tag_text), ident))
            self.remember_inputs(tag_text=tag_text)

    def ancestors(self, ident):
        result, seen = [], set()
        while ident:
            if ident in seen:
                raise ValueError("장비 트리에 순환 참조가 있습니다.")
            seen.add(ident)
            n = self.node(ident)
            result.append(n)
            ident = n["parent"]
        return result

    def path_label(self, ident, nodes_by_id=None):
        if nodes_by_id is not None:
            parts, seen = [], set()
            while ident:
                if ident in seen:
                    raise ValueError("장비 트리에 순환 참조가 있습니다.")
                seen.add(ident)
                node = nodes_by_id[ident]
                parts.append(node["serial"] if node["kind"] == "unit" else node["name"])
                ident = node["parent"]
            return " / ".join(reversed(parts))
        return " / ".join(n["serial"] if n["kind"] == "unit" else n["name"] for n in reversed(self.ancestors(ident)))

    def descendants(self, ids):
        children = {}
        for n in self.nodes():
            children.setdefault(n["parent"], []).append(n["id"])
        found, pending = set(), list(ids)
        while pending:
            ident = pending.pop()
            if ident not in found:
                self.node(ident)
                found.add(ident)
                pending.extend(children.get(ident, []))
        return found

    def duplicate_nodes(self, ids):
        selected = list(dict.fromkeys(ids))
        roots = [ident for ident in selected if not any(a["id"] in selected for a in self.ancestors(ident)[1:])]
        if any(self.node(ident)["kind"] not in ("group", "device") for ident in roots):
            raise ValueError("그룹 또는 장치를 선택해 복제하세요.")
        nodes = self.nodes()
        by_id = {n["id"]: n for n in nodes}
        children = {}
        names = {}
        for n in nodes:
            children.setdefault(n["parent"], []).append(n)
            names.setdefault((n["parent"], n["kind"]), set()).add(n["name"])
        copies = []
        with self.atomic():
            for ident in roots:
                source = by_id[ident]
                base = re.sub(r"\s*\(\d+\)$", "", source["name"]) or source["name"]
                used = names[(source["parent"], source["kind"])]
                index = 1
                while f"{base} ({index})" in used: index += 1
                name = f"{base} ({index})"
                used.add(name)
                pending = [(source, source["parent"], name)]
                for original, parent, copy_name in pending:
                    new_id = self.add_node(original["kind"], original.get("stored_name", copy_name), parent, serial=original["serial"], purpose=original["purpose"], description=original["description"], tag_text=original["tags"])
                    if original["id"] == ident: copies.append(new_id)
                    pending.extend((child, new_id, child["name"]) for child in children.get(original["id"], []))
        return copies

    def move_nodes(self, ids, parent):
        selected = set(ids)
        roots = [i for i in ids if not any(a["id"] in selected for a in self.ancestors(i)[1:])]
        if parent in self.descendants(roots):
            raise ValueError("자기 자신이나 하위 항목 안으로 이동할 수 없습니다.")
        with self.atomic():
            for ident in roots:
                self.validate_parent(self.node(ident)["kind"], parent)
                self.db.execute("UPDATE nodes SET parent=? WHERE id=?", (parent, ident))

    def import_rows(self, parent, rows):
        self.validate_parent("device", parent)
        count, created = 0, []
        with self.atomic():
            for values in rows:
                name, serials = values[:2]
                purpose = values[2] if len(values) > 2 else ""
                existing = self.db.execute("SELECT id FROM nodes WHERE parent IS ? AND kind='device' AND name=? AND deleted IS NULL", (parent, name)).fetchone()
                device = existing[0] if existing else self.add_node("device", name, parent)
                if not existing:
                    created.append(device)
                for serial in serials:
                    created.append(self.add_node("unit", name, device, serial=serial, purpose=purpose))
                    count += 1
        return created, count

    def save_entry(self, *, title, category, day, html, plain, tag_text="", location="", node_ids=(), event=False, ident=None, category_options=(), end_day=None, purpose_after=None):
        if not title.strip() or not category.strip():
            raise ValueError("제목과 기록 종류를 입력하세요.")
        category = category.strip()
        date.fromisoformat(day)
        end_day = end_day or day
        date.fromisoformat(end_day)
        if end_day < day: raise ValueError("종료일은 시작일보다 빠를 수 없습니다.")
        node_ids = list(dict.fromkeys(node_ids))
        if not event and len(node_ids) != 1:
            raise ValueError("일반 기록은 대상 항목 하나가 필요합니다.")
        for n in node_ids:
            self.node(n)
        ident = ident or uuid.uuid4().hex
        with self.atomic():
            self.db.execute("""INSERT INTO entries(id,event,category,day,title,html,plain,tags,location,end_day,purpose_after)
                VALUES(?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET
                category=excluded.category,day=excluded.day,title=excluded.title,html=excluded.html,
                plain=excluded.plain,tags=excluded.tags,location=excluded.location,end_day=excluded.end_day,
                purpose_after=excluded.purpose_after""",
                (ident, int(event), category, day, title.strip(), html, plain, tags(tag_text), location.strip(), end_day, purpose_after))
            self.db.executemany("INSERT OR IGNORE INTO categories VALUES(?)", [(name.strip(),) for name in [*category_options, category] if name.strip()])
            self.remember_inputs(location, tag_text)
            # Retain links to trashed equipment so restoring it recovers its event history.
            self.db.execute("DELETE FROM entry_nodes WHERE entry_id=? AND node_id IN (SELECT id FROM nodes WHERE deleted IS NULL)", (ident,))
            self.db.executemany("INSERT INTO entry_nodes VALUES(?,?)", [(ident, n) for n in node_ids])
        return ident

    def save_record(self, *, node_ids, purpose_options=(), ident=None, **values):
        """Save a record and apply only intentional purpose changes atomically."""
        targets = list(dict.fromkeys(node_ids))
        with self.atomic():
            original = self.entry(ident) if ident else None
            apply_purpose = values["category"] == "용도변경" and (
                original is None or original["category"] != "용도변경"
                or original["purpose_after"] != values.get("purpose_after")
                or set(self.targets(ident)) != set(targets))
            result = self.save_entry(ident=ident, node_ids=targets, **values)
            if apply_purpose:
                self.bulk_update_purpose(targets, values["purpose_after"], purpose_options)
        return result

    def entry(self, ident):
        r = self.db.execute("SELECT * FROM entries WHERE id=? AND deleted IS NULL", (ident,)).fetchone()
        if not r:
            raise ValueError("기록을 찾을 수 없습니다.")
        return dict(r)

    def targets(self, ident):
        return [r[0] for r in self.db.execute("SELECT node_id FROM entry_nodes JOIN nodes ON nodes.id=node_id WHERE entry_id=? AND nodes.deleted IS NULL ORDER BY nodes.name,nodes.serial,nodes.created", (ident,))]

    def record_sources(self, query="", limit=51, offset=0):
        """Small pages of record metadata for reusing historical participants."""
        return [dict(r) for r in self.db.execute("""
            SELECT e.id,e.category,e.day,e.end_day,e.title,
                (SELECT count(*) FROM entry_nodes t JOIN nodes n ON n.id=t.node_id
                 WHERE t.entry_id=e.id AND n.deleted IS NULL) AS target_count
            FROM entries e WHERE e.deleted IS NULL
            AND (?='' OR instr(lower(e.title || ' ' || e.category || ' ' || e.plain || ' ' || e.tags || ' ' || e.location),lower(?))>0)
            ORDER BY e.day DESC,e.created DESC,e.rowid DESC LIMIT ? OFFSET ?
        """, (query.strip(), query.strip(), limit, offset))]

    def link_event(self, ident, ids):
        if not self.entry(ident)["event"]:
            raise ValueError("기록에만 여러 장비를 연결할 수 있습니다.")
        # Capture current leaves, so later additions to a group do not rewrite event participation.
        expanded = self.descendants(ids)
        nodes = {n["id"]: n for n in self.nodes()}
        parents = {n["parent"] for n in nodes.values()}
        targets = [i for i in expanded if nodes[i]["kind"] == "unit" or (nodes[i]["kind"] == "device" and i not in parents)]
        if not targets:
            raise ValueError("연결할 장비가 없습니다. 장치 또는 일련번호 항목을 선택하세요.")
        with self.atomic():
            self.db.executemany("INSERT OR IGNORE INTO entry_nodes VALUES(?,?)", [(ident, n) for n in targets])
        return len(targets)

    def unlink_event(self, ident, ids):
        if not self.entry(ident)["event"]:
            raise ValueError("기록만 연결을 해제할 수 있습니다.")
        with self.atomic():
            self.db.executemany("DELETE FROM entry_nodes WHERE entry_id=? AND node_id=?", [(ident, i) for i in ids])

    def entries(self, node_id=None, inherited=True, events_only=False, query="", category="전체", day_from=None, day_to=None):
        ancestors = {n["id"] for n in self.ancestors(node_id)} if node_id and inherited else ({node_id} if node_id else None)
        conditions, params = ["e.deleted IS NULL"], []
        if ancestors is not None:
            placeholders = ",".join("?" for _ in ancestors)
            conditions.append(f"EXISTS (SELECT 1 FROM entry_nodes t WHERE t.entry_id=e.id AND t.node_id IN ({placeholders}))")
            params.extend(ancestors)
        if events_only:
            conditions.append("e.event=1")
        if category != "전체":
            conditions.append("e.category=?")
            params.append(category)
        if day_from is not None:
            conditions.append("COALESCE(NULLIF(e.end_day,''),e.day)>=?")
            params.append(day_from)
        if day_to is not None:
            conditions.append("e.day<=?")
            params.append(day_to)
        where = " AND ".join(conditions)
        nodes = {n["id"]: n for n in self.nodes()}
        links = {}
        for r in self.db.execute("SELECT t.entry_id,t.node_id FROM entry_nodes t JOIN entries e ON e.id=t.entry_id WHERE " + where + " ORDER BY t.rowid", params):
            if r["node_id"] in nodes:
                links.setdefault(r["entry_id"], []).append(r["node_id"])
        result = []
        for row in self.db.execute("SELECT e.* FROM entries e WHERE " + where + " ORDER BY e.day DESC,e.created DESC,e.rowid DESC", params):
            e = dict(row)
            target = links.get(e["id"], [])
            if not e["event"] and not target:
                continue
            if events_only and not e["event"]:
                continue
            if ancestors is not None and not ancestors.intersection(target):
                continue
            if category != "전체" and e["category"] != category:
                continue
            haystack = " ".join([e["title"], e["plain"], e["tags"], e["location"], *(nodes[t]["name"] + " " + nodes[t]["serial"] + " " + nodes[t]["tags"] for t in target)])
            if query.startswith("#"):
                available = tags(e["tags"] + " " + " ".join(nodes[t]["tags"] for t in target)).casefold().split()
                if not all(q in available for q in tags(query).casefold().split()):
                    continue
            elif query.casefold() not in haystack.casefold():
                continue
            e["targets"] = target
            e["inherited"] = bool(node_id and node_id not in target)
            result.append(e)
        return result

    def target_summary(self, entry, limit=1):
        ids = entry["targets"] if "targets" in entry else self.targets(entry["id"])
        names = []
        for ident in ids[:limit]:
            node = self.node(ident)
            if node["kind"] == "unit":
                names.append(f"{node['name']} · {node['serial']}")
            else:
                names.append(node["name"])
        return (", ".join(names) + (f" 외 {len(ids)-limit}개" if len(ids) > limit else "")) or "연결 장비 없음"

    def restore_moves(self, target, expected):
        with self.atomic():
            for ident, parent in expected.items():
                if self.node(ident)["parent"] != parent:
                    raise ValueError("장비 소속이 변경되어 이 이동을 되돌릴 수 없습니다.")
            for ident, parent in target.items():
                self.move_nodes([ident], parent)

    def current_state(self, ident):
        self.node(ident)
        base = """SELECT e.{column} FROM entries e
            WHERE e.deleted IS NULL AND e.day<=? AND {condition}
            AND EXISTS (SELECT 1 FROM entry_nodes t WHERE t.entry_id=e.id AND t.node_id=?)
            ORDER BY e.day DESC,e.created DESC,e.rowid DESC LIMIT 1"""
        params = (date.today().isoformat(), ident)
        location_row = self.db.execute(base.format(column="location", condition="(e.location!='' OR e.category='위치')"), params).fetchone()
        status_row = self.db.execute(base.format(column="category", condition="e.category IN ('사용','출고','입고')"), params).fetchone()
        location = (location_row[0] if location_row else self.node(ident).get("location", "")) or "미지정"
        status = status_row[0] if status_row else "미기록"
        return location, status

    def remember_inputs(self, location="", tag_text=""):
        stamp = datetime.now().isoformat(timespec="microseconds")
        previous = self.db.execute("SELECT MAX(used_at) FROM recent_inputs").fetchone()[0]
        if previous and stamp <= previous:
            stamp = (datetime.fromisoformat(previous) + timedelta(microseconds=1)).isoformat(timespec="microseconds")
        values = [("location", location.strip())] if location.strip() else []
        values.extend(("tags", tag) for tag in tags(tag_text).split())
        self.db.executemany("INSERT INTO recent_inputs VALUES(?,?,?) ON CONFLICT(kind,value) DO UPDATE SET used_at=excluded.used_at", [(kind, value, stamp) for kind, value in values])

    def recent_suggestions(self, kind, limit=5):
        if kind not in ("location", "tags"): raise ValueError("잘못된 추천 항목입니다.")
        result = [r[0] for r in self.db.execute("SELECT value FROM recent_inputs WHERE kind=? ORDER BY used_at DESC,rowid DESC LIMIT ?", (kind, limit))]
        # Existing projects can recommend their saved values before the first new edit.
        if len(result) < limit:
            for row in self.db.execute(f"SELECT {kind} FROM entries WHERE deleted IS NULL ORDER BY created DESC,rowid DESC"):
                values = row[0].split() if kind == "tags" else [row[0].strip()]
                for value in values:
                    if value and value not in result: result.append(value)
                    if len(result) >= limit: return result[:limit]
        if kind == "tags" and len(result) < limit:
            for value in self.vocab("tags"):
                if value not in result: result.append(value)
        return result[:limit]

    def vocab(self, column):
        if column == "purpose":
            return [r[0] for r in self.db.execute("SELECT name FROM purposes WHERE name!='' ORDER BY name")]
        return sorted({t for table in ("nodes", "entries") for r in self.db.execute(f"SELECT tags FROM {table} WHERE deleted IS NULL") for t in r[0].split()})

    def delete(self, ids, table="nodes"):
        if table not in ("nodes", "entries"):
            raise ValueError("잘못된 삭제 대상")
        affected = self.descendants(ids) if table == "nodes" else set(ids)
        if not affected:
            return
        token = uuid.uuid4().hex
        label = f"{'장비·그룹' if table == 'nodes' else '기록'} {len(affected)}개"
        with self.atomic():
            self.db.execute("INSERT INTO trash VALUES(?,?,?)", (token, label, datetime.now().isoformat(timespec="seconds")))
            self.db.executemany(f"UPDATE {table} SET deleted=? WHERE id=? AND deleted IS NULL", [(token, i) for i in affected])

    def trash(self):
        return [dict(r) for r in self.db.execute("SELECT * FROM trash ORDER BY created DESC,rowid DESC")]

    def restore(self, token):
        with self.atomic():
            all_nodes = {n["id"]: n for n in self.nodes(True)}
            restore = {n["id"] for n in all_nodes.values() if n["deleted"] == token}
            for ident in list(restore):
                parent = all_nodes[ident]["parent"]
                while parent:
                    if all_nodes[parent]["deleted"]:
                        restore.add(parent)
                    parent = all_nodes[parent]["parent"]
            self.db.executemany("UPDATE nodes SET deleted=NULL WHERE id=?", [(i,) for i in restore])
            self.db.execute("UPDATE entries SET deleted=NULL WHERE deleted=?", (token,))
            self.db.execute("DELETE FROM trash WHERE id=?", (token,))

    def backup(self):
        folder = self.root / "backups"
        folder.mkdir(exist_ok=True)
        dest = folder / (datetime.now().strftime("%Y%m%d_%H%M%S_%f") + ".sqlite3")
        out = sqlite3.connect(dest)
        try:
            self.db.backup(out)
        finally:
            out.close()
        return dest

    def close(self):
        self.db.close()


# ============================================================================
# THEME
# ============================================================================

STYLE = """
QWidget { font-family: '맑은 고딕', 'Malgun Gothic', sans-serif; font-size: 13px; color: #26354a; }
QMainWindow, QDialog { background: #f4f6fa; }
QLabel { background: transparent; }
QLabel#Title { font-size: 27px; font-weight: 700; color: #16263c; }
QLabel#Subtitle { color: #788598; font-size: 12px; }
QLabel#Section:focus { border-bottom: 1px solid #12816f; }
QLabel#Brand:focus { border-bottom: 1px solid #b4d1da; }
QLabel#Section { font-size: 15px; font-weight: 700; }
QLabel#DeviceBadge { background: #e3f2ee; border: 1px solid #87bfae; border-radius: 8px; padding: 4px 12px; font-size: 18px; font-weight: 700; color: #126652; }
QLabel#Metric { font-size: 23px; font-weight: 700; color: #143d46; }
QFrame#Card { background: white; border: 1px solid #e3e9f0; border-radius: 12px; }
QFrame#Rail { background: #122539; border: none; }
QFrame#Rail QLabel { color: #a6b8cb; }
QFrame#Rail QLabel#Brand { font-size: 20px; color: #ffffff; font-weight: 800; }
QFrame#Rail QPushButton { background: transparent; color: #b4c6d8; border: none; text-align: left; padding: 7px 10px; border-radius: 8px; }
QFrame#Rail QPushButton:hover { background: #20384f; color: white; }
QFrame#Rail QPushButton:checked { background: #23505b; color: #c5fff1; font-weight: 700; }
QFrame#TreePanel { background: #fbfcfe; border: none; border-right: 1px solid #e2e7ef; }
QLineEdit, QComboBox, QDateEdit, QSpinBox, QTextEdit, QTextBrowser {
  background: white; border: 1px solid #dce3ec; border-radius: 7px; padding: 8px; selection-background-color: #c5eee7; selection-color: #153b39;
}
QLineEdit:focus, QComboBox:focus, QTextEdit:focus { border: 1px solid #22a48c; }
QLineEdit:disabled { color: #8c99a8; background: #eff2f6; }
QComboBox::drop-down, QDateEdit::drop-down { border: none; width: 24px; }
QComboBox QAbstractItemView { background: white; selection-background-color: #ddf4ee; padding: 4px; }
QPushButton { background: white; border: 1px solid #dce3ec; border-radius: 7px; padding: 9px 14px; font-weight: 600; }
QPushButton:hover { background: #eef8f5; border-color: #91cbbb; }
QPushButton:pressed { background: #dbefe9; }
QPushButton:disabled { color: #9eabb9; background: #edf1f5; border-color: #e2e7ed; }
QPushButton#Primary { background: #12816f; color: white; border: 1px solid #12816f; }
QPushButton#Primary:hover { background: #096e60; }
QPushButton#Primary:disabled { background: #a6cbc4; border-color: #a6cbc4; }
QPushButton#Category { padding: 5px 9px; font-size: 12px; }
QPushButton#Category:checked { background: #d9f0e9; border-color: #12816f; color: #096e60; }
QPushButton#Suggestion { background: #f2f6fa; border: 1px solid #9caec1; border-radius: 6px; padding: 3px 7px; color: #40546d; font-size: 11px; font-weight: 500; }
QPushButton#Suggestion:hover, QPushButton#Suggestion:focus { background: #d9f0e9; border-color: #12816f; color: #096e60; }
QPushButton#TagChip { background: #d9f0e9; border: 1px solid #a6d8ca; border-radius: 12px; padding: 3px 9px; color: #096e60; font-size: 12px; }
QCalendarWidget QWidget#qt_calendar_navigationbar { background: #edf5f4; }
QCalendarWidget QToolButton { color: #26354a; background: #edf5f4; }
QPushButton#Danger { color: #ba4353; }
QToolButton { background: white; border: 1px solid #e1e6ed; border-radius: 5px; padding: 6px 10px; }
QToolButton:hover, QToolButton:checked { background: #dff3ed; }
QTreeWidget, QListWidget, QTableWidget { background: white; border: 1px solid #e3e9f0; border-radius: 8px; outline: none; alternate-background-color: #f8fafc; }
QTreeWidget#AssetTree { background: transparent; border: none; font-size: 11px; }
QTreeWidget#AssetTree::item { padding: 0px 3px; height: 18px; }
QTreeWidget::item, QListWidget::item { padding: 2px 6px; border-radius: 5px; }
QTreeWidget::item:hover, QListWidget::item:hover { background: #edf5f4; }
QTreeWidget::item:selected, QListWidget::item:selected { background: #d9f0e9; color: #0a6656; }
QListWidget#Timeline { background: transparent; border: none; }
QListWidget#Timeline::item { background: white; border: 1px solid #e3e8ef; margin-bottom: 9px; padding: 0; border-radius: 10px; }
QListWidget#Timeline::item:selected { background: #eaf7f2; border-color: #69b6a3; }
QHeaderView::section { background: #f0f4f8; color: #6b7a8e; border: none; border-bottom: 1px solid #e1e7ef; padding: 10px; font-weight: 600; }
QTabWidget::pane { border: none; background: transparent; }
QTabBar::tab { background: transparent; padding: 13px 18px; color: #778397; border-bottom: 2px solid transparent; }
QTabBar::tab:selected { color: #087560; border-bottom: 2px solid #168570; font-weight: 700; }
QSplitter::handle { background: #e7ecf2; width: 1px; }
QScrollBar:vertical { background: transparent; width: 8px; margin: 2px; }
QScrollBar::handle:vertical { background: #cbd5e0; border-radius: 3px; min-height: 30px; }
QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical { height: 0; }
QScrollBar::add-page:vertical, QScrollBar::sub-page:vertical { background: transparent; }
QStatusBar { background: #edf2f6; color: #778599; border-top: 1px solid #dee5ec; }
QMenuBar { background: #f8fafc; }
QMenu { background: white; border: 1px solid #dce4ed; padding: 5px; }
QMenu::item { padding: 9px 26px; }
QMenu::item:selected { background: #e2f4ee; }
QToolTip { background: #182e42; color: white; border: none; padding: 7px; }
"""


# ============================================================================
# ICONS
# ============================================================================

PATHS = {
    "edit": '<path d="m15 5 4 4M4 20l4.5-1L20 7.5a2.8 2.8 0 0 0-4-4L4.5 15zM4.5 15l4 4"/>',
    "step-up": '<path d="M12 5v14M5 12h14"/>',
    "step-down": '<path d="M5 12h14"/>',
    "inventory": '<rect x="3" y="3" width="7" height="7" rx="1.5"/><rect x="14" y="3" width="7" height="7" rx="1.5"/><rect x="3" y="14" width="7" height="7" rx="1.5"/><rect x="14" y="14" width="7" height="7" rx="1.5"/>',
    "events": '<rect x="3" y="5" width="18" height="16" rx="3"/><path d="M7 3v4m10-4v4M3 11h18m-13 5h3"/>',
    "search": '<circle cx="10.5" cy="10.5" r="6.5"/><path d="m16 16 5 5"/>',
    "group": '<path d="M3 7V5a2 2 0 0 1 2-2h5l3 4h6a2 2 0 0 1 2 2v10a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2V7z"/>',
    "device": '<rect x="3" y="4" width="18" height="13" rx="2"/><path d="M8 21h8m-4-4v4M7 9h5m-5 4h9"/>',
    "unit": '<path d="m12 3 9 5v9l-9 5-9-5V8l9-5zM3 8l9 5 9-5m-9 5v9M7.5 5.5l9 5"/>',
    "trash": '<path d="M3 6h18M9 6V3h6v3M5 6l1 15h12l1-15M10 10v7m4-7v7"/>',
    "import": '<path d="M12 3v12m-5-5 5 5 5-5M4 16v5h16v-5"/>',
}
_cache = {}


def icon(name, color="#728da0"):
    key = (name, color)
    if key not in _cache:
        svg = f'<svg xmlns="http://www.w3.org/2000/svg" width="24" height="24" viewBox="0 0 24 24"><g fill="none" stroke="{color}" stroke-width="1.65" stroke-linecap="round" stroke-linejoin="round">{PATHS[name]}</g></svg>'
        renderer = QSvgRenderer(QByteArray(svg.encode()))
        pixmap = QPixmap(48, 48)
        pixmap.fill(Qt.GlobalColor.transparent)
        painter = QPainter(pixmap)
        renderer.render(painter)
        painter.end()
        _cache[key] = QIcon(pixmap)
    return _cache[key]


# ============================================================================
# EXCEL
# ============================================================================

HEADERS = ("장치이름", "일련번호 처음", "일련번호 끝")

# Inventory round-trip files are separate from the serial-range registration template.
LIST_HEADERS = ("구분", "소속 그룹", "장치이름", "일련번호", "용도", "현재 위치", "설명", "태그", "최근 상태", "장비 ID", "상위 ID")
LIST_FIELDS = {"name": "장치이름", "serial": "일련번호", "purpose": "용도",
               "location": "현재 위치", "description": "설명", "tags": "태그"}


@dataclass
class EquipmentListRow:
    row: int
    kind: str
    ident: str
    parent: str
    values: dict


@dataclass
class EquipmentChange:
    row: int
    ident: str
    kind: str
    parent: str | None
    before: dict | None
    after: dict

    def differences(self):
        return {key: (self.before.get(key, "") if self.before else "", self.after[key])
                for key in LIST_FIELDS if self.before is None or self.before.get(key, "") != self.after[key]}

    def caption(self):
        return self.after["name"] + (" · " + self.after["serial"] if self.kind == "unit" else "")


@dataclass
class EquipmentComparison:
    project_identity: str
    snapshot: dict
    changes: list
    errors: list
    omitted: list


def equipment_snapshot(store):
    # Resolve latest direct state for all equipment in two queries, avoiding a
    # per-node scan of an ever-growing record history during Excel round trips.
    def latest(column, condition):
        query = f"""SELECT node_id,value FROM (
            SELECT t.node_id,e.{column} AS value,
                row_number() OVER (PARTITION BY t.node_id ORDER BY e.day DESC,e.created DESC,e.rowid DESC) AS rank
            FROM entries e JOIN entry_nodes t ON t.entry_id=e.id
            WHERE e.deleted IS NULL AND e.day<=? AND {condition}) WHERE rank=1"""
        return dict(store.db.execute(query, (date.today().isoformat(),)))
    locations = latest("location", "(e.location!='' OR e.category='위치')")
    statuses = latest("category", "e.category IN ('사용','출고','입고')")
    result = {}
    for node in store.nodes():
        if node["kind"] == "group":
            continue
        location = locations.get(node["id"], node.get("location", ""))
        status = statuses.get(node["id"], "미기록")
        result[node["id"]] = {"kind": node["kind"], "parent": node["parent"],
            **{key: node.get(key, "") for key in LIST_FIELDS},
            "location": location, "status": status}
    return result


def export_equipment_list(store, path):
    """Generate a portable editable list using the Python app's installed Excel library."""
    snapshot = equipment_snapshot(store)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "장비목록"
    sheet.freeze_panes = "D2"
    sheet.sheet_view.showGridLines = False
    widths = [16, 30, 26, 24, 20, 24, 48, 30, 16, 38, 38]
    for column, (header, width) in enumerate(zip(LIST_HEADERS, widths), 1):
        letter = sheet.cell(1, column).column_letter
        sheet.column_dimensions[letter].width = width
        cell = sheet.cell(1, column, header)
        cell.font = Font(name="Malgun Gothic", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="16344A")
        cell.alignment = Alignment(vertical="center", horizontal="center")
    sheet.row_dimensions[1].height = 30
    nodes = store.nodes()
    by_id = {node["id"]: node for node in nodes}
    children = {}
    for node in nodes:
        children.setdefault(node["parent"], []).append(node)
    pending = list(reversed(children.get(None, [])))
    row_number = 2
    while pending:
        node = pending.pop()
        pending.extend(reversed(children.get(node["id"], [])))
        if node["kind"] == "group":
            continue
        values = snapshot[node["id"]]
        group = node["parent"] if node["kind"] == "device" else by_id[node["parent"]]["parent"]
        group_name = store.path_label(group, by_id) if group else "프로젝트 최상위"
        row_values = ["장치" if node["kind"] == "device" else "개별 장비", group_name,
            values["name"], values["serial"], values["purpose"], values["location"],
            values["description"], values["tags"], values["status"], node["id"], node["parent"] or ""]
        for column, value in enumerate(row_values, 1):
            cell = sheet.cell(row_number, column, value)
            cell.data_type = "s"  # Serial IDs and text beginning with '=' remain literal text.
            cell.number_format = "@"
            editable = 3 <= column <= 8
            cell.fill = PatternFill("solid", fgColor="F2F8F6" if editable else "EDF1F5")
            cell.font = Font(name="Malgun Gothic", size=11, color="26354A" if editable else "697586",
                             bold=node["kind"] == "device")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        lines = max(sum(max(1, (sum(1 if char.isascii() else 2 for char in line) + width - 1) // width)
                        for line in str(value).split("\n")) for value, width in zip(row_values, widths))
        sheet.row_dimensions[row_number].height = min(409, max(30, lines * 16 + 8))
        row_number += 1
    sheet.auto_filter.ref = f"A1:K{max(1, row_number - 1)}"
    guide = workbook.create_sheet("사용 안내")
    instructions = [
        "장비 목록 조회 및 변경 반영", "프로젝트: " + store.name,
        "녹색 열의 장치이름·일련번호·용도·현재 위치·설명·태그를 수정하세요.",
        "기존 행의 구분·장비 ID·상위 ID는 유지하세요. 소속 그룹과 최근 상태는 조회용입니다.",
        "새 일련번호: 같은 장치의 행을 복사하고 장비 ID(J열)만 비운 뒤 일련번호를 입력하세요.",
        "새 장치: 구분을 장치로 적고 장비 ID를 비우세요. 상위 ID도 비우면 비교 창에서 선택한 그룹에 등록합니다.",
        "새 장치의 일련번호: 구분은 개별 장비, 장치이름은 동일하게 입력하고 장비 ID와 상위 ID를 비우세요.",
        "일련번호·ID는 텍스트 형식을 유지하세요. 수식은 가져올 수 없습니다.",
        "태그는 #태그1 #태그2처럼 공백으로 구분합니다.",
        "도구 → 엑셀 장비 목록 비교·반영에서 변경 파일을 선택하세요.",
        "변경 전·후 값을 확인하고 적용 날짜를 선택한 다음 반영 버튼을 누르세요.",
        "누락된 행은 삭제하지 않습니다. 반영 직전에 자동 백업하고 변경된 장비에 날짜별 기록을 남깁니다.",
        "용도 등 기본 속성은 반영 즉시 바뀝니다. 현재 위치는 오늘까지의 위치 기록으로 계산합니다."]
    guide.column_dimensions["A"].width = 115
    for index, value in enumerate(instructions, 1):
        cell = guide.cell(index, 1, value)
        cell.font = Font(name="Malgun Gothic", size=14 if index == 1 else 11, bold=index == 1, color="16344A")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        guide.row_dimensions[index].height = 36
    guide.sheet_view.showGridLines = False
    metadata = workbook.create_sheet("SEMP 정보")
    metadata.append(["SEMP-equipment-list", "1"])
    metadata.append(["project_identity", store.project_identity])
    metadata.sheet_state = "hidden"
    workbook.save(path)
    workbook.close()
    return Path(path)


def read_equipment_list(path, store):
    workbook = load_workbook(path, read_only=True, data_only=False)
    iterator = None
    try:
        if "장비목록" not in workbook.sheetnames or "SEMP 정보" not in workbook.sheetnames:
            raise ValueError("장비 목록 내보내기로 생성한 엑셀 파일을 사용하세요.")
        metadata = workbook["SEMP 정보"]
        if metadata["A1"].value != "SEMP-equipment-list" or metadata["B1"].value != "1":
            raise ValueError("지원하지 않는 장비 목록 형식입니다.")
        if metadata["B2"].value != store.project_identity:
            raise ValueError("다른 프로젝트에서 내보낸 장비 목록입니다. 해당 프로젝트를 열어주세요.")
        sheet = workbook["장비목록"]
        iterator = sheet.iter_rows(max_col=len(LIST_HEADERS))
        if tuple(cell_text(cell) for cell in next(iterator, [])) != LIST_HEADERS:
            raise ValueError("장비 목록의 열 제목을 유지하세요.")
        rows, errors = [], []
        for number, cells in enumerate(iterator, 2):
            if all(cell.value is None for cell in cells):
                continue
            try:
                values = [cell_text(cell) for cell in cells]
                kind = {"장치": "device", "개별 장비": "unit", "일련번호": "unit"}.get(values[0])
                if not kind:
                    raise ValueError("구분은 장치 또는 개별 장비로 입력하세요.")
                fields = dict(zip(LIST_FIELDS, values[2:8]))
                fields["tags"] = tags(fields["tags"])
                if not fields["name"]:
                    raise ValueError("장치이름이 필요합니다.")
                if kind == "unit" and not fields["serial"]:
                    raise ValueError("개별 장비는 일련번호가 필요합니다.")
                if kind == "device" and fields["serial"]:
                    raise ValueError("장치 행은 일련번호를 비우세요.")
                rows.append(EquipmentListRow(number, kind, values[9], values[10], fields))
            except ValueError as exc:
                errors.append(f"{number}행: {exc}")
        return rows, errors
    finally:
        if iterator is not None:
            iterator.close()
        workbook.close()


def compare_equipment_list(store, rows, source_errors=(), new_device_parent=None):
    snapshot = equipment_snapshot(store)
    store.validate_parent("device", new_device_parent)
    changes, errors, present, resolved = [], list(source_errors), set(), []
    devices = {ident: values for ident, values in snapshot.items() if values["kind"] == "device"}
    device_keys = {(value["parent"], value["name"]): ident for ident, value in devices.items()}
    unit_keys = {(value["parent"], value["serial"]): ident for ident, value in snapshot.items() if value["kind"] == "unit"}
    device_names = {}
    # Resolve device rows first so new units can refer to new devices regardless of row order.
    for row in [row for row in rows if row.kind == "device"]:
        parent = row.parent or new_device_parent
        ident = row.ident
        if not ident:
            ident = device_keys.get((parent, row.values["name"])) or uuid.uuid4().hex
        before = snapshot.get(ident)
        if row.ident and (before is None or before["kind"] != row.kind):
            errors.append(f"{row.row}행: 장비 ID가 없거나 구분이 다릅니다.")
            continue
        if before and row.parent != (before["parent"] or ""):
            errors.append(f"{row.row}행: 기존 장비의 상위 ID는 변경하지 마세요.")
            continue
        if before:
            parent = before["parent"]
        try:
            store.validate_parent("device", parent)
        except ValueError as exc:
            errors.append(f"{row.row}행: {exc}")
            continue
        resolved.append((row, ident, parent, before))
        devices[ident] = {"kind": "device", "parent": parent, **row.values}
    for ident, value in devices.items():
        device_names.setdefault(value["name"], set()).add(ident)
        if ident in snapshot:
            device_names.setdefault(snapshot[ident]["name"], set()).add(ident)
    for row in [row for row in rows if row.kind == "unit"]:
        before = snapshot.get(row.ident) if row.ident else None
        if row.ident and (before is None or before["kind"] != "unit"):
            errors.append(f"{row.row}행: 장비 ID가 없거나 구분이 다릅니다.")
            continue
        parent = row.parent
        if not parent:
            candidates = device_names.get(row.values["name"], set())
            if len(candidates) != 1:
                errors.append(f"{row.row}행: 상위 장치를 특정할 수 없습니다. 기존 행의 상위 ID를 복사하세요.")
                continue
            parent = next(iter(candidates))
        if parent not in devices:
            errors.append(f"{row.row}행: 상위 장치가 없습니다.")
            continue
        allowed_names = {devices[parent]["name"], snapshot.get(parent, {}).get("name", "")}
        if row.values["name"] not in allowed_names:
            errors.append(f"{row.row}행: 장치 이름은 상위 장치 행에서 수정하세요.")
            continue
        if before and before["parent"] != parent:
            errors.append(f"{row.row}행: 기존 장비의 상위 ID는 변경하지 마세요.")
            continue
        ident = row.ident
        if not ident:
            ident = unit_keys.get((parent, row.values["serial"])) or uuid.uuid4().hex
            before = snapshot.get(ident)
        values = {**row.values, "name": devices[parent]["name"]}
        resolved.append((EquipmentListRow(row.row, row.kind, row.ident, parent, values), ident, parent, before))
    final = {ident: dict(values) for ident, values in snapshot.items()}
    for row, ident, parent, before in resolved:
        if ident in present:
            errors.append(f"{row.row}행: 같은 장비가 여러 행에 있습니다. 새 장비는 장비 ID를 비우세요.")
            continue
        present.add(ident)
        after = dict(row.values)
        effective_before = dict(before) if before else None
        if row.kind == "unit" and effective_before:
            effective_before["name"] = after["name"]  # Renaming a device also renames its displayed units.
        change = EquipmentChange(row.row, ident, row.kind, parent, effective_before, after)
        if change.differences():
            changes.append(change)
        final[ident] = {"kind": row.kind, "parent": parent, **after}
    keys = {}
    for ident, values in final.items():
        key = (values["parent"], values["kind"], values["serial"] if values["kind"] == "unit" else values["name"])
        if key in keys:
            errors.append("반영 후 같은 위치에 장치 이름 또는 일련번호가 중복됩니다: " + key[2])
        keys[key] = ident
    omitted = [ident for ident in snapshot if ident not in present]
    return EquipmentComparison(store.project_identity, snapshot, changes, list(dict.fromkeys(errors)), omitted)


def apply_equipment_changes(store, comparison, day):
    date.fromisoformat(day)
    if comparison.project_identity != store.project_identity:
        raise ValueError("비교한 프로젝트가 변경되었습니다. 다시 비교하세요.")
    if comparison.errors:
        raise ValueError("엑셀 오류를 수정한 뒤 다시 비교하세요.")
    if not comparison.changes:
        return 0, None
    if equipment_snapshot(store) != comparison.snapshot:
        raise ValueError("비교 이후 장비 정보나 위치가 바뀌었습니다. 파일을 다시 비교하세요.")
    backup = store.backup()
    with store.atomic():
        # Temporarily free unique keys so legitimate name/serial swaps can be applied together.
        for change in comparison.changes:
            if change.before:
                if change.kind == "device":
                    store.db.execute("UPDATE nodes SET name=? WHERE id=?", ("__SEMP_SYNC_" + uuid.uuid4().hex, change.ident))
                else:
                    store.db.execute("UPDATE nodes SET serial=? WHERE id=?", ("__SEMP_SYNC_" + uuid.uuid4().hex, change.ident))
        changes = sorted(comparison.changes, key=lambda change: change.kind == "unit")
        created_ids = {}
        for change in changes:
            after = change.after
            parent = created_ids.get(change.parent, change.parent)
            if change.before is None:
                ident = store.add_node(change.kind, after["name"], parent, serial=after["serial"],
                    purpose=after["purpose"], description=after["description"], tag_text=after["tags"])
                created_ids[change.ident] = ident
            else:
                ident = change.ident
                store.update_node(ident, after["name"], after["serial"], after["purpose"], after["description"], after["tags"])
            differences = change.differences()
            if change.before is None:
                body = "엑셀 장비 목록에서 추가\n" + "\n".join(f"{LIST_FIELDS[key]}: {value or '미지정'}" for key, value in after.items())
                store.save_entry(title=change.caption() + " 등록", category="장비등록", day=day,
                    html="<p>" + html.escape(body).replace("\n", "<br>") + "</p>", plain=body,
                    node_ids=[ident], event=True, location=after["location"])
            else:
                for key, (before, value) in differences.items():
                    body = f"엑셀 장비 목록 반영\n{LIST_FIELDS[key]}: {before or '미지정'} → {value or '미지정'}"
                    category = "용도변경" if key == "purpose" else "위치" if key == "location" else "장비정보변경"
                    store.save_entry(title=change.caption() + " " + LIST_FIELDS[key] + " 변경", category=category,
                        day=day, html="<p>" + html.escape(body).replace("\n", "<br>") + "</p>", plain=body,
                        node_ids=[ident], event=True, location=value if key == "location" else "",
                        purpose_after=value if key == "purpose" else None)
    return len(changes), backup


@dataclass
class ImportRow:
    row: int
    name: str
    first: str
    last: str
    serials: list[str]
    error: str = ""
    purpose: str = ""


def cell_text(cell):
    if cell.value is None:
        return ""
    if cell.data_type == "f":
        raise ValueError("수식 대신 텍스트 값을 입력하세요.")
    value = cell.value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if int(value) != value:
            raise ValueError("일련번호 숫자는 정수 또는 텍스트로 입력하세요.")
        fmt = cell.number_format
        if len(str(abs(int(value)))) > 15:
            raise ValueError("15자리를 넘는 일련번호는 텍스트 형식으로 입력하세요.")
        return str(int(value)).zfill(len(fmt)) if fmt and set(fmt) == {"0"} else str(int(value))
    if not isinstance(value, str):
        raise ValueError("날짜나 논리값 대신 텍스트를 입력하세요.")
    return value.strip()


def read_import(path: str | Path):
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        ws = wb["장비등록"] if "장비등록" in wb.sheetnames else wb.worksheets[0]
        iterator = ws.iter_rows(max_col=4)
        first_row = next(iterator, None)
        if not first_row or tuple(cell_text(c) for c in first_row[:3]) != HEADERS:
            raise ValueError("첫 행은 '장치이름 / 일련번호 처음 / 일련번호 끝'이어야 합니다. 제공 템플릿을 사용하세요.")
        has_purpose = cell_text(first_row[3]) == "용도"
        if cell_text(first_row[3]) not in ("", "용도"):
            raise ValueError("네 번째 열 제목은 용도여야 합니다.")
        result, seen, total = [], set(), 0
        for number, cells in enumerate(iterator, 2):
            if number > 20001:
                raise ValueError("한 파일에 최대 20,000행까지 읽을 수 있습니다.")
            cells = cells if has_purpose else cells[:3]
            if all(c.value is None for c in cells):
                continue
            name = first = last = purpose = ""
            serials, error = [], ""
            try:
                name, first, last = [cell_text(c) for c in cells[:3]]
                purpose = cell_text(cells[3]) if has_purpose else ""
                if not name:
                    raise ValueError("장치이름이 필요합니다.")
                serials = serial_range(first, last)
                if purpose and not serials:
                    raise ValueError("용도를 지정하려면 일련번호를 입력하세요. 장치만 등록할 때는 용도를 비워두세요.")
                if any((name, s) in seen for s in serials):
                    raise ValueError("파일 안에 같은 장치의 중복 일련번호가 있습니다.")
                if total + len(serials) > 20000:
                    raise ValueError("파일 전체에서 최대 20,000개의 일련번호까지 미리 볼 수 있습니다.")
                seen.update((name, s) for s in serials)
                total += len(serials)
            except ValueError as exc:
                serials, error = [], str(exc)
            result.append(ImportRow(number, name, first, last, serials, error, purpose))
        return result
    finally:
        wb.close()


# ============================================================================
# WIDGETS
# ============================================================================

NODE_MIME = "application/x-semp-nodes"
IMPORT_MIME = "application/x-semp-import"


class ElidedLabel(QLabel):
    def __init__(self, text):
        super().__init__(text)
        self.setMinimumWidth(0)
        self.setSizePolicy(QSizePolicy.Policy.Ignored, QSizePolicy.Policy.Preferred)
    def paintEvent(self, event):
        if self.wordWrap(): return super().paintEvent(event)
        painter = QPainter(self)
        painter.setPen(self.palette().color(QPalette.ColorRole.WindowText))
        text = self.fontMetrics().elidedText(self.text(), Qt.TextElideMode.ElideRight, self.contentsRect().width())
        painter.drawText(self.contentsRect(), self.alignment(), text)
    def setText(self, text):
        super().setText(text)
        self.setToolTip(text)


class HomeLink(ElidedLabel):
    clicked = Signal()

    def __init__(self, text):
        super().__init__(text)
        self.setTextFormat(Qt.TextFormat.PlainText)
        self.setCursor(Qt.CursorShape.PointingHandCursor)
        self.setFocusPolicy(Qt.FocusPolicy.StrongFocus)
        self.setText(text)

    def setText(self, text):
        super().setText(text)
        self.setToolTip(text + " · 달력 메인 화면으로 이동")
        self.setAccessibleName(text + " · 달력 메인 화면으로 이동")

    def mouseReleaseEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton and self.rect().contains(event.position().toPoint()):
            self.clicked.emit()
        else: super().mouseReleaseEvent(event)

    def keyPressEvent(self, event):
        if event.key() in (Qt.Key.Key_Return, Qt.Key.Key_Enter, Qt.Key.Key_Space):
            self.clicked.emit()
        else: super().keyPressEvent(event)


class DoubleClickLabel(QLabel):
    doubleClicked = Signal()

    def mouseDoubleClickEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            self.doubleClicked.emit()
            event.accept()
        else:
            super().mouseDoubleClickEvent(event)


def label(text, role=None, elide=False):
    w = ElidedLabel(text) if elide else QLabel(text)
    w.setTextFormat(Qt.TextFormat.PlainText)
    if role:
        w.setObjectName(role)
    return w


def button(text, callback, primary=False):
    w = QPushButton(text)
    w.setAutoDefault(False)
    if primary:
        w.setObjectName("Primary")
        w.setDefault(True)
    w.clicked.connect(callback)
    return w


def row(*widgets):
    layout = QHBoxLayout()
    layout.setSpacing(10)
    for w in widgets:
        if w is None:
            layout.addStretch()
        else:
            layout.addWidget(w)
    return layout


def node_caption(node):
    if node["kind"] != "unit": return node["name"]
    name, serial = node["name"], node["serial"]
    return serial if serial.casefold().startswith(name.casefold() + "-") else f"{name}-{serial}"


class FlowLayout(QLayout):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.items = []
        self.setContentsMargins(0, 0, 0, 0)
        self.setSpacing(5)
    def addItem(self, item): self.items.append(item)
    def count(self): return len(self.items)
    def itemAt(self, index): return self.items[index] if index < len(self.items) else None
    def takeAt(self, index): return self.items.pop(index) if index < len(self.items) else None
    def expandingDirections(self): return Qt.Orientation(0)
    def hasHeightForWidth(self): return True
    def heightForWidth(self, width): return self.arrange(QRect(0, 0, width, 0), True)
    def setGeometry(self, rect):
        super().setGeometry(rect)
        self.arrange(rect, False)
    def sizeHint(self): return self.minimumSize()
    def minimumSize(self):
        size = QSize()
        for item in self.items: size = size.expandedTo(item.minimumSize())
        return size
    def arrange(self, rect, measure):
        x, y, height = rect.x(), rect.y(), 0
        for item in self.items:
            size = item.sizeHint()
            if x > rect.x() and x + size.width() > rect.right() + 1:
                x, y, height = rect.x(), y + height + self.spacing(), 0
            if not measure: item.setGeometry(QRect(QPoint(x, y), size))
            x += size.width() + self.spacing()
            height = max(height, size.height())
        return y + height - rect.y()


class CategoryButtons(QWidget):
    changed = Signal(str)

    def __init__(self, choices, selected, title="기록 종류 추가", prompt="새 종류"):
        super().__init__()
        self.prompt_title, self.prompt_label = title, prompt
        self.flow = FlowLayout(self)
        self.group = QButtonGroup(self)
        self.choices = {}
        for choice in choices: self.add_choice(choice)
        self.add_button = button("＋ 추가", self.prompt_choice)
        self.flow.addWidget(self.add_button)
        self.setCurrentText(selected)
    def add_choice(self, text):
        if text in self.choices: return
        b = button(text, lambda: None)
        b.setCheckable(True)
        b.setObjectName("Category")
        b.clicked.connect(lambda checked=False, value=text: self.changed.emit(value))
        self.group.addButton(b)
        self.choices[text] = b
        if hasattr(self, "add_button"): self.flow.removeWidget(self.add_button)
        self.flow.addWidget(b)
        if hasattr(self, "add_button"): self.flow.addWidget(self.add_button)
    def setCurrentText(self, text):
        self.add_choice(text)
        self.choices[text].setChecked(True)
    def currentText(self):
        b = self.group.checkedButton()
        return b.text() if b else "기타"
    def prompt_choice(self):
        text, ok = QInputDialog.getText(self, self.prompt_title, self.prompt_label)
        if ok and text.strip(): self.setCurrentText(text.strip())


class PurposeButtons(CategoryButtons):
    DEFAULTS = ("비행용", "예비용", "환경용", "점검용")

    def __init__(self, vocabulary=(), selected=""):
        choices = list(dict.fromkeys([*self.DEFAULTS, *(value for value in vocabulary if value)]))
        super().__init__(choices, selected, title="용도 추가", prompt="새 용도")

    def setCurrentText(self, text):
        if text:
            super().setCurrentText(text)
        else:
            self.group.setExclusive(False)
            for choice in self.choices.values(): choice.setChecked(False)
            self.group.setExclusive(True)

    def currentText(self):
        selected = self.group.checkedButton()
        return selected.text() if selected else ""


def equipment_ranks(nodes):
    ranks, siblings = {}, {}
    for node in nodes:
        if node["kind"] == "unit":
            siblings.setdefault(node["parent"], []).append(node)
    def natural(value):
        return [(0, int(part)) if part.isdigit() else (1, part.casefold())
                for part in re.split(r"(\d+)", value)]
    def sequence(serial):
        values = []
        for part in re.findall(r"\d+", serial):
            if len(part) == 4 and 1900 <= int(part) <= 2099:
                continue
            if len(part) > 4 and 1900 <= int(part[:4]) <= 2099:
                values.append(int(part[4:]))
            elif len(part) == 5 and 20 <= int(part[:2]) <= 99:
                values.append(int(part[-3:]))
            else:
                values.append(int(part))
        return values[-1] if values else None
    for units in siblings.values():
        for rank, node in enumerate(sorted(units, key=lambda item: natural(item["serial"])), 1):
            value = sequence(node["serial"])
            # Years can be separate or joined to the sequence. Thus 2024-A001,
            # A001-2025 and T-26001 all resolve to sequence number 1.
            ranks[node["id"]] = value if value is not None else rank
    return ranks


def matching_equipment(nodes, query="", purpose="", ordinal=0, purposes=None, ordinals=None, kinds=None):
    by_id = {n["id"]: n for n in nodes}
    ranks = equipment_ranks(nodes)
    selected_purposes = set(purposes or ([purpose] if purpose else []))
    selected_ordinals = set(ordinals or ([ordinal] if ordinal else []))
    selected_kinds = set(kinds or [])
    matches = set()
    for n in nodes:
        if selected_purposes and (n["purpose"] or "미지정") not in selected_purposes: continue
        if selected_ordinals and ranks.get(n["id"]) not in selected_ordinals: continue
        if selected_kinds and n["kind"] not in selected_kinds: continue
        text = " ".join(str(n.get(key, "")) for key in ("name", "serial", "purpose", "description", "tags"))
        parent = n["parent"]
        while parent in by_id:
            text += " " + by_id[parent]["name"]
            parent = by_id[parent]["parent"]
        if not query or query.casefold() in text.casefold(): matches.add(n["id"])
    return matches


def with_ancestors(nodes, matches):
    by_id = {n["id"]: n for n in nodes}
    visible = set(matches)
    for ident in matches:
        parent = by_id[ident]["parent"]
        while parent in by_id:
            visible.add(parent)
            parent = by_id[parent]["parent"]
    return visible


class VisibleSpinBox(QSpinBox):
    """Explicit step controls remain visible with the application's stylesheet."""
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setButtonSymbols(QSpinBox.ButtonSymbols.NoButtons)
        self.setStyleSheet("QSpinBox { padding-right: 29px; }")
        self.up_button = QToolButton(self)
        self.down_button = QToolButton(self)
        for control, text, step, title in (
                (self.up_button, "+", 1, "순번 증가"),
                (self.down_button, "−", -1, "순번 감소")):
            control.setIcon(icon("step-up" if step > 0 else "step-down", "#126652"))
            control.setIconSize(QSize(12, 12))
            control.setAccessibleName(title)
            control.setToolTip(title)
            control.setFocusPolicy(Qt.FocusPolicy.NoFocus)
            control.setAutoRepeat(True)
            control.setStyleSheet("QToolButton { padding: 0; border: none; border-radius: 3px; background: #e7f2ef; color: #126652; font-size: 14px; font-weight: 700; } QToolButton:hover { background: #c5e6dc; } QToolButton:disabled { color: #a2b1ad; background: #f0f4f3; }")
            control.clicked.connect(lambda checked=False, delta=step: self.stepBy(delta))
        self.valueChanged.connect(self.update_step_buttons)
        self.update_step_buttons()

    def update_step_buttons(self):
        self.up_button.setEnabled(self.value() < self.maximum() or self.wrapping())
        self.down_button.setEnabled(self.value() > self.minimum() or self.wrapping())

    def resizeEvent(self, event):
        super().resizeEvent(event)
        height = max(1, (self.height() - 6) // 2)
        self.up_button.setGeometry(self.width() - 26, 2, 23, height)
        self.down_button.setGeometry(self.width() - 26, height + 3, 23, height)
        self.update_step_buttons()


class EquipmentFilters(QWidget):
    changed = Signal()
    KIND_LABELS = {"group": "그룹", "device": "장치", "unit": "개별 장비"}

    def __init__(self, nodes=(), detailed=False):
        super().__init__()
        self.detailed = detailed
        self.nodes = list(nodes)
        self.purpose_buttons = {}
        self.ordinal_buttons = {}
        self.kind_buttons = {}
        if detailed:
            self.build_detailed()
            self.refresh(nodes)
            return
        layout = QHBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(5)
        self.purpose = QComboBox()
        self.purpose.setMinimumWidth(80)
        self.ordinal = VisibleSpinBox()
        self.ordinal.setRange(0, 99999)
        self.ordinal.setSpecialValueText("전체 순번")
        self.ordinal.setSuffix("번째")
        self.ordinal.setToolTip("1900~2099의 연도 숫자를 제외하고 마지막 숫자 묶음을 순번으로 사용합니다. 값이 1이면 모두 첫 번째입니다.")
        layout.addWidget(self.purpose, 1)
        layout.addWidget(self.ordinal, 1)
        self.refresh(nodes)
        self.purpose.currentIndexChanged.connect(self.changed)
        self.ordinal.valueChanged.connect(self.changed)

    def build_detailed(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(7)
        layout.addWidget(label("용도 · 여러 개 선택 가능", "Subtitle"))
        self.purpose_area = QWidget()
        self.purpose_flow = FlowLayout(self.purpose_area)
        self.purpose_flow.setSpacing(4)
        layout.addWidget(self.purpose_area)
        layout.addWidget(label("연도를 제외한 마지막 숫자 기준", "Subtitle"))
        self.ordinal_area = QWidget()
        self.ordinal_flow = FlowLayout(self.ordinal_area)
        self.ordinal_flow.setSpacing(4)
        layout.addWidget(self.ordinal_area)
        self.ordinal_input = VisibleSpinBox()
        self.ordinal_input.setRange(1, 99999)
        self.ordinal_input.setPrefix("순번 ")
        self.ordinal_input.setSuffix("번째")
        layout.addLayout(row(self.ordinal_input, button("＋ 순번 추가", self.add_ordinal)))
        layout.addWidget(label("구분", "Subtitle"))
        self.kind_area = QWidget()
        self.kind_flow = FlowLayout(self.kind_area)
        self.kind_flow.setSpacing(4)
        layout.addWidget(self.kind_area)
        for kind, text in self.KIND_LABELS.items():
            self.kind_buttons[kind] = self.add_filter_button(self.kind_flow, kind, text)
        layout.addWidget(button("필터 초기화", self.reset))

    def add_filter_button(self, flow, value, text):
        control = button(text, lambda: None)
        control.setCheckable(True)
        control.setObjectName("Category")
        control.setProperty("filter_value", value)
        control.clicked.connect(self.changed)
        flow.addWidget(control)
        return control

    def add_ordinal(self):
        value = self.ordinal_input.value()
        if value not in self.ordinal_buttons:
            self.ordinal_buttons[value] = self.add_filter_button(self.ordinal_flow, value, f"{value}번째")
            self.update_counts()
        self.ordinal_buttons[value].setChecked(True)
        self.changed.emit()

    @staticmethod
    def selected(buttons):
        return {value for value, control in buttons.items() if control.isChecked()}

    def refresh(self, nodes):
        self.nodes = list(nodes)
        if self.detailed:
            selected = self.selected(self.purpose_buttons)
            values = list(dict.fromkeys([*PurposeButtons.DEFAULTS,
                                         *(n["purpose"] for n in nodes if n["purpose"]), "미지정"]))
            for value in values:
                if value not in self.purpose_buttons:
                    self.purpose_buttons[value] = self.add_filter_button(self.purpose_flow, value, value)
                self.purpose_buttons[value].setChecked(value in selected)
            for value in (1, 2, 3):
                if value not in self.ordinal_buttons:
                    self.ordinal_buttons[value] = self.add_filter_button(self.ordinal_flow, value, f"{value}번째")
            self.update_counts()
            return
        old = self.purpose.currentData()
        self.purpose.blockSignals(True)
        self.purpose.clear()
        self.purpose.addItem("전체 용도", "")
        for value in dict.fromkeys([*PurposeButtons.DEFAULTS, *(n["purpose"] for n in nodes if n["purpose"])]):
            self.purpose.addItem(value, value)
        self.purpose.setCurrentIndex(max(0, self.purpose.findData(old)))
        self.purpose.blockSignals(False)

    def update_counts(self):
        ranks = equipment_ranks(self.nodes)
        for value, control in self.purpose_buttons.items():
            count = sum(n["kind"] == "unit" and (n["purpose"] or "미지정") == value for n in self.nodes)
            control.setText(f"{value} {count}")
        for value, control in self.ordinal_buttons.items():
            count = sum(rank == value for rank in ranks.values())
            control.setText(f"{value}번째 {count}")
        for value, control in self.kind_buttons.items():
            count = sum(n["kind"] == value for n in self.nodes)
            control.setText(f"{self.KIND_LABELS[value]} {count}")

    def matches(self, nodes, query=""):
        if self.detailed:
            purposes = self.selected(self.purpose_buttons)
            ordinals = self.selected(self.ordinal_buttons)
            kinds = self.selected(self.kind_buttons)
            base = matching_equipment(nodes, query.strip(), purposes=purposes, ordinals=ordinals)
            if not kinds:
                return base
            by_parent = {}
            by_id = {node["id"]: node for node in nodes}
            for node in nodes:
                by_parent.setdefault(node["parent"], []).append(node["id"])
            result = set()
            for node in nodes:
                if node["kind"] not in kinds:
                    continue
                pending = [node["id"]]
                while pending:
                    ident = pending.pop()
                    children = by_parent.get(ident, [])
                    if children:
                        pending.extend(children)
                    elif ident in base and by_id[ident]["kind"] in ("device", "unit"):
                        result.add(ident)
            return result
        return matching_equipment(nodes, query.strip(), self.purpose.currentData() or "", self.ordinal.value())

    def active(self):
        if self.detailed:
            return bool(self.selected(self.purpose_buttons) or self.selected(self.ordinal_buttons) or
                        self.selected(self.kind_buttons))
        return bool(self.purpose.currentData() or self.ordinal.value())

    def reset(self):
        if self.detailed:
            for control in (*self.purpose_buttons.values(), *self.ordinal_buttons.values(), *self.kind_buttons.values()):
                control.setChecked(False)
            self.changed.emit()
            return
        self.purpose.setCurrentIndex(0)
        self.ordinal.setValue(0)


class RecentProjectCard(QWidget):
    def __init__(self, name, path, exists=True):
        super().__init__()
        layout = QVBoxLayout(self)
        layout.setContentsMargins(10, 5, 10, 5)
        layout.setSpacing(2)
        title = label(name, "Section", elide=True)
        address = label(path + ("" if exists else " · 파일 없음"), "Subtitle", elide=True)
        layout.addWidget(title)
        layout.addWidget(address)
        self.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
        self.setEnabled(exists)


class SuggestionRow(QWidget):
    def __init__(self, values, callback, title="최근 사용"):
        super().__init__()
        flow = FlowLayout(self)
        caption = label(title, "Subtitle")
        caption.setStyleSheet("font-size: 11px; color: #788598;")
        flow.addWidget(caption)
        self.buttons = []
        for value in values[:5]:
            b = button(value if len(value) <= 18 else value[:18] + "…", lambda checked=False, v=value: callback(v))
            b.setObjectName("Suggestion")
            b.setToolTip(value)
            b.setAccessibleName(value)
            flow.addWidget(b)
            self.buttons.append(b)
        if not values: flow.addWidget(label("아직 없습니다", "Subtitle"))


class TagEdit(QWidget):
    def __init__(self, vocabulary=(), parent=None, recent=()):
        super().__init__(parent)
        self.tokens = []
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(5)
        self.chips = QWidget()
        self.flow = FlowLayout(self.chips)
        layout.addWidget(self.chips)
        self.input = QLineEdit()
        self.input.setPlaceholderText("#해시태그 입력 후 공백 · 태그의 ×로 삭제")
        layout.addWidget(self.input)
        self.suggestions = SuggestionRow(recent, self.add_suggestion, "최근 태그")
        layout.addWidget(self.suggestions)
        self.popup = QCompleter(list(vocabulary), self.input)
        self.popup.setCaseSensitivity(Qt.CaseSensitivity.CaseInsensitive)
        self.input.setCompleter(self.popup)
        self.input.textEdited.connect(self.commit_delimited)
        self.popup.activated.connect(self.complete_tag)
        self.setFocusProxy(self.input)
    def add_suggestion(self, text):
        self.setText(self.text() + " " + text)
        self.input.setFocus()

    def commit_delimited(self, text):
        if text and (text[-1].isspace() or text[-1] == ','):
            self.setText(self.text())
    def complete_tag(self, text):
        self.setText(" ".join(self.tokens + [text]))
    def text(self):
        return tags(" ".join(self.tokens + [self.input.text()]))
    def setText(self, text):
        self.tokens = tags(text).split()
        self.input.clear()
        self.render_chips()
    def remove_tag(self, token):
        self.tokens.remove(token)
        self.render_chips()
    def render_chips(self):
        while self.flow.count():
            item = self.flow.takeAt(0)
            item.widget().hide()
            item.widget().deleteLater()
        for token in self.tokens:
            b = button((token[:24] + "…" if len(token) > 24 else token) + " ×", lambda checked=False, t=token: self.remove_tag(t))
            b.setObjectName("TagChip")
            b.setToolTip(token + " · 클릭하여 삭제")
            self.flow.addWidget(b)
        self.chips.setVisible(bool(self.tokens))
        self.flow.invalidate()


class EventCalendar(QCalendarWidget):
    def __init__(self):
        super().__init__()
        self.events_by_day = {}
        self.setGridVisible(True)
        self.setVerticalHeaderFormat(QCalendarWidget.VerticalHeaderFormat.NoVerticalHeader)
        self.setFirstDayOfWeek(Qt.DayOfWeek.Monday)
        self.setMinimumHeight(260)
        self.setToolTip("날짜 안의 작은 숫자는 해당 날짜의 기록 수입니다.")
    def paintCell(self, painter, rect, day):
        entries = self.events_by_day.get(day.toString("yyyy-MM-dd"), [])
        painter.save()
        selected = day == self.selectedDate()
        current_month = day.month() == self.monthShown() and day.year() == self.yearShown()
        painter.fillRect(rect, QColor("#d9f0e9" if selected else "#ffffff" if current_month else "#f7f9fb"))
        painter.setPen(QColor("#e2e7ef"))
        painter.drawRect(rect.adjusted(0, 0, -1, -1))
        painter.setPen(QColor("#26354a" if current_month else "#9da8b5"))
        font = painter.font()
        font.setPixelSize(12)
        font.setBold(selected)
        painter.setFont(font)
        painter.drawText(rect.adjusted(7, 2, -5, -2), Qt.AlignmentFlag.AlignTop | Qt.AlignmentFlag.AlignLeft, str(day.day()))
        if entries:
            font.setPixelSize(10)
            font.setBold(True)
            painter.setFont(font)
            count = str(len(entries))
            width = max(18, painter.fontMetrics().horizontalAdvance(count) + 10)
            badge = QRect(rect.right() - width - 4, rect.bottom() - 18, width, 16)
            painter.setPen(Qt.PenStyle.NoPen)
            painter.setBrush(QColor("#c6e9de"))
            painter.drawRoundedRect(badge, 8, 8)
            painter.setPen(QColor("#08634f"))
            painter.drawText(badge, Qt.AlignmentFlag.AlignCenter, count)
        painter.restore()


class PortableTextEdit(QTextEdit):
    def insertFromMimeData(self, source):
        if source.hasHtml():
            # Notes support text formatting; linked/pasted images are intentionally omitted.
            html = re.sub(r"<(img|object|iframe)\b[^>]*>.*?</\1>|<img\b[^>]*>", "", source.html(), flags=re.I | re.S)
            self.insertHtml(html)
        elif source.hasText():
            self.insertPlainText(source.text())


class ColorPaletteDialog(QDialog):
    # Hue columns and separate text/background shades follow the organization in
    # https://atlassian.design/foundations/color/color-picker-swatches
    HUES = ("회색", "빨강", "주황", "초록", "청록", "파랑", "보라", "분홍")
    TEXT_ROWS = (
        ("#26354a", "#7f1d1d", "#7c2d12", "#14532d", "#134e4a", "#1e3a8a", "#581c87", "#831843"),
        ("#596579", "#b91c1c", "#9a3412", "#166534", "#0f766e", "#1d4ed8", "#7e22ce", "#be185d"),
    )
    HIGHLIGHT_ROWS = (
        ("#f1f5f9", "#fee2e2", "#ffedd5", "#dcfce7", "#ccfbf1", "#dbeafe", "#f3e8ff", "#fce7f3"),
        ("#e2e8f0", "#fecaca", "#fed7aa", "#bbf7d0", "#99f6e4", "#bfdbfe", "#e9d5ff", "#fbcfe8"),
        ("#cbd5e1", "#fca5a5", "#fdba74", "#86efac", "#5eead4", "#93c5fd", "#d8b4fe", "#f9a8d4"),
    )

    def __init__(self, kind, current, parent=None):
        super().__init__(parent)
        self.kind = kind
        self.selected_color = QColor(current)
        self.setWindowTitle("글자색" if kind == "color" else "강조색")
        self.setFixedWidth(440)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(22, 20, 22, 20)
        layout.setSpacing(14)
        layout.addWidget(label(self.windowTitle(), "Section"))
        layout.addWidget(label("색상을 누르면 선택한 글자에 바로 적용합니다.", "Subtitle"))
        self.swatches = []
        grid = QGridLayout()
        grid.setSpacing(8)
        for col, name in enumerate(self.HUES):
            header = label(name, "Subtitle")
            header.setAlignment(Qt.AlignmentFlag.AlignCenter)
            grid.addWidget(header, 0, col)
        rows = self.TEXT_ROWS if kind == "color" else self.HIGHLIGHT_ROWS
        for r, colors in enumerate(rows):
            for col, value in enumerate(colors):
                swatch = button("", lambda checked=False, v=value: self.choose(v))
                swatch.setFixedSize(38, 32)
                swatch.setCheckable(True)
                swatch.setChecked(QColor(value) == current)
                swatch.setStyleSheet(f"QPushButton {{ background: {value}; border: 1px solid #cbd5e1; border-radius: 5px; padding: 0; }} QPushButton:hover, QPushButton:focus, QPushButton:checked {{ border: 3px solid #12816f; }}")
                swatch.setToolTip(f"{self.HUES[col]} · {r + 1}단계 · {value.upper()}")
                swatch.setAccessibleName(swatch.toolTip())
                swatch.setProperty("colorValue", value)
                self.swatches.append(swatch)
                grid.addWidget(swatch, r + 1, col)
        layout.addLayout(grid)
        self.preview = label("Aa  기록 내용 미리보기")
        self.preview.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.preview.setMinimumHeight(50)
        layout.addWidget(self.preview)
        self.hex_edit = QLineEdit(current.name() if current.isValid() and current.alpha() else "")
        self.hex_edit.setPlaceholderText("직접 입력 · #RRGGBB")
        self.hex_edit.textChanged.connect(self.update_preview)
        layout.addLayout(row(self.hex_edit, button("적용", self.apply_hex, True)))
        self.error = label("", "Subtitle")
        layout.addWidget(self.error)
        default = "기본 글자색" if kind == "color" else "강조색 없음"
        layout.addLayout(row(button(default, lambda: self.choose("#26354a" if kind == "color" else "transparent")), None, button("취소", self.reject)))
        self.update_preview(self.hex_edit.text())

    def update_preview(self, text):
        color = QColor(text) if re.fullmatch(r"#[0-9a-fA-F]{6}", text) else QColor("#26354a" if self.kind == "color" else "#ffffff")
        foreground = color.name() if self.kind == "color" else "#26354a"
        background = "#ffffff" if self.kind == "color" else color.name()
        self.preview.setStyleSheet(f"color: {foreground}; background: {background}; border: 1px solid #e2e8f0; border-radius: 6px; font-size: 16px;")

    def apply_hex(self):
        text = self.hex_edit.text().strip()
        if not re.fullmatch(r"#[0-9a-fA-F]{6}", text):
            self.error.setText("#과 여섯 자리 색상 코드를 입력하세요. 예: #12816F")
            return
        self.choose(text)

    def choose(self, value):
        self.selected_color = QColor(value)
        self.accept()


class RichEditor(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        tools = QHBoxLayout()
        self.editor = PortableTextEdit()
        self.editor.setPlaceholderText("무슨 일이 있었나요?\n날짜와 함께 남겨두고 싶은 내용을 자유롭게 작성하세요.")
        self.editor.setMinimumHeight(180)
        self.format_buttons = []
        for text, method in (("B", "bold"), ("I", "italic"), ("U", "underline")):
            b = QToolButton()
            b.setText(text)
            b.setCheckable(True)
            b.setFocusPolicy(Qt.FocusPolicy.NoFocus)
            b.setToolTip({"B": "굵게 (Ctrl+B)", "I": "기울임 (Ctrl+I)", "U": "밑줄 (Ctrl+U)"}[text])
            b.clicked.connect(lambda checked, m=method: self.apply_format(m, checked))
            tools.addWidget(b)
            self.format_buttons.append(b)
        self.size = QComboBox()
        self.size.addItems(["10", "11", "12", "14", "16", "18", "24"])
        self.size.setCurrentText("11")
        self.size.setFixedWidth(75)
        self.size.activated.connect(lambda: self.apply_format("size", float(self.size.currentText())))
        tools.addWidget(self.size)
        for text, kind in (("글자색", "color"), ("강조색", "highlight")):
            b = QToolButton()
            b.setText(text)
            b.setFocusPolicy(Qt.FocusPolicy.NoFocus)
            b.clicked.connect(lambda checked=False, k=kind: self.choose_color(k))
            tools.addWidget(b)
        tools.addStretch()
        tools.addWidget(label("선택한 글자에 서식 적용", "Subtitle"))
        layout.addLayout(tools)
        layout.addWidget(self.editor)
        self.editor.currentCharFormatChanged.connect(self.sync)
        from PySide6.QtGui import QShortcut
        for key, b in zip(("Ctrl+B", "Ctrl+I", "Ctrl+U"), self.format_buttons):
            shortcut = QShortcut(QKeySequence(key), self.editor)
            shortcut.activated.connect(b.click)

    def sync(self, fmt):
        for b, value in zip(self.format_buttons, (fmt.fontWeight() >= QFont.Weight.Bold, fmt.fontItalic(), fmt.fontUnderline())):
            b.setChecked(value)

    def apply_format(self, kind, value):
        fmt = QTextCharFormat()
        if kind == "bold": fmt.setFontWeight(QFont.Weight.Bold if value else QFont.Weight.Normal)
        elif kind == "italic": fmt.setFontItalic(value)
        elif kind == "underline": fmt.setFontUnderline(value)
        elif kind == "size": fmt.setFontPointSize(value)
        elif kind == "color": fmt.setForeground(value)
        elif kind == "highlight": fmt.setBackground(value)
        self.editor.mergeCurrentCharFormat(fmt)
        self.editor.setFocus()

    def choose_color(self, kind):
        cursor = self.editor.textCursor()
        current = cursor.charFormat().foreground().color() if kind == "color" else cursor.charFormat().background().color()
        dialog = ColorPaletteDialog(kind, current, self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.editor.setTextCursor(cursor)
            self.apply_format(kind, dialog.selected_color)


class AssetTree(QTreeWidget):
    dropPreviewChanged = Signal(str)
    duplicateRequested = Signal()
    moveRequested = Signal(list, object)
    importRequested = Signal(object, object)
    deleteRequested = Signal()
    renameRequested = Signal()

    def __init__(self):
        super().__init__()
        self.setObjectName("AssetTree")
        self.setHeaderHidden(True)
        self.setSelectionMode(QTreeWidget.SelectionMode.ExtendedSelection)
        self.setDragEnabled(True)
        self.setAcceptDrops(True)
        self.setDropIndicatorShown(False)
        self.drop_target = None
        self.drop_active = False
        self.accept_node_drops = True
        self.setIndentation(18)
        self.setUniformRowHeights(True)

    def ids(self):
        return [i.data(0, Qt.ItemDataRole.UserRole) for i in self.selectedItems() if not i.isHidden()]

    def startDrag(self, supported):
        mime = QMimeData()
        mime.setData(NODE_MIME, json.dumps(self.ids()).encode())
        drag = QDrag(self)
        drag.setMimeData(mime)
        drag.exec(Qt.DropAction.CopyAction)

    def valid_drag(self, e):
        return e.mimeData().hasFormat(IMPORT_MIME) or (self.accept_node_drops and e.mimeData().hasFormat(NODE_MIME))

    def resolve_drop_target(self, point, importing):
        item = self.itemAt(point)
        if importing:
            while item and item.data(0, Qt.ItemDataRole.UserRole + 1) != "group": item = item.parent()
        while item and item.parent() and point.x() < self.visualItemRect(item).left(): item = item.parent()
        return item

    def dragEnterEvent(self, e):
        if self.valid_drag(e): e.acceptProposedAction()
        else: e.ignore()

    def dragMoveEvent(self, e):
        if not self.valid_drag(e):
            e.ignore()
            return
        self.drop_target = self.resolve_drop_target(e.position().toPoint(), e.mimeData().hasFormat(IMPORT_MIME))
        self.drop_active = True
        names, item = [], self.drop_target
        while item:
            names.insert(0, item.text(0))
            item = item.parent()
        self.dropPreviewChanged.emit("등록 위치: " + (" / ".join(names) if names else "프로젝트 최상위"))
        self.viewport().update()
        e.acceptProposedAction()

    def clear_drop_preview(self):
        self.drop_target = None
        self.drop_active = False
        self.viewport().update()
        self.dropPreviewChanged.emit("")

    def dragLeaveEvent(self, e):
        self.clear_drop_preview()
        e.accept()

    def paintEvent(self, event):
        super().paintEvent(event)
        if not self.drop_active: return
        item = self.drop_target
        if item:
            rect = self.visualItemRect(item)
            bottom, following = rect.bottom(), self.itemBelow(item)
            while following:
                ancestor = following.parent()
                while ancestor and ancestor != item: ancestor = ancestor.parent()
                if ancestor != item: break
                bottom = self.visualItemRect(following).bottom()
                following = self.itemBelow(following)
            rect = QRect(max(1, rect.left()), rect.top(), self.viewport().width() - max(1, rect.left()) - 2, min(bottom, self.viewport().height() - 2) - rect.top() + 1)
        else:
            rect = QRect(1, max(0, self.viewport().height() - 30), self.viewport().width() - 2, 28)
        painter = QPainter(self.viewport())
        painter.setPen(QColor("#12816f"))
        painter.setBrush(QColor(18, 129, 111, 35))
        painter.drawRoundedRect(rect, 4, 4)
        if not item: painter.drawText(rect, Qt.AlignmentFlag.AlignCenter, "프로젝트 최상위에 등록")
        painter.end()

    def dropEvent(self, e):
        if not self.valid_drag(e):
            e.ignore()
            return
        importing = e.mimeData().hasFormat(IMPORT_MIME)
        target = self.resolve_drop_target(e.position().toPoint(), importing)
        ident = target.data(0, Qt.ItemDataRole.UserRole) if target else None
        fmt = IMPORT_MIME if importing else NODE_MIME
        try: data = json.loads(bytes(e.mimeData().data(fmt)))
        except (ValueError, UnicodeDecodeError):
            self.clear_drop_preview()
            e.ignore()
            return
        self.clear_drop_preview()
        if importing: self.importRequested.emit(data, ident)
        else: self.moveRequested.emit(data, ident)
        e.acceptProposedAction()

    def keyPressEvent(self, e):
        if e.key() == Qt.Key.Key_D and e.modifiers() & Qt.KeyboardModifier.ControlModifier: self.duplicateRequested.emit()
        elif e.key() == Qt.Key.Key_Delete: self.deleteRequested.emit()
        elif e.key() == Qt.Key.Key_F2: self.renameRequested.emit()
        else: super().keyPressEvent(e)


class ActionTree(QTreeWidget):
    deleteRequested = Signal()
    renameRequested = Signal()
    def __init__(self):
        super().__init__()
        self.setSelectionMode(QTreeWidget.SelectionMode.ExtendedSelection)
    def keyPressEvent(self, e):
        if e.key() == Qt.Key.Key_Delete: self.deleteRequested.emit()
        elif e.key() == Qt.Key.Key_F2: self.renameRequested.emit()
        else: super().keyPressEvent(e)


class ActionList(QListWidget):
    deleteRequested = Signal()
    def __init__(self):
        super().__init__()
        self.setSelectionMode(QListWidget.SelectionMode.ExtendedSelection)
    def keyPressEvent(self, e):
        if e.key() == Qt.Key.Key_Delete: self.deleteRequested.emit()
        else: super().keyPressEvent(e)


class EventTargets(ActionList):
    nodesDropped = Signal(list)
    def __init__(self):
        super().__init__()
        self.setAcceptDrops(True)
    def dragEnterEvent(self, e):
        if e.mimeData().hasFormat(NODE_MIME): e.acceptProposedAction()
    def dragMoveEvent(self, e): self.dragEnterEvent(e)
    def dropEvent(self, e):
        if e.mimeData().hasFormat(NODE_MIME):
            self.nodesDropped.emit(json.loads(bytes(e.mimeData().data(NODE_MIME))))
            e.acceptProposedAction()


def record_period(entry):
    end = entry.get("end_day") or entry["day"]
    return entry["day"] if end == entry["day"] else f"{entry['day']} ~ {end}"


class NoteCard(QWidget):
    def __init__(self, entry, source):
        super().__init__()
        layout = QVBoxLayout(self)
        layout.setContentsMargins(16, 12, 16, 12)
        layout.setSpacing(5)
        category = label(entry["category"])
        color = {"이슈": "#bb5264", "출고": "#b77d22", "입고": "#198271"}.get(entry["category"], "#497995")
        category.setStyleSheet(f"color: {color}; font-weight: 700; font-size: 18px;")
        layout.addLayout(row(category, None, label(record_period(entry), "Subtitle")))
        plain = " ".join((entry["plain"] or entry["title"]).split())
        preview = label((plain[:100] + "…") if len(plain) > 100 else plain or "메모 없음", "Subtitle", elide=True)
        layout.addWidget(preview)
        footer = label(source + ("   " + entry["tags"] if entry["tags"] else ""), "Subtitle", elide=True)
        layout.addWidget(footer)
        self.setToolTip("더블클릭하여 전체 내용 확인 / 수정")
        # Selection/keyboard belong to the list, not to the display-only card.
        self.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)


# ============================================================================
# DIALOGS
# ============================================================================

class PurposeDelegate(QStyledItemDelegate):
    def __init__(self, choices, parent=None):
        super().__init__(parent)
        self.choices = choices

    def createEditor(self, parent, option, index):
        editor = QComboBox(parent)
        editor.addItem("미지정", "")
        for value in self.choices:
            editor.addItem(value, value)
        editor.activated.connect(lambda *_: self.commitData.emit(editor))
        QTimer.singleShot(0, editor.showPopup)
        return editor

    def setEditorData(self, editor, index):
        value = index.data(Qt.ItemDataRole.EditRole) or ""
        found = editor.findData(value)
        if found < 0:
            editor.addItem(value, value)
            found = editor.count() - 1
        editor.setCurrentIndex(found)

    def setModelData(self, editor, model, index):
        model.setData(index, editor.currentData() or "", Qt.ItemDataRole.EditRole)

    def paint(self, painter, option, index):
        combo = QStyleOptionComboBox()
        combo.rect = option.rect.adjusted(2, 2, -2, -2)
        combo.state = option.state
        combo.currentText = index.data(Qt.ItemDataRole.DisplayRole) or "미지정"
        combo.palette = option.palette
        style = option.widget.style() if option.widget else QApplication.style()
        style.drawComplexControl(QStyle.ComplexControl.CC_ComboBox, combo, painter, option.widget)
        style.drawControl(QStyle.ControlElement.CE_ComboBoxLabel, combo, painter, option.widget)


class SerialRegistrationDialog(QDialog):
    def __init__(self, store, device, parent=None):
        super().__init__(parent)
        self.store, self.device = store, device
        self.purpose_by_serial = {}
        self.setWindowTitle("일련번호별 용도 등록")
        self.resize(720, 720)
        layout = QVBoxLayout(self)
        layout.addWidget(label(device["name"] + " · 일련번호 등록", "Section"))
        self.first, self.last = QLineEdit(), QLineEdit()
        self.first.setObjectName("SerialFirst")
        self.last.setObjectName("SerialLast")
        self.first.setPlaceholderText("예: A-001")
        self.last.setPlaceholderText("예: A-020 · 하나만 등록하려면 비워두세요")
        form = QFormLayout()
        form.addRow("처음 *", self.first)
        form.addRow("끝", self.last)
        layout.addLayout(form)
        self.purpose_choices = list(dict.fromkeys([*PurposeButtons.DEFAULTS, *store.vocab("purpose")]))
        self.purpose = QComboBox()
        self.purpose.addItem("미지정", "")
        for value in self.purpose_choices:
            self.purpose.addItem(value, value)
        self.purpose.setObjectName("UnitPurpose")
        layout.addWidget(label("공통 용도 · 필요할 때 아래 전체 적용 버튼을 누르세요.", "Subtitle"))
        layout.addLayout(row(self.purpose, button("＋ 용도 추가", self.add_purpose)))
        layout.addWidget(button("목록 전체에 용도 적용", self.apply_all))
        self.preview = label("범위를 입력하면 아래에 일련번호가 표시됩니다.", "Subtitle")
        layout.addWidget(self.preview)
        self.table = QTableWidget(0, 2)
        self.table.setHorizontalHeaderLabels(["일련번호", "용도 · 클릭하여 선택"])
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.EditKeyPressed)
        self.table.cellClicked.connect(lambda r, c: self.table.editItem(self.table.item(r, c)) if c == 1 else None)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.delegate = PurposeDelegate(self.purpose_choices, self.table)
        self.table.setItemDelegateForColumn(1, self.delegate)
        self.table.itemChanged.connect(self.remember_purpose)
        layout.addWidget(self.table, 1)
        self.description = QTextEdit()
        self.description.setObjectName("UnitDescription")
        self.description.setAcceptRichText(False)
        self.description.setPlainText(device["description"])
        self.description.setMaximumHeight(90)
        layout.addWidget(label("공통 설명", "Subtitle"))
        layout.addWidget(self.description)
        layout.addLayout(row(None, button("취소", self.reject), button("등록", self.save, True)))
        self.timer = QTimer(self)
        self.timer.setSingleShot(True)
        self.timer.setInterval(200)
        self.timer.timeout.connect(self.rebuild)
        self.first.textChanged.connect(lambda: self.timer.start())
        self.last.textChanged.connect(lambda: self.timer.start())

    def add_purpose(self):
        value, accepted = QInputDialog.getText(self, "용도 추가", "새 용도 이름")
        value = value.strip()
        if not accepted or not value:
            return
        if value not in self.purpose_choices:
            self.purpose_choices.append(value)
            self.purpose.addItem(value, value)
        self.purpose.setCurrentIndex(self.purpose.findData(value))
        self.table.viewport().update()

    def remember_purpose(self, item):
        if item.column() == 1:
            self.purpose_by_serial[self.table.item(item.row(), 0).text()] = item.text().strip()

    def rebuild(self):
        self.timer.stop()
        try:
            serials = serial_range(self.first.text(), self.last.text())
            self.preview.setText(f"{len(serials):,}개 · 각 행의 용도를 지정하세요. 빈 용도는 미지정으로 등록됩니다.")
        except ValueError as exc:
            serials = []
            self.preview.setText(str(exc))
        self.table.blockSignals(True)
        try:
            self.table.setRowCount(len(serials))
            for row_index, serial in enumerate(serials):
                item = QTableWidgetItem(serial)
                item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                self.table.setItem(row_index, 0, item)
                self.table.setItem(row_index, 1, QTableWidgetItem(self.purpose_by_serial.get(serial, "")))
        finally:
            self.table.blockSignals(False)

    def apply_all(self):
        self.rebuild()
        value = self.purpose.currentData() or ""
        for row_index in range(self.table.rowCount()):
            self.table.item(row_index, 1).setText(value)

    def save(self):
        self.timer.stop()
        try:
            serials = serial_range(self.first.text(), self.last.text())
            if not serials:
                raise ValueError("처음 일련번호를 입력하세요.")
            with self.store.atomic():
                for serial in serials:
                    self.store.add_node("unit", self.device["name"], self.device["id"],
                        serial=serial, purpose=self.purpose_by_serial.get(serial, ""),
                        purpose_options=self.purpose_choices, description=self.description.toPlainText())
        except Exception as exc:
            message = "이미 등록된 일련번호가 있습니다. 범위를 확인하세요." if isinstance(exc, sqlite3.IntegrityError) else str(exc)
            QMessageBox.warning(self, "등록 실패", message + "\n입력 내용은 유지됩니다.")
            return
        self.accept()


class NodeDialog(QDialog):
    def __init__(self, store, kind, node=None, parent=None):
        super().__init__(parent)
        self.kind = kind
        self.setWindowTitle(KINDS[kind] + (" 정보 수정" if node else " 등록"))
        self.setMinimumWidth(490)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(26, 24, 26, 24)
        layout.addWidget(label(self.windowTitle(), "Title"))
        form = QFormLayout()
        form.setSpacing(14)
        node = node or {}
        self.name = QLineEdit(node.get("stored_name", node.get("name", "")))
        self.serial = QLineEdit(node.get("serial", ""))
        self.purpose = PurposeButtons(store.vocab("purpose"), node.get("purpose", ""))
        self.description = QTextEdit()
        self.description.setAcceptRichText(False)
        self.description.setPlainText(node.get("description", ""))
        self.description.setPlaceholderText("장비의 특징, 관리 메모 등 설명을 입력하세요.")
        self.description.setMaximumHeight(130)
        self.tag_edit = TagEdit(store.vocab("tags"), recent=store.recent_suggestions("tags"))
        self.tag_edit.setText(node.get("tags", ""))
        form.addRow("이름 *", self.name)
        if kind == "unit": form.addRow("일련번호 *", self.serial)
        if kind == "unit": form.addRow("용도", self.purpose)
        else: self.purpose.hide()
        form.addRow("해시태그", self.tag_edit)
        form.addRow("설명", self.description)
        layout.addLayout(form)
        layout.addLayout(row(None, button("취소", self.reject), button("저장", self.validate, True)))

    def validate(self):
        if not self.name.text().strip() or (self.kind == "unit" and not self.serial.text().strip()):
            QMessageBox.warning(self, "입력 확인", "이름과 필요한 일련번호를 입력하세요.")
            return
        self.accept()

    def values(self):
        return dict(name=self.name.text(), serial=self.serial.text(), purpose=self.purpose.currentText(), purpose_options=list(self.purpose.choices) if self.kind == "unit" else (), description=self.description.toPlainText(), tag_text=self.tag_edit.text())


class PurposeChangeDialog(QDialog):
    def __init__(self, store, target_ids, parent=None):
        super().__init__(parent)
        self.setWindowTitle("용도 변경")
        self.setMinimumWidth(560)
        targets = [store.node(ident) for ident in store.descendants(target_ids)
                   if store.node(ident)["kind"] == "unit"]
        current = {node["purpose"] for node in targets}
        selected = next(iter(current)) if len(current) == 1 else ""
        layout = QVBoxLayout(self)
        layout.setContentsMargins(26, 24, 26, 24)
        layout.addWidget(label("용도 일괄 변경", "Title"))
        layout.addWidget(label(f"선택한 항목 아래 개별 장비 {len(targets):,}개의 용도를 함께 변경합니다.", "Subtitle"))
        if len(current) > 1:
            layout.addWidget(label("현재 여러 용도가 섞여 있습니다. 아래에서 새 용도를 선택하세요.", "Subtitle"))
        self.purpose = PurposeButtons(store.vocab("purpose"), selected)
        layout.addWidget(self.purpose)
        self.clear_purpose = QCheckBox("용도 지정 해제")
        self.clear_purpose.toggled.connect(lambda checked: self.purpose.setEnabled(not checked))
        layout.addWidget(self.clear_purpose)
        layout.addLayout(row(None, button("취소", self.reject), button("변경", self.validate, True)))

    def validate(self):
        if not self.clear_purpose.isChecked() and not self.purpose.currentText():
            QMessageBox.warning(self, "입력 확인", "변경할 용도를 선택하거나 용도 지정 해제를 선택하세요.")
            return
        self.accept()

    def value(self):
        return "" if self.clear_purpose.isChecked() else self.purpose.currentText()


class EntryDialog(QDialog):
    def __init__(self, store, event=False, entry=None, category="기타", parent=None, selected_ids=(), day=None):
        super().__init__(parent)
        self.save_handler = None
        entry = entry or {}
        self.setWindowTitle("기록" + (" 수정" if entry else " 작성"))
        self.resize(1440 if event else 760, 820 if event else 680)
        if event:
            self.setMinimumSize(1250, 720)
        outer = QHBoxLayout(self)
        layout = QVBoxLayout()
        outer.addLayout(layout, 2)
        layout.setContentsMargins(26, 24, 26, 24)
        layout.setSpacing(14)
        layout.addWidget(label(self.windowTitle(), "Title"))
        self.original_title = entry.get("title", "")
        self.category = CategoryButtons(store.categories(), entry.get("category", category))
        self.day = QDateEdit()
        self.day.setCalendarPopup(True)
        self.day.setDisplayFormat("yyyy-MM-dd")
        self.day.setDate(QDate.fromString(entry["day"], "yyyy-MM-dd") if entry else day or QDate.currentDate())
        layout.addWidget(self.category)
        self.purpose_change_panel = QFrame()
        self.purpose_change_panel.setObjectName("Card")
        purpose_layout = QVBoxLayout(self.purpose_change_panel)
        purpose_layout.setContentsMargins(16, 14, 16, 14)
        purpose_layout.setSpacing(8)
        purpose_layout.addWidget(label("변경할 용도", "Section"))
        purpose_layout.addWidget(label("저장하면 선택한 모든 개별 장비의 용도가 함께 변경됩니다.", "Subtitle"))
        self.purpose_after = PurposeButtons(store.vocab("purpose"), entry.get("purpose_after") or "")
        purpose_layout.addWidget(self.purpose_after)
        self.clear_purpose = QCheckBox("용도 지정 해제")
        self.clear_purpose.setChecked(
            entry.get("category") == "용도변경" and entry.get("purpose_after") == "")
        self.clear_purpose.toggled.connect(lambda checked: self.purpose_after.setEnabled(not checked))
        self.purpose_after.changed.connect(lambda value: self.clear_purpose.setChecked(False))
        self.purpose_after.setEnabled(not self.clear_purpose.isChecked())
        purpose_layout.addWidget(self.clear_purpose)
        layout.addWidget(self.purpose_change_panel)
        self.category.changed.connect(self.update_purpose_change)
        self.update_purpose_change()
        self.mini_calendar = QCalendarWidget(self)
        self.mini_calendar.setWindowFlags(Qt.WindowType.Popup)
        self.mini_calendar.setVerticalHeaderFormat(QCalendarWidget.VerticalHeaderFormat.NoVerticalHeader)
        self.mini_calendar.setSelectedDate(self.day.date())
        self.mini_calendar.resize(330, 250)
        self.mini_calendar.hide()
        def choose_day(value):
            self.day.setDate(value)
            self.mini_calendar.hide()
        self.mini_calendar.clicked.connect(choose_day)
        self.day.dateChanged.connect(self.mini_calendar.setSelectedDate)
        def open_calendar():
            self.mini_calendar.move(self.day.mapToGlobal(QPoint(0, self.day.height())))
            self.mini_calendar.show()
        self.calendar_button = button("", open_calendar)
        self.calendar_button.setIcon(icon("events", "#12816f"))
        self.calendar_button.setIconSize(QSize(20, 20))
        self.calendar_button.setFixedSize(38, 38)
        self.calendar_button.setToolTip("달력에서 날짜 선택")
        self.calendar_button.setAccessibleName("달력에서 날짜 선택")
        self.end_day = QDateEdit()
        self.end_day.setCalendarPopup(True)
        self.end_day.setDisplayFormat("yyyy-MM-dd")
        self.end_day.setDate(QDate.fromString(entry.get("end_day") or entry.get("day", ""), "yyyy-MM-dd") if entry else self.day.date())
        def advance_end(value):
            if self.end_day.date() < value: self.end_day.setDate(value)
        self.day.dateChanged.connect(advance_end)
        self.end_mini_calendar = QCalendarWidget(self)
        self.end_mini_calendar.setWindowFlags(Qt.WindowType.Popup)
        self.end_mini_calendar.setVerticalHeaderFormat(QCalendarWidget.VerticalHeaderFormat.NoVerticalHeader)
        self.end_mini_calendar.setSelectedDate(self.end_day.date())
        self.end_mini_calendar.resize(330, 250)
        self.end_mini_calendar.hide()
        self.end_day.dateChanged.connect(self.end_mini_calendar.setSelectedDate)
        def choose_end(value):
            self.end_day.setDate(value)
            self.end_mini_calendar.hide()
        self.end_mini_calendar.clicked.connect(choose_end)
        def open_end_calendar():
            self.end_mini_calendar.move(self.end_day.mapToGlobal(QPoint(0, self.end_day.height())))
            self.end_mini_calendar.show()
        self.end_calendar_button = button("", open_end_calendar)
        self.end_calendar_button.setIcon(icon("events", "#12816f"))
        self.end_calendar_button.setFixedSize(38, 38)
        self.end_calendar_button.setToolTip("종료일 선택")
        self.end_calendar_button.setAccessibleName("종료일 선택")
        layout.addLayout(row(label("시작일"), self.day, self.calendar_button, label("종료일"), self.end_day, self.end_calendar_button))
        self.location = QLineEdit(entry.get("location", ""))
        self.location.setPlaceholderText("현 위치(선택)")
        layout.addWidget(self.location)
        self.location_suggestions = SuggestionRow(store.recent_suggestions("location"), self.location.setText, "최근 위치")
        layout.addWidget(self.location_suggestions)
        self.rich = RichEditor()
        if entry: self.rich.editor.setHtml(entry["html"])
        layout.addWidget(self.rich, 1)
        self.tag_edit = TagEdit(store.vocab("tags"), recent=store.recent_suggestions("tags"))
        self.tag_edit.setText(entry.get("tags", ""))
        layout.addWidget(self.tag_edit)
        if event:
            initial = store.targets(entry["id"]) if entry else selected_ids
            self.picker = TargetPicker(store, initial, bool(entry))
            self.target_items = self.picker.target_items
            self.target_search = self.picker.target_search
            outer.addWidget(self.picker, 1)
        layout.addLayout(row(None, button("취소", self.reject), button("기록 저장", self.validate, True)))

    def target_ids(self): return self.picker.target_ids()
    def clear_targets(self): self.picker.clear_targets()

    def update_purpose_change(self, *_):
        self.purpose_change_panel.setVisible(self.category.currentText() == "용도변경")

    def validate(self):
        if self.end_day.date() < self.day.date():
            QMessageBox.warning(self, "입력 확인", "종료일은 시작일보다 빠를 수 없습니다.")
            return
        if not self.rich.editor.toPlainText().strip():
            QMessageBox.warning(self, "입력 확인", "기록 내용을 입력하세요.")
            return
        if self.category.currentText() == "용도변경":
            if hasattr(self, "picker") and not self.target_ids():
                QMessageBox.warning(self, "입력 확인", "용도를 변경할 그룹이나 장비를 선택하세요.")
                return
            if not self.clear_purpose.isChecked() and not self.purpose_after.currentText():
                QMessageBox.warning(self, "입력 확인", "변경할 용도를 선택하거나 용도 지정 해제를 선택하세요.")
                return
        if self.save_handler is not None:
            try:
                self.save_handler()
            except Exception as exc:
                message = "같은 이름 또는 일련번호가 이미 있습니다." if isinstance(exc, sqlite3.IntegrityError) else str(exc)
                QMessageBox.warning(self, "저장 실패", message + "\n\n작성 내용은 유지됩니다. 확인 후 다시 저장하세요.")
                return
        self.accept()

    def values(self):
        purpose_after = None
        if self.category.currentText() == "용도변경":
            purpose_after = "" if self.clear_purpose.isChecked() else self.purpose_after.currentText()
        return dict(title=" ".join(self.rich.editor.toPlainText().split())[:60] or self.original_title, category=self.category.currentText(), category_options=list(self.category.choices), day=self.day.date().toString("yyyy-MM-dd"), end_day=self.end_day.date().toString("yyyy-MM-dd"), html=self.rich.editor.toHtml(), plain=self.rich.editor.toPlainText(), location=self.location.text(), tag_text=self.tag_edit.text(), purpose_after=purpose_after)


class HistorySourceDialog(QDialog):
    def __init__(self, store, parent=None):
        super().__init__(parent)
        self.store = store
        self.selected_record = None
        self.offset = 0
        self.setWindowTitle("이력에서 장비 불러오기")
        self.resize(700, 560)
        layout = QVBoxLayout(self)
        layout.addWidget(label("이력에서 장비 불러오기", "Title"))
        hint = label("기록을 선택하면 연결된 장비를 현재 선택에 추가합니다. 기존 기록은 변경되지 않습니다.", "Subtitle")
        hint.setWordWrap(True)
        layout.addWidget(hint)
        self.search = QLineEdit()
        self.search.setPlaceholderText("기록 내용 · 종류 · 위치 · 해시태그 검색")
        layout.addWidget(self.search)
        self.listing = QListWidget()
        layout.addWidget(self.listing, 1)
        self.preview = label("기록을 선택하세요.", "Subtitle")
        self.preview.setWordWrap(True)
        layout.addWidget(self.preview)
        self.more = button("이력 더 보기", lambda: self.load_page(False))
        layout.addWidget(self.more)
        self.apply_button = button("이 기록의 장비 추가", self.choose, True)
        self.apply_button.setEnabled(False)
        layout.addLayout(row(None, button("취소", self.reject), self.apply_button))
        self.timer = QTimer(self)
        self.timer.setSingleShot(True)
        self.timer.setInterval(250)
        self.timer.timeout.connect(self.load_page)
        self.search.textChanged.connect(lambda: self.timer.start())
        self.search.returnPressed.connect(self.load_page)
        self.listing.currentItemChanged.connect(self.preview_record)
        self.listing.itemDoubleClicked.connect(lambda *_: self.choose())
        self.load_page()

    def load_page(self, reset=True):
        self.timer.stop()
        if reset:
            self.offset = 0
            self.listing.clear()
        records = self.store.record_sources(self.search.text(), offset=self.offset)
        for record in records[:50]:
            item = QListWidgetItem(f"{record_period(record)} · {record['category']} · 장비 {record['target_count']}개\n{record['title']}")
            item.setData(Qt.ItemDataRole.UserRole, record)
            self.listing.addItem(item)
        self.offset += min(50, len(records))
        self.more.setVisible(len(records) > 50)
        if not self.listing.count():
            self.preview.setText("불러올 기록이 없습니다.")

    def preview_record(self, item=None, previous=None):
        self.apply_button.setEnabled(item is not None)
        if item is None:
            self.preview.setText("기록을 선택하세요.")
            return
        record = item.data(Qt.ItemDataRole.UserRole)
        self.preview.setText(self.store.target_summary(record, limit=3))

    def choose(self):
        item = self.listing.currentItem()
        if item is not None:
            self.selected_record = item.data(Qt.ItemDataRole.UserRole)
            self.accept()


class TargetPicker(QWidget):
    def __init__(self, store, initial=(), editing=False, collections=True):
        super().__init__()
        panel = self
        panel.setMinimumWidth(720)
        layout = QVBoxLayout(panel)
        layout.setContentsMargins(0, 0, 0, 0)
        split = QSplitter()
        tree_panel = QWidget()
        tree_layout = QVBoxLayout(tree_panel)
        tree_layout.setContentsMargins(0, 0, 6, 0)
        controls = QWidget()
        controls.setMinimumWidth(350)
        control_layout = QVBoxLayout(controls)
        control_layout.setContentsMargins(6, 0, 0, 0)
        control_scroll = QScrollArea()
        control_scroll.setWidgetResizable(True)
        control_scroll.setFrameShape(QFrame.Shape.NoFrame)
        control_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        control_scroll.setWidget(controls)
        split.addWidget(tree_panel)
        split.addWidget(control_scroll)
        split.setSizes([620, 350])
        split.setStretchFactor(0, 1)
        layout.addWidget(split, 1)
        tree_layout.addWidget(label("대상 그룹 · 장치", "Section"))
        hint = label("그룹·장치를 체크하면 현재 하위 장비를 함께 선택합니다. 개별 장비는 따로 해제할 수 있습니다.", "Subtitle")
        hint.setWordWrap(True)
        tree_layout.addWidget(hint)
        if collections:
            self.collection_choice = QComboBox()
            self.collection_choice.addItem("장비 모음 선택", None)
            for collection in store.collections():
                self.collection_choice.addItem(collection["name"], collection["id"])
            self.collection_apply = button("모음 장비 선택", lambda: self.load_collection(store))
            self.collection_apply.setToolTip("모음에 저장한 장비를 아래 기록 대상에 한 번에 체크합니다.")
            self.collection_apply.setEnabled(False)
            self.collection_choice.currentIndexChanged.connect(lambda: self.collection_apply.setEnabled(self.collection_choice.currentData() is not None))
            self.collection_choice.setMaxVisibleItems(12)
            control_layout.addWidget(label("장비 모음", "Section"))
            control_layout.addLayout(row(self.collection_choice, self.collection_apply))
            self.collection_feedback = label("모음을 고르면 저장된 장비들을 한 번에 선택할 수 있습니다.", "Subtitle")
            self.collection_feedback.setWordWrap(True)
            control_layout.addWidget(self.collection_feedback)
            control_layout.addWidget(button("이력에서 장비 불러오기", lambda: self.load_history(store)))
            self.history_feedback = label("이전 기록의 대상 장비를 현재 선택에 추가합니다.", "Subtitle")
            self.history_feedback.setWordWrap(True)
            control_layout.addWidget(self.history_feedback)
        self.target_search = QLineEdit()
        self.target_search.setPlaceholderText("이름 / 일련번호 검색")
        self.target_search.setClearButtonEnabled(True)
        control_layout.addWidget(label("대상 필터", "Section"))
        control_layout.addWidget(self.target_search)
        self.target_tree = QTreeWidget()
        self.target_tree.setHeaderHidden(True)
        self.target_tree.setUniformRowHeights(True)
        self.target_tree.setMinimumHeight(500)
        self.target_tree.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)
        self.target_items = {}
        nodes = store.nodes()
        nodes_by_id = {node["id"]: node for node in nodes}
        self.filter_nodes = nodes
        self.filters = EquipmentFilters(nodes, detailed=True)
        self.filters.changed.connect(lambda: self.filter_targets(self.target_search.text()))
        control_layout.addWidget(self.filters)
        self.select_filtered = button("필터 결과 선택", self.select_visible_targets)
        control_layout.addWidget(self.select_filtered)
        parents = {n["parent"] for n in nodes}
        self.target_leaves = {n["id"] for n in nodes if n["kind"] == "unit" or (n["kind"] == "device" and n["id"] not in parents)}
        for n in nodes:
            item = QTreeWidgetItem([n["serial"] if n["kind"] == "unit" else n["name"]])
            item.setToolTip(0, store.path_label(n["id"], nodes_by_id))
            item.setFlags(item.flags() | Qt.ItemFlag.ItemIsUserCheckable)
            item.setCheckState(0, Qt.CheckState.Unchecked)
            self.target_items[n["id"]] = item
        for n in nodes:
            item = self.target_items[n["id"]]
            parent = self.target_items.get(n["parent"])
            if parent is not None: parent.addChild(item)
            else: self.target_tree.addTopLevelItem(item)
            if n["id"] not in parents and n["kind"] == "group":
                item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsUserCheckable)
        selected = set(initial) if editing else store.descendants(initial)
        for ident in selected & self.target_leaves:
            self.target_items[ident].setCheckState(0, Qt.CheckState.Checked)
        self.target_tree.expandAll()
        tree_layout.addWidget(self.target_tree, 1)
        self.target_summary = label("", "Subtitle")
        control_layout.addStretch()
        control_layout.addWidget(self.target_summary)
        control_layout.addWidget(button("선택 모두 해제", self.clear_targets))
        self.target_tree.itemChanged.connect(self.target_changed)
        self.target_search.textChanged.connect(self.filter_targets)
        self.update_parent_checks()
        self.update_target_summary()

    def target_ids(self):
        return [ident for ident, item in self.target_items.items() if ident in self.target_leaves and item.checkState(0) == Qt.CheckState.Checked]

    def clear_targets(self):
        self.target_tree.blockSignals(True)
        for ident in self.target_leaves:
            self.target_items[ident].setCheckState(0, Qt.CheckState.Unchecked)
        self.target_tree.blockSignals(False)
        self.update_parent_checks()
        self.update_target_summary()

    def update_target_summary(self, *args):
        self.target_summary.setText(f"선택 대상 {len(self.target_ids()):,}개 · 저장 시 함께 연결")

    def filter_targets(self, text):
        matches = self.filters.matches(self.filter_nodes, text)
        visible = with_ancestors(self.filter_nodes, matches)
        for ident, item in self.target_items.items():
            item.setHidden(ident not in visible)
            if ident in visible: item.setExpanded(True)
        count = len(matches & self.target_leaves)
        self.select_filtered.setText(f"필터 결과 {count}개 선택")
        self.select_filtered.setEnabled(count > 0)

    def select_visible_targets(self):
        self.target_tree.blockSignals(True)
        for ident in self.target_leaves:
            if not self.target_items[ident].isHidden(): self.target_items[ident].setCheckState(0, Qt.CheckState.Checked)
        self.target_tree.blockSignals(False)
        self.update_parent_checks()
        self.update_target_summary()

    def update_parent_checks(self):
        self.target_tree.blockSignals(True)
        def update(item):
            if not item.childCount(): return item.checkState(0)
            states = [update(item.child(i)) for i in range(item.childCount())]
            state = states[0] if all(s == states[0] for s in states) else Qt.CheckState.PartiallyChecked
            item.setCheckState(0, state)
            return state
        for i in range(self.target_tree.topLevelItemCount()): update(self.target_tree.topLevelItem(i))
        self.target_tree.blockSignals(False)

    def target_changed(self, item, column):
        if item.childCount():
            state = item.checkState(0)
            self.target_tree.blockSignals(True)
            pending = [item.child(i) for i in range(item.childCount())]
            for child in pending:
                if child.isHidden(): continue
                if child.childCount(): pending.extend(child.child(i) for i in range(child.childCount()))
                else: child.setCheckState(0, state)
            self.target_tree.blockSignals(False)
        self.update_parent_checks()
        self.update_target_summary()

    def load_collection(self, store):
        ident = self.collection_choice.currentData()
        if not ident: return
        self.filters.reset()
        self.target_search.clear()
        for node_id in store.collection_targets(ident):
            item = self.target_items.get(node_id)
            if item: item.setCheckState(0, Qt.CheckState.Checked)
        self.filters.reset()
        self.target_search.clear()
        count = len(store.collection_targets(ident))
        self.collection_feedback.setText(f"{self.collection_choice.currentText()} · 장비 {count}개 선택에 반영" if count else "이 모음에는 현재 선택할 수 있는 장비가 없습니다.")

    def apply_history_targets(self, store, ident):
        # Reuse exact surviving participants, never expand historical groups.
        targets = set(store.targets(ident))
        available = targets & self.target_leaves
        before = set(self.target_ids())
        self.filters.reset()
        self.target_search.clear()
        self.target_tree.blockSignals(True)
        try:
            for node_id in available:
                self.target_items[node_id].setCheckState(0, Qt.CheckState.Checked)
        finally:
            self.target_tree.blockSignals(False)
        self.update_parent_checks()
        self.update_target_summary()
        added = len(available - before)
        self.history_feedback.setText(f"이력에서 {added}개 추가 · 이미 선택됨 {len(available & before)}개. 삭제된 장비와 그룹 등 선택할 수 없는 대상은 제외합니다.")
        return added

    def load_history(self, store):
        dialog = HistorySourceDialog(store, self)
        try:
            if dialog.exec() == QDialog.DialogCode.Accepted:
                self.apply_history_targets(store, dialog.selected_record["id"])
        finally:
            dialog.deleteLater()


class CollectionDialog(QDialog):
    def __init__(self, store, collection=None, selected_ids=(), parent=None):
        super().__init__(parent)
        self.setWindowTitle("장비 모음 수정" if collection else "장비 모음 만들기")
        self.resize(1180, 800)
        self.setMinimumSize(1040, 720)
        layout = QVBoxLayout(self)
        layout.addWidget(label(self.windowTitle(), "Title"))
        hint = label("여러 그룹의 장비를 묶어 기록에서 불러옵니다. 원래 장비 위치는 유지됩니다.", "Subtitle")
        hint.setWordWrap(True)
        layout.addWidget(hint)
        self.name = QLineEdit(collection["name"] if collection else "")
        self.name.setPlaceholderText("모음 이름 · 예: 현장 A 시험 세트")
        layout.addWidget(self.name)
        self.picker = TargetPicker(store, store.collection_targets(collection["id"]) if collection else selected_ids, bool(collection), collections=False)
        layout.addWidget(self.picker, 1)
        layout.addLayout(row(None, button("취소", self.reject), button("모음 저장", self.validate, True)))
    def validate(self):
        if not self.name.text().strip() or not self.picker.target_ids():
            QMessageBox.warning(self, "입력 확인", "모음 이름과 대상 장비를 입력하세요.")
            return
        self.accept()


class ImportTable(QTableWidget):
    def keyPressEvent(self, e):
        if e.key() != Qt.Key.Key_Delete:
            return super().keyPressEvent(e)
        selected = {i.row() for i in self.selectedIndexes()}
        dialog = self.parent_dialog
        kept = [(i,r) for i,r in enumerate(dialog.rows) if i not in selected]
        dialog.imported = {new for new,(old,r) in enumerate(kept) if old in dialog.imported}
        dialog.rows = [r for old,r in kept]
        dialog.render()

    def startDrag(self, actions):
        indices = sorted({i.row() for i in self.selectedIndexes()})
        valid = [i for i in indices if not self.parent_dialog.rows[i].error and i not in self.parent_dialog.imported]
        if not valid: return
        mime = QMimeData()
        mime.setData(IMPORT_MIME, json.dumps({"session": self.parent_dialog.session, "rows": valid}).encode())
        drag = QDrag(self)
        drag.setMimeData(mime)
        drag.exec(Qt.DropAction.CopyAction)


class ImportDialog(QDialog):
    def __init__(self, main):
        super().__init__(main)
        import uuid
        self.main, self.rows, self.imported = main, [], set()
        self.session = uuid.uuid4().hex
        self.setWindowTitle("엑셀로 장비 불러오기")
        self.resize(1180, 720)
        self.setMinimumSize(960, 620)
        self.setModal(False)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 24, 24, 24)
        layout.addWidget(label("엑셀로 장비 불러오기", "Title"))
        layout.addWidget(label("왼쪽 목록에서 행을 선택해 오른쪽 등록 위치 트리로 드래그하세요. 초록색 영역 안에 등록됩니다.", "Subtitle"))
        layout.addLayout(row(button("엑셀 불러오기", self.load), button("템플릿 저장", main.save_template), None))
        self.table = ImportTable(0, 7)
        self.table.parent_dialog = self
        self.table.setHorizontalHeaderLabels(["장치이름", "처음", "끝", "용도", "수량", "검증", "상세"])
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.table.setSelectionMode(QTableWidget.SelectionMode.ExtendedSelection)
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.table.setDragEnabled(True)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        self.table.horizontalHeader().setSectionResizeMode(6, QHeaderView.ResizeMode.Stretch)
        split = QSplitter()
        split.addWidget(self.table)
        panel = QWidget()
        panel_layout = QVBoxLayout(panel)
        panel_layout.setContentsMargins(10, 0, 0, 0)
        panel_layout.addLayout(row(label("등록 위치", "Section"), None,
                                   button("＋ 그룹", lambda: self.add_import_node("group")),
                                   button("＋ 장치", lambda: self.add_import_node("device"))))
        self.import_tree = AssetTree()
        self.import_tree.setDragEnabled(False)
        self.import_tree.accept_node_drops = False
        self.import_tree.importRequested.connect(self.drop_rows)
        self.import_tree.itemSelectionChanged.connect(self.pick_destination)
        panel_layout.addWidget(self.import_tree, 1)
        self.drop_hint = label("그룹 안쪽은 하위, 왼쪽 들여쓰기 영역은 상위 그룹으로 등록합니다.", "Subtitle")
        self.drop_hint.setWordWrap(True)
        self.import_tree.dropPreviewChanged.connect(lambda text: self.drop_hint.setText(text or "그룹을 선택하거나 목록을 드래그하세요."))
        panel_layout.addWidget(self.drop_hint)
        split.addWidget(panel)
        split.setSizes([730, 350])
        layout.addWidget(split, 1)
        self.preview = label("파일을 불러오면 일련번호 생성 결과가 여기에 표시됩니다.", "Subtitle")
        self.preview.setWordWrap(True)
        self.preview.setMinimumHeight(44)
        layout.addWidget(self.preview)
        self.table.itemSelectionChanged.connect(self.show_preview)
        self.destination = QComboBox()
        self.refresh_destinations()
        layout.addLayout(row(label("등록 위치"), self.destination, button("선택 행 등록", self.register_selected, True)))
        layout.addLayout(row(label("Ctrl: 다중 선택   Shift: 연속 선택", "Subtitle"), None, button("닫기", self.close)))

    def refresh_destinations(self):
        old = self.destination.currentData()
        self.destination.clear()
        self.destination.addItem("프로젝트 최상위", None)
        self.import_tree.blockSignals(True)
        self.import_tree.clear_drop_preview()
        self.import_tree.clear()
        self.import_items = {}
        root = QTreeWidgetItem([self.main.store.name])
        root.setData(0, Qt.ItemDataRole.UserRole, None)
        root.setData(0, Qt.ItemDataRole.UserRole + 1, "group")
        self.import_tree.addTopLevelItem(root)
        nodes = self.main.store.nodes()
        for n in nodes:
            if n["kind"] == "unit": continue
            item = QTreeWidgetItem([n["name"]])
            item.setData(0, Qt.ItemDataRole.UserRole, n["id"])
            item.setData(0, Qt.ItemDataRole.UserRole + 1, n["kind"])
            item.setIcon(0, icon(n["kind"]))
            self.import_items[n["id"]] = item
        for n in nodes:
            if n["kind"] == "unit": continue
            self.import_items.get(n["parent"], root).addChild(self.import_items[n["id"]])
        self.import_tree.expandAll()
        self.import_tree.setCurrentItem(self.import_items.get(old, root))
        self.import_tree.blockSignals(False)
        for n in self.main.store.nodes():
            if n["kind"] == "group": self.destination.addItem(self.main.store.path_label(n["id"]), n["id"])
        index = self.destination.findData(old)
        self.destination.setCurrentIndex(max(index, 0))

    def pick_destination(self):
        item = self.import_tree.currentItem()
        while item and item.data(0, Qt.ItemDataRole.UserRole + 1) != "group": item = item.parent()
        ident = item.data(0, Qt.ItemDataRole.UserRole) if item else None
        self.destination.setCurrentIndex(max(0, self.destination.findData(ident)))

    def drop_rows(self, data, parent):
        if data.get("session") == self.session: self.register(data.get("rows", []), parent)

    def add_import_node(self, kind):
        dialog = NodeDialog(self.main.store, kind, parent=self)
        if dialog.exec() != QDialog.DialogCode.Accepted: return
        try:
            ident = self.main.store.add_node(kind, parent=self.destination.currentData(), **dialog.values())
            self.main.refresh()
            self.refresh_destinations()
            self.import_tree.setCurrentItem(self.import_items[ident])
        except Exception as exc: self.main.show_error(exc)

    def add_import_group(self):
        self.add_import_node("group")

    def load(self):
        path, _ = QFileDialog.getOpenFileName(self, "장비 목록 불러오기", "", "Excel (*.xlsx)")
        if not path: return
        try:
            loaded = read_import(path)
        except Exception as exc:
            QMessageBox.warning(self, "엑셀 읽기 실패", str(exc))
            return
        self.rows, self.imported = loaded, set()
        self.render()

    def render(self):
        self.table.setRowCount(len(self.rows))
        for i, r in enumerate(self.rows):
            done = i in self.imported
            details = r.error or ("일련번호 없이 장치 등록" if not r.serials else ", ".join(r.serials[:3]) + (" … " + r.serials[-1] if len(r.serials) > 3 else ""))
            for col, text in enumerate((r.name, r.first, r.last, r.purpose, str(len(r.serials)), "등록됨" if done else "확인 필요" if r.error else "등록 가능", details)):
                item = QTableWidgetItem(text)
                item.setToolTip(details)
                if r.error: item.setForeground(QColor("#b44355"))
                elif done: item.setForeground(QColor("#91a09e"))
                self.table.setItem(i, col, item)
        self.show_preview()

    def show_preview(self):
        indices = sorted({i.row() for i in self.table.selectedIndexes()})
        if not indices:
            self.preview.setText(f"총 {len(self.rows):,}행 · 등록할 행을 선택하세요.")
            return
        count = sum(len(self.rows[i].serials) for i in indices if not self.rows[i].error and i not in self.imported)
        r = self.rows[indices[0]]
        serials = ", ".join(r.serials[:15]) + (" … " + r.serials[-1] if len(r.serials) > 15 else "")
        self.preview.setText(f"{len(indices)}행 선택 · 개별 장비 {count:,}개 생성 예정\n{r.error or serials or '일련번호 없이 장치만 등록'}")

    def register_selected(self):
        self.register(sorted({i.row() for i in self.table.selectedIndexes()}), self.destination.currentData())

    def register(self, indices, parent):
        indices = [i for i in indices if 0 <= i < len(self.rows) and i not in self.imported]
        if not indices:
            QMessageBox.information(self, "등록", "아직 등록하지 않은 행을 선택하세요.")
            return
        if any(self.rows[i].error for i in indices):
            QMessageBox.warning(self, "입력 확인", "오류가 있는 행은 제외하고 선택하세요. 원본 엑셀을 수정한 뒤 다시 불러올 수 있습니다.")
            return
        try:
            self.main.store.import_rows(parent, [(self.rows[i].name, self.rows[i].serials, self.rows[i].purpose) for i in indices])
        except Exception as exc:
            self.main.show_error(exc)
            return
        self.imported.update(indices)
        self.main.refresh()
        self.render()
        self.main.statusBar().showMessage(f"엑셀 {len(indices)}행 등록 및 저장 완료", 6000)


# ============================================================================
# APP
# ============================================================================

ROOT = Path(__file__).resolve().parent


class EquipmentCompareDialog(QDialog):
    def __init__(self, main):
        super().__init__(main)
        self.main = main
        self.project_identity = main.store.project_identity
        self.comparison = None
        self.rows = []
        self.source_errors = []
        self.loaded = False
        self.setWindowTitle("엑셀 장비 목록 비교·반영")
        self.resize(1180, 760)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 24, 24, 24)
        layout.addWidget(label("엑셀 변경사항 확인", "Title"))
        hint = label("수정한 장비 목록을 선택하고 변경 전·후 값을 확인하세요. 누락된 행은 삭제하지 않습니다.", "Subtitle")
        hint.setWordWrap(True)
        layout.addWidget(hint)
        layout.addLayout(row(button("현재 장비 목록 내보내기", main.export_equipment),
                             button("수정한 목록 불러오기", self.load), None))
        self.filename = label("아직 파일을 선택하지 않았습니다.", "Subtitle")
        self.filename.setWordWrap(True)
        layout.addWidget(self.filename)
        self.destination = QComboBox()
        self.destination.addItem("프로젝트 최상위", None)
        for node in main.store.nodes():
            if node["kind"] == "group":
                self.destination.addItem(main.store.path_label(node["id"]), node["id"])
        self.destination.currentIndexChanged.connect(self.compare)
        layout.addLayout(row(label("새 장치 등록 그룹"), self.destination))
        self.summary = label("파일을 불러오면 추가·변경된 장비를 표시합니다.", "Section")
        layout.addWidget(self.summary)
        self.table = QTableWidget(0, 5)
        self.table.setHorizontalHeaderLabels(["엑셀 행", "대상 장비", "속성", "현재 값", "엑셀 값"])
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setWordWrap(True)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.ResizeToContents)
        layout.addWidget(self.table, 1)
        self.errors = QTextBrowser()
        self.errors.setMaximumHeight(120)
        self.errors.hide()
        layout.addWidget(self.errors)
        self.day = QDateEdit(QDate.currentDate())
        self.day.setDisplayFormat("yyyy-MM-dd")
        self.day.setCalendarPopup(True)
        self.mini_calendar = QCalendarWidget(self)
        self.mini_calendar.setWindowFlags(Qt.WindowType.Popup)
        self.mini_calendar.setVerticalHeaderFormat(QCalendarWidget.VerticalHeaderFormat.NoVerticalHeader)
        self.mini_calendar.resize(330, 250)
        self.mini_calendar.hide()
        def choose_day(value):
            self.day.setDate(value)
            self.mini_calendar.hide()
        self.mini_calendar.clicked.connect(choose_day)
        self.day.dateChanged.connect(self.mini_calendar.setSelectedDate)
        def open_calendar():
            self.mini_calendar.setSelectedDate(self.day.date())
            self.mini_calendar.move(self.day.mapToGlobal(QPoint(0, self.day.height())))
            self.mini_calendar.show()
        self.calendar_button = button("", open_calendar)
        self.calendar_button.setIcon(icon("events", "#12816f"))
        self.calendar_button.setFixedSize(38, 38)
        self.calendar_button.setToolTip("적용 날짜 선택")
        self.calendar_button.setAccessibleName("적용 날짜 선택")
        layout.addLayout(row(label("적용 날짜"), self.day, self.calendar_button,
                             label("변경 이력을 이 날짜로 남깁니다.", "Subtitle"), None))
        self.apply_button = button("변경사항 반영", self.apply, True)
        self.apply_button.setEnabled(False)
        layout.addLayout(row(None, button("닫기", self.reject), self.apply_button))

    def load(self):
        path, _ = QFileDialog.getOpenFileName(self, "수정한 장비 목록 선택", "", "Excel (*.xlsx)")
        if not path:
            return
        self.loaded = False
        try:
            self.rows, self.source_errors = read_equipment_list(path, self.main.store)
            self.loaded = True
            self.filename.setText(path)
            self.compare()
        except Exception as exc:
            self.comparison = None
            self.rows, self.source_errors = [], []
            self.filename.setText(path)
            self.apply_button.setEnabled(False)
            self.table.setRowCount(0)
            self.errors.hide()
            self.summary.setText("파일을 읽지 못했습니다.")
            self.main.show_error(exc)

    def compare(self):
        if not self.loaded:
            return
        try:
            self.comparison = compare_equipment_list(self.main.store, self.rows, self.source_errors, self.destination.currentData())
            plan = self.comparison
            additions = sum(change.before is None for change in plan.changes)
            self.summary.setText(f"추가 {additions}개 · 수정 {len(plan.changes) - additions}개 · 오류 {len(plan.errors)}개 · 파일에 없는 장비 {len(plan.omitted)}개 (삭제 안 함)")
            items = []
            for change in plan.changes:
                for key, (before, after) in change.differences().items():
                    if change.before is None and not after:
                        continue
                    items.append([str(change.row), change.caption(),
                                  ("추가 · " if change.before is None else "") + LIST_FIELDS[key],
                                  before or "미지정", after or "미지정"])
            self.table.setRowCount(len(items))
            for row_number, values in enumerate(items):
                for column, value in enumerate(values):
                    item = QTableWidgetItem(value)
                    item.setToolTip(value)
                    self.table.setItem(row_number, column, item)
            self.table.resizeRowsToContents()
            self.errors.setPlainText("\n".join(plan.errors))
            self.errors.setVisible(bool(plan.errors))
            self.apply_button.setEnabled(bool(plan.changes) and not plan.errors)
        except Exception as exc:
            self.apply_button.setEnabled(False)
            self.main.show_error(exc)

    def apply(self):
        if not self.comparison:
            return
        try:
            if not self.main.store or self.main.store.project_identity != self.project_identity:
                raise ValueError("프로젝트가 변경되었습니다. 비교 창을 다시 열어주세요.")
            count, backup = apply_equipment_changes(self.main.store, self.comparison, self.day.date().toString("yyyy-MM-dd"))
            self.main.refresh()
            self.main.statusBar().showMessage(f"장비 {count}개 변경 반영 완료 · 자동 백업: {backup}", 10000)
            self.accept()
        except Exception as exc:
            self.apply_button.setEnabled(False)
            self.main.show_error(exc)


class MainWindow(QMainWindow):
    def __init__(self, settings=None):
        super().__init__()
        self.settings = settings if settings is not None else QSettings("CSG", "SEMP")
        self.store = None
        self.project_lock = None
        self.selected_id = None
        self.event_id = None
        self.mode = "inventory"
        self.import_dialog = None
        self.note_limit = 100
        self.record_query = None
        self.undo_moves = []
        self.redo_moves = []
        self.changing_view = False
        self.setWindowTitle("SEMP · System Equipment Management Program")
        self.resize(1440, 900)
        self.setMinimumSize(1120, 720)
        self.build_menu()
        self.build_ui()
        self.refresh()

    def build_menu(self):
        file = self.menuBar().addMenu("프로젝트")
        for text, key, callback in (("새 프로젝트", "Ctrl+N", self.new_project), ("프로젝트 열기", "Ctrl+O", self.open_project), ("백업 저장", "Ctrl+S", self.backup_project)):
            action = QAction(text, self)
            action.setShortcut(QKeySequence(key))
            action.triggered.connect(callback)
            file.addAction(action)
        self.recent_menu = file.addMenu("최근 프로젝트")
        file.addSeparator()
        file.addAction("프로젝트 폴더 열기", self.open_folder)
        file.addAction("프로젝트 닫기", self.close_project)
        tools = self.menuBar().addMenu("도구")
        tools.addAction("엑셀로 장비 불러오기", self.open_import)
        tools.addAction("엑셀 템플릿 저장", self.save_template)
        tools.addAction("엑셀 장비 목록 내보내기", self.export_equipment)
        tools.addAction("엑셀 장비 목록 비교·반영", self.open_equipment_comparison)
        tools.addAction("휴지통", self.show_trash)
        help_menu = self.menuBar().addMenu("도움말")
        help_menu.addAction("사용 안내", self.show_help)
        search = QAction(self)
        search.setShortcut(QKeySequence("Ctrl+F"))
        search.triggered.connect(lambda: self.search.setFocus())
        self.addAction(search)

    def build_ui(self):
        shell = QWidget()
        shell_layout = QVBoxLayout(shell)
        shell_layout.setContentsMargins(0, 0, 0, 0)
        shell_layout.setSpacing(0)
        rail = QFrame()
        rail.setObjectName("Rail")
        nav = QHBoxLayout(rail)
        nav.setContentsMargins(16, 6, 16, 6)
        nav.setSpacing(6)
        self.logo_link = HomeLink("SEMP")
        self.logo_link.setObjectName("Brand")
        self.logo_link.setFixedWidth(60)
        self.logo_link.setSizePolicy(QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Preferred)
        self.logo_link.clicked.connect(self.show_home)
        nav.addWidget(self.logo_link)
        self.nav_buttons = {}
        for key, text in (("inventory", "  장비 관리"), ("search", "  통합 검색")):
            b = button(text, lambda checked=False, k=key: self.set_mode(k))
            b.setIcon(icon(key, "#b4d1da"))
            b.setIconSize(QSize(17,17))
            b.setCheckable(True)
            self.nav_buttons[key] = b
            nav.addWidget(b)
        nav.addWidget(button("장비 모음", self.manage_collections))
        nav.addStretch()
        import_btn = button("  엑셀로 장비 불러오기", self.open_import)
        import_btn.setIcon(icon("import", "#b4d1da"))
        nav.addWidget(import_btn)
        trash_btn = button("  휴지통", self.show_trash)
        trash_btn.setIcon(icon("trash", "#b4d1da"))
        nav.addWidget(trash_btn)
        nav.addWidget(button("?   사용 안내", self.show_help))
        self.saved_label = label("프로젝트를 열어주세요")
        self.saved_label.setStyleSheet("font-size: 10px; padding: 0 6px;")
        nav.addWidget(self.saved_label)
        shell_layout.addWidget(rail)

        workspace = QWidget()
        wl = QVBoxLayout(workspace)
        wl.setContentsMargins(0, 0, 0, 0)
        wl.setSpacing(0)
        top = QWidget()
        top.setObjectName("TopBar")
        top.setStyleSheet("QWidget#TopBar { background: white; }")
        tl = QHBoxLayout(top)
        tl.setContentsMargins(16, 8, 16, 8)
        self.project_label = HomeLink("프로젝트 없음")
        self.project_label.setObjectName("Section")
        self.project_label.clicked.connect(self.show_home)
        self.project_label.setMinimumWidth(120)
        self.search = QLineEdit()
        self.search.setPlaceholderText("장비, 일련번호, #해시태그 검색    Ctrl+F")
        self.search.setMinimumWidth(280)
        self.search.setMaximumWidth(410)
        self.search.setClearButtonEnabled(True)
        self.search.returnPressed.connect(lambda: self.set_mode("search"))
        self.search.textChanged.connect(lambda: self.refresh_search() if self.mode == "search" else None)
        self.new_event_btn = button("＋ 기록 작성", self.new_event, True)
        tl.addWidget(self.project_label)
        tl.addStretch()
        tl.addWidget(self.search)
        tl.addWidget(self.new_event_btn)
        wl.addWidget(top)

        splitter = QSplitter()
        panel = QFrame()
        panel.setObjectName("TreePanel")
        panel.setMinimumWidth(225)
        panel.setMaximumWidth(480)
        pl = QVBoxLayout(panel)
        pl.setContentsMargins(10, 10, 8, 8)
        pl.setSpacing(6)
        self.calendar_view_btn = button("달력 보기", self.clear_selection)
        self.calendar_view_btn.setToolTip("선택을 해제하고 전체 기록 달력으로 이동합니다.")
        pl.addLayout(row(label("장비 탐색", "Section"), None, self.calendar_view_btn))
        self.tree_filter = QLineEdit()
        self.tree_filter.setPlaceholderText("트리 내 이름 / 일련번호 찾기")
        self.tree_filter.setClearButtonEnabled(True)
        self.tree_filter.textChanged.connect(self.filter_tree)
        pl.addWidget(self.tree_filter)
        self.equipment_filters = EquipmentFilters()
        self.equipment_filters.changed.connect(self.filter_tree)
        pl.addWidget(self.equipment_filters)
        self.add_group_btn = button("＋ 그룹", lambda: self.add_node("group"))
        self.add_device_btn = button("＋ 장치", lambda: self.add_node("device"))
        pl.addLayout(row(self.add_group_btn, self.add_device_btn))
        self.tree = AssetTree()
        self.tree.itemSelectionChanged.connect(self.tree_selection)
        self.tree.itemDoubleClicked.connect(lambda item, col: self.edit_node())
        self.tree.moveRequested.connect(self.move_nodes)
        self.tree.importRequested.connect(self.drop_import)
        self.tree.deleteRequested.connect(self.delete_nodes)
        self.tree.renameRequested.connect(self.edit_node)
        self.tree.duplicateRequested.connect(self.duplicate_nodes)
        self.tree.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.tree.customContextMenuRequested.connect(self.tree_menu)
        pl.addWidget(self.tree, 1)
        self.tree_hint = label("그룹을 만들고 장비를 등록하세요.", "Subtitle")
        self.tree_hint.setWordWrap(True)
        pl.addWidget(self.tree_hint)
        pl.addWidget(label("Ctrl 다중 선택 · Shift 연속 선택\n드래그 이동 · F2 수정 · Delete 삭제", "Subtitle"))
        splitter.addWidget(panel)

        self.stack = QStackedWidget()
        self.stack.addWidget(self.build_inventory())
        self.stack.addWidget(self.build_events())
        self.stack.addWidget(self.build_search())
        content = QWidget()
        content_layout = QVBoxLayout(content)
        content_layout.setContentsMargins(0, 0, 0, 0)
        self.view_tabs = QTabBar()
        self.view_tabs.setStyleSheet("QTabBar::tab { padding: 7px 12px; }")
        self.view_tabs.setExpanding(False)
        self.view_tabs.setTabsClosable(True)
        self.view_tabs.setUsesScrollButtons(True)
        self.view_tabs.addTab("홈")
        self.view_tabs.setTabData(0, {"key": ("home",), "mode": "inventory", "node": None, "entry": None, "query": None})
        self.view_tabs.setTabButton(0, QTabBar.ButtonPosition.RightSide, None)
        self.view_tabs.currentChanged.connect(self.activate_view)
        self.view_tabs.tabCloseRequested.connect(self.close_view)
        self.close_all_tabs_btn = button("모든 탭 닫기", self.close_all_views)
        self.close_all_tabs_btn.setToolTip("홈을 제외하고 열려 있는 장치·검색·기록 탭을 모두 닫습니다.")
        self.close_all_tabs_btn.setEnabled(False)
        tabs = QWidget()
        tabs_layout = QHBoxLayout(tabs)
        tabs_layout.setContentsMargins(0, 0, 8, 0)
        tabs_layout.setSpacing(6)
        tabs_layout.addWidget(self.view_tabs, 1)
        tabs_layout.addWidget(self.close_all_tabs_btn)
        content_layout.addWidget(tabs)
        content_layout.addWidget(self.stack, 1)
        splitter.addWidget(content)
        edit_menu = self.menuBar().addMenu("편집")
        self.undo_action = QAction("트리 이동 취소", self.tree)
        self.redo_action = QAction("트리 이동 다시 실행", self.tree)
        for action, shortcut, redo in ((self.undo_action, "Ctrl+Z", False), (self.redo_action, "Ctrl+Y", True)):
            action.setShortcut(QKeySequence(shortcut))
            action.setShortcutContext(Qt.ShortcutContext.WidgetWithChildrenShortcut)
            action.triggered.connect(lambda checked=False, r=redo: self.undo_tree_move(r))
            self.tree.addAction(action)
            edit_menu.addAction(action)
        self.update_move_actions()
        splitter.setSizes([290, 970])
        splitter.setStretchFactor(1, 1)
        splitter.setChildrenCollapsible(False)
        wl.addWidget(splitter, 1)
        self.project_pages = QStackedWidget()
        welcome = QWidget()
        intro = QVBoxLayout(welcome)
        intro.setContentsMargins(80, 50, 80, 50)
        intro.addStretch()
        intro.addWidget(label("SEMP · 장비 이력 관리", "Title"))
        intro.addWidget(label("프로젝트 파일을 찾아 불러오세요.", "Section"))
        intro.addWidget(label("기존 프로젝트 폴더의 project.semp 파일을 선택하면 장비와 기록을 불러옵니다.", "Subtitle"))
        intro.addLayout(row(button("프로젝트 파일 불러오기", self.open_project, True), button("새 프로젝트", self.new_project), None))
        intro.addWidget(label("최근에 열어본 프로젝트", "Section"))
        self.recent_list = QListWidget()
        self.recent_list.setMaximumHeight(210)
        self.recent_list.itemClicked.connect(lambda item: self.load_project(item.data(Qt.ItemDataRole.UserRole)))
        intro.addWidget(self.recent_list)
        self.refresh_recent()
        intro.addStretch()
        self.project_pages.addWidget(welcome)
        self.project_pages.addWidget(workspace)
        self.navigation = rail
        shell_layout.addWidget(self.project_pages, 1)
        self.setCentralWidget(shell)
        self.statusBar().showMessage("새 프로젝트를 만들거나 project.semp 파일을 열어 시작하세요.")

    def build_inventory(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.setContentsMargins(28, 16, 28, 16)
        layout.setSpacing(10)
        self.breadcrumb = label("WORKSPACE / EQUIPMENT", "Subtitle", elide=True)
        layout.addWidget(self.breadcrumb)
        self.detail_title = label("장비의 모든 이력, 한곳에서", "Title", elide=True)
        self.add_unit_btn = button("＋ 일련번호", self.add_serials)
        self.device_badge = label("", "DeviceBadge", elide=True)
        self.device_badge.setMaximumWidth(240)
        self.device_badge.setSizePolicy(QSizePolicy.Policy.Preferred, QSizePolicy.Policy.Fixed)
        self.device_badge.setFixedHeight(38)
        title_row = row(self.device_badge, self.detail_title, self.add_unit_btn)
        title_row.setStretch(1, 1)
        layout.addLayout(title_row)
        self.detail_hint = label("", "Subtitle")
        self.detail_hint.setWordWrap(True)
        self.edit_btn = button("정보 수정", self.edit_node)
        self.edit_btn.setIcon(icon("edit", "#61798d"))
        self.edit_btn.setIconSize(QSize(16, 16))
        self.edit_btn.setToolTip("이름 · 설명 · 해시태그 수정 (개별 장비는 일련번호 · 용도 포함)")
        layout.addLayout(row(self.detail_hint, None, self.edit_btn))
        self.metrics_panel = QWidget()
        metrics = QHBoxLayout(self.metrics_panel)
        metrics.setContentsMargins(0, 0, 0, 0)
        metrics.setSpacing(12)
        self.metrics = []
        for _ in range(3):
            card = QFrame()
            card.setObjectName("Card")
            cl = QVBoxLayout(card)
            cl.setContentsMargins(18, 15, 18, 15)
            name = label("", "Subtitle")
            value = DoubleClickLabel("")
            value.setObjectName("Metric")
            value.setWordWrap(True)
            cl.addWidget(name)
            cl.addWidget(value)
            self.metrics.append((name, value))
            metrics.addWidget(card, 1)
        self.metrics[2][1].doubleClicked.connect(self.change_detail_purpose)
        layout.addWidget(self.metrics_panel)
        self.calendar_panel = QWidget()
        calendar_layout = QVBoxLayout(self.calendar_panel)
        calendar_layout.setContentsMargins(0, 0, 0, 0)
        calendar_layout.addWidget(label("날짜를 선택해 기록을 확인하고 새 기록을 추가하세요.", "Subtitle"))
        self.calendar = EventCalendar()
        self.calendar.selectionChanged.connect(self.refresh_calendar_day)
        self.calendar.currentPageChanged.connect(lambda y, m: self.refresh_calendar())
        calendar_layout.addWidget(self.calendar, 1)
        self.calendar_day_label = label("", "Section")
        calendar_layout.addWidget(self.calendar_day_label)
        self.calendar_entries = QListWidget()
        self.calendar_entries.setMaximumHeight(120)
        self.calendar_entries.itemDoubleClicked.connect(lambda item: self.open_record(item.data(Qt.ItemDataRole.UserRole)))
        calendar_layout.addWidget(self.calendar_entries)
        self.calendar_create_btn = button("＋ 선택 날짜에 기록 작성", lambda: self.new_event(day=self.calendar.selectedDate()), True)
        calendar_layout.addLayout(row(None, self.calendar_create_btn))
        layout.addWidget(self.calendar_panel, 1)
        self.tabs = QTabWidget()
        self.tabs.currentChanged.connect(self.remember_detail_tab)
        history = QWidget()
        hl = QVBoxLayout(history)
        hl.setContentsMargins(0, 14, 0, 0)
        self.inherited = QCheckBox("상위 그룹 이력 포함")
        self.inherited.setChecked(True)
        self.inherited.toggled.connect(self.refresh_history)
        self.category_filter = QComboBox()
        self.category_filter.addItems(["전체", *CATEGORIES])
        self.category_filter.currentTextChanged.connect(self.refresh_history)
        hl.addLayout(row(self.inherited, None, label("종류", "Subtitle"), self.category_filter))
        self.timeline = ActionList()
        self.timeline.setObjectName("Timeline")
        self.timeline.itemDoubleClicked.connect(lambda i: self.open_record(i.data(Qt.ItemDataRole.UserRole)))
        self.timeline.deleteRequested.connect(self.delete_entries)
        self.timeline.setSpacing(1)
        hl.addWidget(self.timeline, 1)
        self.history_empty = label("", "Subtitle")
        self.history_empty.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.history_empty.setWordWrap(True)
        hl.addWidget(self.history_empty)
        self.more_btn = button("이력 더 보기", self.more_history)
        hl.addWidget(self.more_btn)
        self.tabs.addTab(history, "활동 이력")
        basic = QWidget()
        bl = QVBoxLayout(basic)
        bl.setContentsMargins(0, 16, 0, 0)
        self.basic_text = QTextBrowser()
        self.basic_text.setOpenLinks(False)
        bl.addWidget(self.basic_text)
        self.tabs.addTab(basic, "기본 정보")
        child = QWidget()
        chl = QVBoxLayout(child)
        chl.setContentsMargins(0, 16, 0, 0)
        self.children = ActionTree()
        self.children.setHeaderLabels(["이름 / 일련번호", "종류", "용도 / 설명"])
        self.children.setRootIsDecorated(False)
        self.children.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self.children.itemDoubleClicked.connect(lambda i, c: self.select_node(i.data(0, Qt.ItemDataRole.UserRole)))
        self.children.deleteRequested.connect(lambda: self.delete_nodes([i.data(0, Qt.ItemDataRole.UserRole) for i in self.children.selectedItems()]))
        self.children.renameRequested.connect(self.edit_child)
        chl.addWidget(self.children)
        self.tabs.addTab(child, "하위 항목")
        layout.addWidget(self.tabs, 1)
        return page

    def build_events(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.setContentsMargins(28, 26, 28, 20)
        layout.setSpacing(14)
        layout.addWidget(label("WORKSPACE / RECORDS", "Subtitle"))
        layout.addWidget(label("기록", "Title"))
        self.records_hint = label("기록의 내용과 연결된 장비를 확인하세요.", "Subtitle")
        layout.addWidget(self.records_hint)
        split = QSplitter()
        self.events = ActionList()
        self.events.setMinimumWidth(165)
        self.events.setMaximumWidth(380)
        self.events.setWordWrap(True)
        self.events.itemSelectionChanged.connect(self.event_selection)
        self.events.itemDoubleClicked.connect(lambda i: self.edit_entry(i.data(Qt.ItemDataRole.UserRole)))
        self.events.deleteRequested.connect(lambda: self.delete_entries(self.events))
        split.addWidget(self.events)
        detail = QWidget()
        dl = QVBoxLayout(detail)
        dl.setContentsMargins(16, 0, 0, 0)
        self.event_title = label("기록을 만들어보세요", "Section")
        self.event_title.setWordWrap(True)
        self.event_meta = label("", "Subtitle")
        self.event_meta.setWordWrap(True)
        dl.addWidget(self.event_title)
        dl.addWidget(self.event_meta)
        self.event_body = QTextBrowser()
        self.event_body.setOpenLinks(False)
        self.event_body.setMinimumHeight(130)
        dl.addWidget(self.event_body, 1)
        self.event_edit_btn = button("기록 수정", lambda: self.edit_entry(self.event_id))
        dl.addLayout(row(None, self.event_edit_btn))
        self.target_count = label("연결된 장비", "Section")
        dl.addWidget(self.target_count)
        hint = label("왼쪽 트리에서 여기에 드래그\n그룹을 넣으면 현재 하위 장비를 연결합니다.", "Subtitle")
        hint.setWordWrap(True)
        dl.addWidget(hint)
        self.targets = EventTargets()
        self.targets.nodesDropped.connect(self.link_targets)
        self.targets.deleteRequested.connect(self.unlink_targets)
        self.targets.itemDoubleClicked.connect(lambda i: self.select_node(i.data(Qt.ItemDataRole.UserRole)))
        dl.addWidget(self.targets, 1)
        self.link_btn = button("선택 장비 연결", lambda: self.link_targets(self.tree.ids()), True)
        self.unlink_btn = button("선택 연결 해제", self.unlink_targets)
        dl.addLayout(row(self.link_btn, self.unlink_btn))
        split.addWidget(detail)
        split.setSizes([270, 500])
        split.setStretchFactor(1, 1)
        split.setChildrenCollapsible(False)
        layout.addWidget(split, 1)
        return page

    def build_search(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.setContentsMargins(28, 26, 28, 20)
        layout.setSpacing(15)
        layout.addWidget(label("WORKSPACE / SEARCH", "Subtitle"))
        layout.addWidget(label("필요한 기록을 빠르게", "Title"))
        layout.addWidget(label("상단 검색창에서 장비·일련번호·메모를 검색하세요. #태그는 정확히 일치하는 태그를 찾습니다.", "Subtitle"))
        self.search_count = label("검색어를 입력하세요.", "Section")
        layout.addWidget(self.search_count)
        self.search_results = ActionTree()
        self.search_results.setHeaderLabels(["구분", "검색 결과", "위치 / 연결 대상", "날짜"])
        self.search_results.setRootIsDecorated(False)
        self.search_results.setAlternatingRowColors(True)
        self.search_results.header().setSectionResizeMode(1, self.search_results.header().ResizeMode.Stretch)
        self.search_results.itemDoubleClicked.connect(self.activate_search)
        self.search_results.deleteRequested.connect(self.delete_search_results)
        layout.addWidget(self.search_results, 1)
        return page

    def recent_paths(self):
        paths = self.settings.value("recentProjects", [])
        return [paths] if isinstance(paths, str) else list(paths or [])

    def remember_project(self, path):
        path = str(Path(path).resolve())
        paths = [p for p in self.recent_paths() if p.casefold() != path.casefold()]
        self.settings.setValue("recentProjects", [path, *paths][:8])
        self.settings.sync()
        self.refresh_recent()

    def refresh_recent(self):
        self.recent_list.clear()
        self.recent_menu.clear()
        for path in self.recent_paths():
            exists = Path(path).is_file()
            text = f"{Path(path).parent.name}  ·  {path}" + ("  (파일 없음)" if not exists else "")
            item = QListWidgetItem()
            item.setData(Qt.ItemDataRole.AccessibleTextRole, text)
            item.setToolTip(path)
            item.setData(Qt.ItemDataRole.UserRole, path)
            if not exists: item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsEnabled)
            item.setSizeHint(QSize(320, 58))
            self.recent_list.addItem(item)
            self.recent_list.setItemWidget(item, RecentProjectCard(Path(path).parent.name, path, exists))
            action = self.recent_menu.addAction(text, lambda checked=False, p=path: self.load_project(p))
            action.setEnabled(exists)
        if not self.recent_paths():
            item = QListWidgetItem("최근에 열어본 프로젝트가 없습니다.")
            item.setFlags(Qt.ItemFlag.NoItemFlags)
            self.recent_list.addItem(item)
        self.recent_menu.setEnabled(bool(self.recent_paths()))

    def refresh_calendar(self):
        self.calendar.events_by_day = {}
        month = date(self.calendar.yearShown(), self.calendar.monthShown(), 1)
        first = month - timedelta(days=7)
        last = month + timedelta(days=38)
        for entry in self.store.entries(day_from=first.isoformat(), day_to=last.isoformat()) if self.store else []:
            # Cache only the visible month and adjacent calendar cells, even for long periods.
            start = max(date.fromisoformat(entry["day"]), first)
            end = min(date.fromisoformat(entry.get("end_day") or entry["day"]), last)
            while start <= end:
                self.calendar.events_by_day.setdefault(start.isoformat(), []).append(entry)
                start += timedelta(days=1)
        self.calendar.updateCells()
        self.refresh_calendar_day()

    def refresh_calendar_day(self):
        day = self.calendar.selectedDate().toString("yyyy-MM-dd")
        entries = self.calendar.events_by_day.get(day, [])
        self.calendar_day_label.setText(f"{day} · 기록 {len(entries)}건")
        self.calendar_create_btn.setText(f"＋ {day} 기록 작성")
        self.calendar_create_btn.setEnabled(self.store is not None)
        self.calendar_entries.clear()
        for entry in entries:
            item = QListWidgetItem(f"{entry['category']} · {self.store.target_summary(entry)} · {record_period(entry)}\n{entry['title']}")
            item.setToolTip("\n".join(self.store.path_label(i) for i in entry["targets"]))
            item.setData(Qt.ItemDataRole.UserRole, entry["id"])
            self.calendar_entries.addItem(item)
        if not entries:
            item = QListWidgetItem("기록 없음 · 아래 작성 버튼으로 새 기록을 추가하세요.")
            item.setFlags(Qt.ItemFlag.NoItemFlags)
            self.calendar_entries.addItem(item)

    def new_collection(self):
        self.edit_collection()

    def edit_collection(self, collection=None):
        if not self.require_project(): return
        dialog = CollectionDialog(self.store, collection, self.tree.ids(), self)
        if dialog.exec() != QDialog.DialogCode.Accepted: return
        try:
            self.store.save_collection(dialog.name.text(), dialog.picker.target_ids(), collection["id"] if collection else None)
            self.statusBar().showMessage("장비 모음 저장 완료", 5000)
        except Exception as exc: self.show_error(exc)

    def manage_collections(self):
        if not self.require_project(): return
        dialog = QDialog(self)
        dialog.setWindowTitle("장비 모음")
        dialog.resize(600, 450)
        layout = QVBoxLayout(dialog)
        layout.addWidget(label("장비 모음", "Title"))
        layout.addWidget(label("여러 그룹의 장비를 묶어 기록 작성 시 한 번에 불러오세요.", "Subtitle"))
        listing = QListWidget()
        listing.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)
        listing.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)
        layout.addWidget(listing, 1)
        def populate():
            listing.clear()
            for collection in self.store.collections():
                item = QListWidgetItem(f"{collection['name']} · 장비 {len(self.store.collection_targets(collection['id']))}개")
                item.setData(Qt.ItemDataRole.UserRole, collection)
                listing.addItem(item)
        def edit(new=False):
            item = listing.currentItem()
            if not new and item is None: return
            self.edit_collection(None if new else item.data(Qt.ItemDataRole.UserRole))
            populate()
        def remove():
            item = listing.currentItem()
            if item is None: return
            collection = item.data(Qt.ItemDataRole.UserRole)
            if QMessageBox.question(dialog, "모음 삭제", f"'{collection['name']}' 모음을 삭제할까요? 원래 장비는 유지됩니다.") != QMessageBox.StandardButton.Yes: return
            self.store.delete_collection(collection["id"])
            populate()
        listing.itemDoubleClicked.connect(lambda item: edit())
        layout.addLayout(row(button("＋ 모음 만들기", lambda: edit(True)), button("수정", edit), button("삭제", remove), None, button("닫기", dialog.reject)))
        populate()
        dialog.exec()

    def require_project(self):
        if self.store: return True
        QMessageBox.information(self, "프로젝트 열기", "먼저 새 프로젝트를 만들거나 기존 프로젝트를 열어주세요.")
        return False

    def show_error(self, exc):
        message = str(exc)
        if isinstance(exc, sqlite3.IntegrityError):
            message = "같은 위치에 동일한 이름 또는 일련번호가 이미 있습니다. 이름이나 등록 위치를 확인하세요. 변경은 저장되지 않았습니다."
        QMessageBox.warning(self, "작업을 완료하지 못했습니다", message)

    def new_project(self):
        parent = QFileDialog.getExistingDirectory(self, "프로젝트 폴더를 만들 위치 선택")
        if not parent: return
        name, ok = QInputDialog.getText(self, "새 프로젝트", "프로젝트 이름")
        if not ok or not name.strip(): return
        try:
            project = Store.create(parent, name)
            path = project.path
            project.close()
            self.load_project(path)
        except Exception as exc: self.show_error(exc)

    def open_project(self):
        path, _ = QFileDialog.getOpenFileName(self, "SEMP 프로젝트 열기", "", "SEMP 프로젝트 (*.semp)")
        if path: self.load_project(path)

    def load_project(self, path):
        path = Path(path).resolve()
        if self.store and path == self.store.path: return
        lock = QLockFile(str(path.parent / ".semp.lock"))
        lock.setStaleLockTime(0)
        if not lock.tryLock(0):
            QMessageBox.warning(self, "프로젝트 사용 중", "다른 SEMP에서 이 프로젝트를 사용 중이거나 폴더에 쓰기 권한이 없습니다.")
            return
        try:
            incoming = Store(path)
        except Exception as exc:
            lock.unlock()
            self.show_error(exc)
            return
        self.close_project()
        self.store, self.project_lock = incoming, lock
        self.remember_project(path)
        self.selected_id = self.event_id = None
        self.setWindowTitle(f"{self.store.name} — SEMP")
        self.refresh()
        self.set_mode("inventory")

    def close_project(self):
        if self.import_dialog:
            self.import_dialog.close()
            self.import_dialog = None
        if self.store:
            self.store.close()
            self.store = None
        if self.project_lock:
            self.project_lock.unlock()
            self.project_lock = None
        self.selected_id = self.event_id = None
        self.record_query = None
        self.undo_moves.clear()
        self.redo_moves.clear()
        self.update_move_actions()
        self.view_tabs.blockSignals(True)
        while self.view_tabs.count() > 1: self.view_tabs.removeTab(1)
        self.view_tabs.setCurrentIndex(0)
        self.view_tabs.blockSignals(False)
        self.close_all_tabs_btn.setEnabled(False)
        self.setWindowTitle("SEMP · System Equipment Management Program")
        if hasattr(self, "stack"): self.refresh()

    def closeEvent(self, event):
        self.close_project()
        event.accept()

    def backup_project(self):
        if not self.require_project(): return
        try:
            path = self.store.backup()
            self.statusBar().showMessage(f"백업 저장 완료 · {path.name}", 7000)
        except Exception as exc: self.show_error(exc)

    def open_folder(self):
        if self.require_project():
            from PySide6.QtGui import QDesktopServices
            from PySide6.QtCore import QUrl
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(self.store.root)))

    def refresh(self):
        enabled = bool(self.store)
        self.project_pages.setCurrentIndex(1 if enabled else 0)
        self.navigation.setVisible(enabled)
        for b in (self.add_group_btn, self.add_device_btn, self.new_event_btn): b.setEnabled(enabled)
        self.project_label.setText(self.store.name if self.store else "프로젝트 없음")
        self.saved_label.setText("● 변경 내용 자동 저장" if self.store else "프로젝트를 열어주세요")
        current_category = self.category_filter.currentText()
        self.category_filter.blockSignals(True)
        self.category_filter.clear()
        self.category_filter.addItems(["전체", *(self.store.categories() if self.store else CATEGORIES)])
        self.category_filter.setCurrentText(current_category)
        self.category_filter.blockSignals(False)
        self.refresh_tree()
        self.refresh_inventory()
        self.refresh_events()
        self.refresh_search()
        if self.import_dialog: self.import_dialog.refresh_destinations()
        if self.store:
            self.statusBar().showMessage(f"자동 저장됨 · {self.store.root}")
        else:
            self.statusBar().showMessage("프로젝트 파일을 불러오거나 새 프로젝트를 만들어 시작하세요.")

    def open_view(self, key, title, mode, node=None, entry=None, query=None):
        for index in range(self.view_tabs.count()):
            if self.view_tabs.tabData(index)["key"] == key:
                self.view_tabs.setCurrentIndex(index)
                self.activate_view(index)
                return
        self.view_tabs.blockSignals(True)
        index = self.view_tabs.addTab(title)
        self.view_tabs.setTabData(index, {"key": key, "mode": mode, "node": node, "entry": entry, "query": query})
        self.view_tabs.blockSignals(False)
        self.close_all_tabs_btn.setEnabled(True)
        self.view_tabs.setCurrentIndex(index)

    def activate_view(self, index):
        if index < 0 or self.changing_view: return
        state = self.view_tabs.tabData(index)
        self.changing_view = True
        try:
            self.selected_id = state["node"] if state["node"] in self.tree_items else None
            self.event_id = state["entry"]
            self.record_query = state["query"]
            self.tree.blockSignals(True)
            self.tree.clearSelection()
            self.tree.setCurrentItem(self.tree_items.get(self.selected_id))
            self.tree.blockSignals(False)
            self.refresh_inventory()
            self.tabs.setCurrentIndex(state.get("detail_tab", 0))
            if state["mode"] == "events": self.refresh_events()
            self.set_mode(state["mode"])
        finally:
            self.changing_view = False

    def remember_detail_tab(self, index):
        if self.changing_view or not hasattr(self, "view_tabs"): return
        state = self.view_tabs.tabData(self.view_tabs.currentIndex())
        if state and state["node"]:
            state["detail_tab"] = index
            self.view_tabs.setTabData(self.view_tabs.currentIndex(), state)

    def close_view(self, index):
        if index > 0:
            self.view_tabs.removeTab(index)
            self.close_all_tabs_btn.setEnabled(self.view_tabs.count() > 1)

    def close_all_views(self):
        self.view_tabs.blockSignals(True)
        while self.view_tabs.count() > 1:
            self.view_tabs.removeTab(1)
        self.view_tabs.setCurrentIndex(0)
        self.view_tabs.blockSignals(False)
        self.close_all_tabs_btn.setEnabled(False)
        self.activate_view(0)

    def open_record(self, ident, query=None):
        if not self.store or not ident: return
        entry = self.store.entry(ident)
        self.open_view(("record", ident, query), entry["category"] + " · " + self.store.target_summary(entry), "events", entry=ident, query=query)

    def set_mode(self, mode):
        if mode == "inventory" and not self.changing_view:
            if self.store and self.selected_id:
                node = self.store.node(self.selected_id)
                self.open_view(("node", self.selected_id), node_caption(node), "inventory", node=self.selected_id)
            else:
                self.show_home()
            return
        if mode == "search" and not self.changing_view:
            self.open_view(("search",), "통합 검색", "search")
            return
        self.mode = mode
        self.stack.setCurrentIndex({"inventory": 0, "events": 1, "search": 2}[mode])
        for key, b in self.nav_buttons.items(): b.setChecked(key == mode)
        if mode == "search":
            self.refresh_search()
            self.search.setFocus()

    def refresh_tree(self):
        expanded = {i for i, item in getattr(self, "tree_items", {}).items() if item.isExpanded()}
        selected = self.tree.ids()
        self.tree.blockSignals(True)
        self.tree.clear()
        self.tree_items = {}
        nodes = self.store.nodes() if self.store else []
        self.equipment_filters.refresh(nodes)
        for n in nodes:
            text = n["serial"] if n["kind"] == "unit" else n["name"]
            item = QTreeWidgetItem([text])
            item.setData(0, Qt.ItemDataRole.UserRole, n["id"])
            item.setData(0, Qt.ItemDataRole.UserRole + 1, n["kind"])
            summary = (n["purpose"] + "\n" if n["kind"] == "unit" else "") + n["description"]
            item.setToolTip(0, f"{KINDS[n['kind']]} · {n['name']}\n{summary}\n{n['tags']}")
            item.setIcon(0, icon(n["kind"], "#3f8e84" if n["kind"] == "unit" else "#71879d"))
            self.tree_items[n["id"]] = item
        for n in nodes:
            item = self.tree_items[n["id"]]
            parent = self.tree_items.get(n["parent"])
            if parent: parent.addChild(item)
            else: self.tree.addTopLevelItem(item)
            item.setExpanded(n["id"] in expanded or not expanded)
            if n["id"] in selected: item.setSelected(True)
        if self.selected_id not in self.tree_items: self.selected_id = None
        if self.selected_id:
            self.tree.setCurrentItem(self.tree_items[self.selected_id], 0, self.tree.selectionModel().SelectionFlag.NoUpdate)
        self.tree.blockSignals(False)
        self.tree_hint.setText(f"{len(nodes):,}개 항목" if nodes else "아직 장비가 없어요.\n＋ 그룹 또는 ＋ 장치로 시작하세요.")
        self.filter_tree()

    def filter_tree(self):
        query = self.tree_filter.text().strip().casefold()
        if not self.store: return
        nodes = self.store.nodes()
        matches = self.equipment_filters.matches(nodes, query)
        visible = with_ancestors(nodes, matches)
        for ident, item in self.tree_items.items():
            item.setHidden(ident not in visible)
            if ident in visible and (query or self.equipment_filters.active()): item.setExpanded(True)
        self.tree_hint.setText(f"{len(matches):,}개 일치 · 전체 {len(nodes):,}개")

    def tree_selection(self):
        current = self.tree.currentItem()
        self.selected_id = current.data(0, Qt.ItemDataRole.UserRole) if current and self.tree.selectedItems() else None
        self.note_limit = 100
        if self.selected_id and not self.changing_view:
            node = self.store.node(self.selected_id)
            selected = self.tree.ids()
            self.open_view(("node", self.selected_id), node_caption(node), "inventory", node=self.selected_id)
            self.tree.blockSignals(True)
            for ident in selected: self.tree_items[ident].setSelected(True)
            self.tree.blockSignals(False)
        else:
            self.refresh_inventory()

    def show_home(self):
        self.view_tabs.setCurrentIndex(0)
        self.activate_view(0)

    def clear_selection(self):
        self.show_home()

    def select_node(self, ident):
        if ident not in self.tree_items: return
        self.tree_filter.clear()
        self.equipment_filters.reset()
        self.tree.clearSelection()
        self.tree.setCurrentItem(self.tree_items[ident])
        self.tree.scrollToItem(self.tree_items[ident])
        self.set_mode("inventory")

    def parent_for_container(self):
        if not self.selected_id: return None
        n = self.store.node(self.selected_id)
        if n["kind"] == "group": return n["id"]
        if n["kind"] == "unit": n = self.store.node(n["parent"])
        return n["parent"]

    def add_node(self, kind):
        if not self.require_project(): return
        parent = self.parent_for_container()
        dialog = NodeDialog(self.store, kind, parent=self)
        if dialog.exec() != QDialog.DialogCode.Accepted: return
        try:
            self.selected_id = self.store.add_node(kind, parent=parent, **dialog.values())
            self.refresh()
            self.select_node(self.selected_id)
        except Exception as exc: self.show_error(exc)

    def edit_node(self):
        if not self.store or not self.selected_id: return
        n = self.store.node(self.selected_id)
        dialog = NodeDialog(self.store, n["kind"], n, self)
        if dialog.exec() != QDialog.DialogCode.Accepted: return
        try:
            self.store.update_node(n["id"], **dialog.values())
            self.refresh()
        except Exception as exc: self.show_error(exc)

    def edit_child(self):
        item = self.children.currentItem()
        if item:
            self.select_node(item.data(0, Qt.ItemDataRole.UserRole))
            self.edit_node()

    def change_detail_purpose(self):
        if not self.store or not self.selected_id: return
        node = self.store.node(self.selected_id)
        if node["kind"] != "unit": return
        dialog = PurposeChangeDialog(self.store, [self.selected_id], self)
        if dialog.exec() != QDialog.DialogCode.Accepted: return
        try:
            self.store.bulk_update_purpose([self.selected_id], dialog.value(), list(dialog.purpose.choices))
            self.refresh()
            self.statusBar().showMessage("개별 장비의 용도를 변경했습니다.", 4000)
        except Exception as exc: self.show_error(exc)

    def add_serials(self):
        if not self.store or not self.selected_id: return
        n = self.store.node(self.selected_id)
        if n["kind"] == "unit": n = self.store.node(n["parent"])
        if n["kind"] != "device": return
        dialog = SerialRegistrationDialog(self.store, n, self)
        try:
            if dialog.exec() == QDialog.DialogCode.Accepted:
                self.refresh()
        finally:
            dialog.deleteLater()

    def move_nodes(self, ids, parent):
        if not self.store: return
        try:
            selected = set(ids)
            roots = [i for i in ids if not any(a["id"] in selected for a in self.store.ancestors(i)[1:])]
            before = {i: self.store.node(i)["parent"] for i in roots if self.store.node(i)["parent"] != parent}
            self.store.move_nodes(ids, parent)
            if before:
                self.undo_moves.append((before, {i: parent for i in before}))
                self.redo_moves.clear()
            self.update_move_actions()
            self.refresh()
        except Exception as exc: self.show_error(exc)

    def update_move_actions(self):
        self.undo_action.setEnabled(bool(self.undo_moves))
        self.redo_action.setEnabled(bool(self.redo_moves))

    def undo_tree_move(self, redo=False):
        source, destination = (self.redo_moves, self.undo_moves) if redo else (self.undo_moves, self.redo_moves)
        if not self.store or not source: return
        before, after = source[-1]
        try:
            self.store.restore_moves(after if redo else before, before if redo else after)
            destination.append(source.pop())
            self.refresh()
            self.update_move_actions()
            self.statusBar().showMessage("트리 이동 다시 실행 완료" if redo else "트리 이동 취소 완료", 4000)
        except Exception as exc: self.show_error(exc)

    def tree_menu(self, pos):
        item = self.tree.itemAt(pos)
        if item and not item.isSelected(): self.tree.setCurrentItem(item)
        menu = QMenu(self)
        menu.addAction("그룹 추가", lambda: self.add_node("group"))
        menu.addAction("장치 추가", lambda: self.add_node("device"))
        if self.selected_id:
            menu.addAction("일련번호 등록", self.add_serials)
            menu.addAction("정보 수정  F2", self.edit_node)
            duplicate_action = menu.addAction("복제  Ctrl+D", self.duplicate_nodes)
            duplicate_action.setEnabled(any(self.store.node(ident)["kind"] in ("group", "device") for ident in self.tree.ids()))
            menu.addAction("선택 장비로 모음 만들기", self.new_collection)
            if self.event_id: menu.addAction("현재 기록에 연결", lambda: self.link_targets(self.tree.ids()))
            menu.addSeparator()
            menu.addAction("선택 항목 삭제  Delete", lambda: self.delete_nodes())
        menu.exec(self.tree.viewport().mapToGlobal(pos))

    def duplicate_nodes(self):
        if not self.store: return
        ids = [ident for ident in self.tree.ids() if self.store.node(ident)["kind"] in ("group", "device")]
        if not ids: return
        try:
            copies = self.store.duplicate_nodes(ids)
            self.refresh()
            if copies:
                self.select_node(copies[0])
                for ident in copies: self.tree_items[ident].setSelected(True)
            self.statusBar().showMessage(f"{len(copies)}개 항목과 하위 장비 복제 완료", 5000)
        except Exception as exc: self.show_error(exc)

    def delete_nodes(self, ids=None):
        if not self.store: return
        ids = self.tree.ids() if ids is None else ids
        if not ids: return
        count = len(self.store.descendants(ids))
        if QMessageBox.question(self, "휴지통으로 이동", f"선택 항목과 하위 항목 {count:,}개를 삭제할까요?\n휴지통에서 복원할 수 있습니다.") != QMessageBox.StandardButton.Yes: return
        try:
            self.store.delete(ids)
            self.refresh()
        except Exception as exc: self.show_error(exc)

    def refresh_inventory(self):
        n = self.store.node(self.selected_id) if self.store and self.selected_id else None
        self.breadcrumb.setVisible(n is not None)
        self.detail_hint.setVisible(n is not None)
        self.device_badge.setVisible(n is not None and n["kind"] == "unit")
        self.metrics_panel.setVisible(n is not None)
        self.tabs.setVisible(n is not None)
        self.calendar_panel.setVisible(n is None)
        self.add_unit_btn.setVisible(n is not None)
        self.edit_btn.setVisible(n is not None)
        if n is None:
            self.refresh_calendar()
        self.add_unit_btn.setEnabled(n is not None and n["kind"] in ("device", "unit"))
        self.edit_btn.setEnabled(n is not None)
        self.inherited.setEnabled(n is not None)
        if n:
            self.breadcrumb.setText(self.store.path_label(n["id"]))
            self.detail_title.setText(n["serial"] if n["kind"] == "unit" else n["name"])
            if n["kind"] == "unit":
                self.device_badge.setText(self.store.node(n["parent"])["name"])
            self.detail_hint.setText(KINDS[n["kind"]] + (f" · {n['tags']}" if n["tags"] else ""))
            location, status = self.store.current_state(n["id"])
            stats = [("현재 위치", location), ("최근 상태", status), ("용도", n["purpose"] or "미지정") if n["kind"] == "unit" else ("설명", n["description"] or "미입력")]
            esc = html.escape
            rows = [("종류", KINDS[n["kind"]]), ("이름", n["name"]), ("일련번호", n["serial"] or "—"), ("용도", n["purpose"] or "미지정"), ("해시태그", n["tags"] or "없음"), ("설명", n["description"] or "아직 설명이 없습니다.")]
            if n["kind"] != "unit": rows = [(key, value) for key, value in rows if key != "용도"]
            self.basic_text.setHtml("<style>body{color:#334155;}h3{color:#788598;font-size:11px;}p{font-size:14px;}</style>" + "".join(f"<h3>{esc(k)}</h3><p>{esc(v).replace(chr(10), '<br>')}</p>" for k, v in rows))
        else:
            nodes = self.store.nodes() if self.store else []
            self.breadcrumb.setText("WORKSPACE / EQUIPMENT")
            self.detail_title.setText("전체 장비" if self.store else "장비의 모든 이력, 한곳에서")
            self.detail_hint.setText("장비를 선택해 위치와 이력을 확인하세요. 작성한 내용은 프로젝트에 자동 저장됩니다." if self.store else "새 프로젝트를 만들고, 장비의 첫 번째 기록을 시작해보세요.")
            stats = [("등록 장치", f"{sum(n['kind'] == 'device' for n in nodes):,}"), ("개별 장비", f"{sum(n['kind'] == 'unit' for n in nodes):,}"), ("기록", f"{len(self.store.entries(events_only=True)) if self.store else 0:,}")]
            self.basic_text.setPlainText("왼쪽 트리에서 항목을 선택하면 설명과 용도를 확인할 수 있습니다.")
        for (name, value), (key, val) in zip(self.metrics, stats):
            name.setText(key)
            display = " ".join(val.split()) if key == "설명" else val
            value.setText(display[:80] + "…" if key == "설명" and len(display) > 80 else display)
            value.setToolTip(val)
            value.setStyleSheet("font-size: 13px; font-weight: 400;" if key == "설명" else "")
        purpose_value = self.metrics[2][1]
        if n is not None and n["kind"] == "unit":
            purpose_value.setCursor(Qt.CursorShape.PointingHandCursor)
            purpose_value.setToolTip((n["purpose"] or "미지정") + " · 더블클릭하여 용도 변경")
            purpose_value.setAccessibleName("용도 · 더블클릭하여 변경")
        else:
            purpose_value.setCursor(Qt.CursorShape.ArrowCursor)
            purpose_value.setAccessibleName("")
        self.children.clear()
        if self.store:
            for child in self.store.nodes():
                if child["parent"] == self.selected_id:
                    i = QTreeWidgetItem([child["serial"] if child["kind"] == "unit" else child["name"], KINDS[child["kind"]], child["purpose"] if child["kind"] == "unit" else child["description"]])
                    i.setData(0, Qt.ItemDataRole.UserRole, child["id"])
                    self.children.addTopLevelItem(i)
        self.refresh_history()

    def refresh_history(self):
        self.timeline.clear()
        entries = self.store.entries(self.selected_id, self.inherited.isChecked(), category=self.category_filter.currentText()) if self.store else []
        for e in entries[:self.note_limit]:
            source = self.store.target_summary(e)
            if e["inherited"]: source = "상위 이력 · " + source
            item = QListWidgetItem()
            item.setData(Qt.ItemDataRole.UserRole, e["id"])
            item.setSizeHint(QSize(320, 102))
            self.timeline.addItem(item)
            self.timeline.setItemWidget(item, NoteCard(e, source))
        self.history_empty.setText(f"{len(entries):,}개 기록 · 더블클릭하여 내용 확인 / 수정" if entries else "아직 기록이 없습니다.\n상단의 기록 작성에서 대상 장비와 기록을 추가하세요.")
        self.more_btn.setVisible(len(entries) > self.note_limit)

    def more_history(self):
        self.note_limit += 100
        self.refresh_history()

    def new_event(self, checked=False, day=None):
        if not self.require_project(): return
        selected_ids = self.tree.ids() if self.mode == "inventory" and self.selected_id and day is None else []
        dialog = EntryDialog(self.store, event=True, parent=self, selected_ids=selected_ids, day=day)
        dialog.save_handler = lambda: self.store.save_record(
            event=True, node_ids=dialog.target_ids(),
            purpose_options=list(dialog.purpose_after.choices), **dialog.values())
        try:
            if dialog.exec() == QDialog.DialogCode.Accepted:
                self.refresh()
                self.statusBar().showMessage("기록을 저장했습니다.", 4000)
        finally:
            dialog.save_handler = None
            dialog.deleteLater()

    def edit_entry(self, ident):
        if not self.store or not ident: return
        e = self.store.entry(ident)
        dialog = EntryDialog(self.store, event=e["event"], entry=e, parent=self)
        dialog.save_handler = lambda: self.store.save_record(
            ident=ident, event=e["event"],
            node_ids=dialog.target_ids() if e["event"] else self.store.targets(ident),
            purpose_options=list(dialog.purpose_after.choices), **dialog.values())
        try:
            if dialog.exec() == QDialog.DialogCode.Accepted:
                self.refresh()
        finally:
            dialog.save_handler = None
            dialog.deleteLater()

    def delete_entries(self, widget=None):
        if not self.store: return
        widget = widget or self.timeline
        ids = [i.data(Qt.ItemDataRole.UserRole) for i in widget.selectedItems()]
        if not ids: return
        if QMessageBox.question(self, "기록 삭제", f"선택한 기록 {len(ids)}개를 휴지통으로 이동할까요?\n삭제한 기록은 연결된 모든 장비의 조회 결과에서 제외됩니다.") != QMessageBox.StandardButton.Yes: return
        try:
            self.store.delete(ids, "entries")
            self.refresh()
        except Exception as exc: self.show_error(exc)

    def refresh_events(self):
        self.events.blockSignals(True)
        self.events.clear()
        entries = self.store.entries(query=self.record_query or "") if self.store else []
        if self.record_query is None: entries = [e for e in entries if e["id"] == self.event_id]
        self.events.setVisible(self.record_query is not None)
        self.records_hint.setText(f"검색어 ‘{self.record_query}’에 해당하는 기록 {len(entries)}개" if self.record_query is not None else "기록의 내용과 연결된 장비를 확인하세요.")
        found = False
        for e in entries:
            i = QListWidgetItem(f"{e['category']} · {record_period(e)}\n{self.store.target_summary(e)}\n{e['plain'][:70]}")
            i.setData(Qt.ItemDataRole.UserRole, e["id"])
            i.setToolTip(self.store.target_summary(e, limit=100))
            i.setSizeHint(QSize(230, 114))
            self.events.addItem(i)
            if e["id"] == self.event_id:
                self.events.setCurrentItem(i)
                found = True
        if not found:
            self.event_id = entries[0]["id"] if entries else None
            if entries: self.events.setCurrentRow(0)
        self.events.blockSignals(False)
        self.refresh_event_detail()

    def event_selection(self):
        item = self.events.currentItem()
        self.event_id = item.data(Qt.ItemDataRole.UserRole) if item else None
        if self.view_tabs.count():
            state = self.view_tabs.tabData(self.view_tabs.currentIndex())
            if state["mode"] == "events":
                state["entry"] = self.event_id
                state["key"] = ("record", self.event_id, self.record_query)
                if self.event_id:
                    e = self.store.entry(self.event_id)
                    self.view_tabs.setTabText(self.view_tabs.currentIndex(), e["category"] + " · " + self.store.target_summary(e))
                self.view_tabs.setTabData(self.view_tabs.currentIndex(), state)
        self.refresh_event_detail()

    def refresh_event_detail(self):
        self.targets.clear()
        enabled = bool(self.store and self.event_id)
        for b in (self.event_edit_btn, self.link_btn, self.unlink_btn): b.setEnabled(enabled)
        if not enabled:
            self.event_title.setText("기록을 만들어보세요")
            self.event_meta.clear()
            self.event_body.setPlainText("현장 시험, 정기 점검, 반출 일정 등을 만들고 관련 장비를 연결하세요.")
            self.target_count.setText("연결된 장비 0개")
            return
        e = self.store.entry(self.event_id)
        self.event_title.setText(e["category"])
        self.event_meta.setText(f"{record_period(e)} · {e['category']} · {e['location']}\n{e['tags']}")
        self.event_body.setHtml(e["html"] or html.escape(e["plain"] or e["title"]))
        targets = self.store.targets(self.event_id)
        for ident in targets:
            i = QListWidgetItem(self.store.path_label(ident))
            i.setData(Qt.ItemDataRole.UserRole, ident)
            self.targets.addItem(i)
        self.target_count.setText(f"연결된 장비 {len(targets):,}개")

    def link_targets(self, ids):
        if not self.store or not self.event_id: return
        try:
            self.store.link_event(self.event_id, ids)
            self.refresh()
        except Exception as exc: self.show_error(exc)

    def unlink_targets(self):
        if not self.store or not self.event_id: return
        ids = [i.data(Qt.ItemDataRole.UserRole) for i in self.targets.selectedItems()]
        try:
            self.store.unlink_event(self.event_id, ids)
            self.refresh()
        except Exception as exc: self.show_error(exc)

    def refresh_search(self):
        self.search_results.clear()
        query = self.search.text().strip()
        if not self.store or not query:
            self.search_count.setText("검색어를 입력하세요.")
            return
        count = 0
        for n in self.store.nodes():
            matched = all(t in n["tags"].casefold().split() for t in tags(query).casefold().split()) if query.startswith("#") else query.casefold() in (node_caption(n) + " " + " ".join(n[k] for k in ("name", "serial", "purpose", "description", "tags"))).casefold()
            if matched:
                i = QTreeWidgetItem([KINDS[n["kind"]], node_caption(n), self.store.path_label(n["id"]), ""])
                i.setData(0, Qt.ItemDataRole.UserRole, ("node", n["id"]))
                self.search_results.addTopLevelItem(i)
                count += 1
        for e in self.store.entries(query=query):
            names = [self.store.path_label(i) for i in e["targets"]]
            source = ", ".join(names[:2]) + (f" 외 {len(names)-2}개" if len(names) > 2 else "")
            i = QTreeWidgetItem(["기록" if e["event"] else "이력", e["title"], source, record_period(e)])
            i.setToolTip(2, "\n".join(names))
            i.setData(0, Qt.ItemDataRole.UserRole, ("entry", e["id"]))
            self.search_results.addTopLevelItem(i)
            count += 1
        self.search_count.setText(f"{count:,}개 결과")

    def activate_search(self, item, col):
        kind, ident = item.data(0, Qt.ItemDataRole.UserRole)
        if kind == "node": self.select_node(ident)
        else:
            self.open_record(ident, self.search.text().strip())

    def delete_search_results(self):
        if not self.store: return
        selected = [i.data(0, Qt.ItemDataRole.UserRole) for i in self.search_results.selectedItems()]
        if not selected: return
        if QMessageBox.question(self, "검색 결과 삭제", "선택한 원본 장비·기록을 휴지통으로 이동할까요?\n장비의 하위 항목도 함께 삭제됩니다.") != QMessageBox.StandardButton.Yes: return
        try:
            with self.store.atomic():
                self.store.delete([ident for kind, ident in selected if kind == "node"])
                self.store.delete([ident for kind, ident in selected if kind == "entry"], "entries")
            self.refresh()
        except Exception as exc: self.show_error(exc)

    def open_import(self):
        if not self.require_project(): return
        if self.import_dialog and self.import_dialog.isVisible():
            self.import_dialog.raise_()
            return
        self.import_dialog = ImportDialog(self)
        self.import_dialog.show()

    def export_equipment(self):
        if not self.require_project(): return
        dest, _ = QFileDialog.getSaveFileName(self, "전체 장비 목록 내보내기",
            f"{self.store.name}_장비목록_{date.today().isoformat()}.xlsx", "Excel (*.xlsx)")
        if not dest: return
        try:
            if not dest.lower().endswith(".xlsx"): dest += ".xlsx"
            export_equipment_list(self.store, dest)
            self.statusBar().showMessage("전체 장치·일련번호 목록 내보내기 완료", 7000)
        except Exception as exc: self.show_error(exc)

    def open_equipment_comparison(self):
        if self.require_project():
            EquipmentCompareDialog(self).exec()

    def drop_import(self, data, parent):
        if not self.import_dialog or data.get("session") != self.import_dialog.session: return
        self.import_dialog.register(data["rows"], parent)

    def save_template(self):
        source = ROOT / "outputs" / "semp" / "SEMP_장비등록_템플릿.xlsx"
        if not source.exists():
            QMessageBox.warning(self, "템플릿 없음", "프로그램의 outputs/semp 폴더에 템플릿 파일이 없습니다.")
            return
        dest, _ = QFileDialog.getSaveFileName(self, "엑셀 템플릿 저장", source.name, "Excel (*.xlsx)")
        if dest:
            try:
                if not dest.lower().endswith(".xlsx"): dest += ".xlsx"
                if Path(dest).resolve() != source.resolve(): shutil.copyfile(source, dest)
                self.statusBar().showMessage("엑셀 템플릿 저장 완료", 5000)
            except Exception as exc: self.show_error(exc)

    def show_trash(self):
        if not self.require_project(): return
        dialog = QDialog(self)
        dialog.setWindowTitle("휴지통")
        dialog.resize(560, 400)
        layout = QVBoxLayout(dialog)
        layout.setContentsMargins(24, 24, 24, 24)
        layout.addWidget(label("삭제한 항목 복원", "Title"))
        listing = ActionList()
        listing.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        layout.addWidget(listing)
        def populate():
            listing.clear()
            for r in self.store.trash():
                i = QListWidgetItem(f"{r['label']}   ·   {r['created'].replace('T',' ')}")
                i.setData(Qt.ItemDataRole.UserRole, r["id"])
                listing.addItem(i)
        def restore():
            i = listing.currentItem()
            if not i: return
            try:
                self.store.restore(i.data(Qt.ItemDataRole.UserRole))
                self.refresh()
                populate()
            except Exception as exc: self.show_error(exc)
        populate()
        layout.addLayout(row(label("삭제된 상위 폴더도 필요하면 함께 복원됩니다.", "Subtitle"), None))
        layout.addLayout(row(None, button("닫기", dialog.reject), button("선택 복원", restore, True)))
        dialog.exec()

    def show_help(self):
        QMessageBox.information(self, "SEMP 사용 안내", "1. 새 프로젝트에서 저장 위치와 이름을 선택합니다.\n2. 그룹 → 장치 → 일련번호 순서로 등록합니다.\n3. 달력에서 날짜별 기록을 확인하고 기록 작성 버튼을 누릅니다.\n\nCtrl/Shift: 다중 선택 · Delete: 삭제 · F2: 정보 수정\nCtrl+N: 새 프로젝트 · Ctrl+O: 열기 · Ctrl+S: 백업 · Ctrl+F: 검색\n\n기록과 정보는 저장 버튼을 누르면 즉시 저장됩니다.\n프로젝트를 닫고 프로젝트 폴더 전체를 다른 PC로 옮기세요.\n\n현재 위치와 상태는 오늘까지의 직접 기록·연결 기록으로 계산합니다.\n그룹 기록은 하위 장비의 상태를 자동 변경하지 않습니다.\n같은 날짜는 나중에 생성한 기록을 우선합니다.\n\n자세한 사용법과 현재 범위는 README.md를 참고하세요.")


def run(project=None):
    app = QApplication.instance() or QApplication(sys.argv)
    app.setApplicationName("SEMP")
    app.setOrganizationName("SEMP")
    app.setStyle("Fusion")
    app.setFont(QFont("Malgun Gothic", 10))
    app.setStyleSheet(STYLE)
    window = MainWindow()
    window.set_mode("inventory")
    if project: window.load_project(project)
    window.show()
    return app.exec()


if __name__ == "__main__":
    raise SystemExit(run(sys.argv[1] if len(sys.argv) > 1 else None))
