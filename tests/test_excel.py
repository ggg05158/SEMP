import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from SEMP import create_import_template, read_import

TEMPLATE = Path(__file__).resolve().parent.parent / "outputs/semp/SEMP_장비등록_템플릿.xlsx"
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


class ExcelTests(unittest.TestCase):
    def fixture(self, path, rows):
        # Inject controlled test cells into a copy of the shipped template.
        with ZipFile(TEMPLATE) as source, ZipFile(path,"w") as dest:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    root = ET.fromstring(data)
                    sheet_data = root.find(f"{{{NS}}}sheetData")
                    for number, values in enumerate(rows, 2):
                        row = next((r for r in sheet_data if r.attrib["r"] == str(number)), None)
                        if row is None:
                            row = ET.SubElement(sheet_data,f"{{{NS}}}row",r=str(number))
                        for col, value in zip("ABCD",values):
                            address = f"{col}{number}"
                            for existing in list(row):
                                if existing.attrib.get("r") == address: row.remove(existing)
                            if value is not None:
                                cell = ET.SubElement(row,f"{{{NS}}}c",r=address,t="inlineStr")
                                inline = ET.SubElement(cell,f"{{{NS}}}is")
                                ET.SubElement(inline,f"{{{NS}}}t").text = value
                    data = ET.tostring(root,encoding="utf-8",xml_declaration=True)
                dest.writestr(item,data)

    def test_shipped_template_is_empty_and_text_formatted(self):
        self.assertEqual(read_import(TEMPLATE), [])
        wb = load_workbook(TEMPLATE)
        try:
            self.assertEqual(wb.active["B2"].number_format, "@")
            self.assertEqual(wb.active["C101"].number_format, "@")
            self.assertEqual(wb.active.freeze_panes, "A2")
        finally: wb.close()

    def test_template_is_generated_without_bundled_xlsx(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "generated.xlsx"
            create_import_template(path)
            self.assertTrue(path.is_file())
            wb = load_workbook(path)
            try:
                ws = wb["장비등록"]
                self.assertEqual([ws.cell(1, column).value for column in range(1, 8)],
                                 ["장치이름", "일련번호 처음", "일련번호 끝", "용도", "위치", "설명", "태그"])
                self.assertEqual(ws.freeze_panes, "A2")
                self.assertFalse(ws.sheet_view.showGridLines)
                self.assertEqual(ws["B2"].number_format, "@")
                self.assertEqual(ws["C101"].number_format, "@")
                self.assertIn("점검용", ws["I15"].value)
                ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "GCU", "001", "003", "비행용"
                wb.save(path)
            finally:
                wb.close()
            rows = read_import(path)
            self.assertEqual(rows[0].serials, ["001", "002", "003"])
            self.assertEqual(rows[0].purpose, "비행용")

    def test_ranges_and_empty_serial_in_actual_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp)/"input.xlsx"
            self.fixture(path, [("장치A","001","003"),("장치B","PSU-09","PSU-11"),("케이블",None,None),("장치C","SN-X",None)])
            rows = read_import(path)
            self.assertEqual(len(rows),4)
            self.assertEqual(rows[0].serials,["001","002","003"])
            self.assertEqual(rows[1].serials,["PSU-09","PSU-10","PSU-11"])
            self.assertEqual(rows[2].serials,[])
            self.assertEqual(rows[3].serials,["SN-X"])
            self.assertFalse(any(r.error for r in rows))

    def test_purpose_column_and_custom_value(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "purposes.xlsx"
            self.fixture(path, [("GCU", "001", "002", "점검용"), ("PSU", "A001", None, "교육용"), ("케이블", None, None, "비행용")])
            rows = read_import(path)
            self.assertEqual(rows[0].purpose, "점검용")
            self.assertEqual(rows[1].purpose, "교육용")
            self.assertTrue(rows[2].error)
        wb = load_workbook(TEMPLATE)
        self.assertEqual(wb.active["D1"].value, "용도")
        self.assertEqual(wb.active["D2"].number_format, "@")
        self.assertIn("점검용", wb.active["F14"].value)
        wb.close()

    def test_tree_rows_and_extended_equipment_fields(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "tree.xlsx"
            create_import_template(path)
            wb = load_workbook(path)
            ws = wb["장비등록"]
            ws.append(["장치A", "2501", None, "실전용", "격납고 A", "00할 때 쓰는 장비", "#주력 #비행"])
            ws.append([None, "2502", None, "예비용", "격납고 B", "예비 장비", "예비"])
            wb.save(path)
            wb.close()
            rows = read_import(path)
            self.assertEqual([row.name for row in rows], ["장치A", "장치A"])
            self.assertEqual(rows[0].location, "격납고 A")
            self.assertEqual(rows[0].description, "00할 때 쓰는 장비")
            self.assertEqual(rows[0].tag_text, "#주력 #비행")
            self.assertEqual(rows[1].tag_text, "#예비")

    def test_validation_keeps_valid_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp)/"input.xlsx"
            self.fixture(path, [("장치A","001","003"),("장치A","002",None),("장치B",None,"005"),(None,"001",None),("장치C","020","010")])
            rows = read_import(path)
            self.assertFalse(rows[0].error)
            self.assertTrue(rows[1].error)
            self.assertTrue(rows[2].error)
            self.assertFalse(rows[3].error)  # Blank device name inherits 장치B from the row above.
            self.assertEqual(rows[3].name, "장치B")
            self.assertTrue(rows[4].error)


if __name__ == "__main__": unittest.main()
