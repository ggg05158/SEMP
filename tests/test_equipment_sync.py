"""Round-trip comparisons run only against disposable project databases."""
import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from PySide6.QtCore import QSettings, QDate
from PySide6.QtGui import QFontDatabase, QFont
from PySide6.QtWidgets import QApplication
from SEMP import (Store, MainWindow, EquipmentCompareDialog, EquipmentListRow,
                  export_equipment_list, read_equipment_list, compare_equipment_list,
                  apply_equipment_changes, equipment_snapshot, LIST_HEADERS, STYLE)


class EquipmentSyncTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.store = Store.create(self.tmp.name, "엑셀 비교", default_groups=False)
        self.group = self.store.add_node("group", "장비")
        self.device = self.store.add_node("device", "GCU", self.group, description="제어장치")
        self.unit = self.store.add_node("unit", "GCU", self.device, serial="001", purpose="비행용", tag_text="#시험")
        self.path = Path(self.tmp.name) / "list.xlsx"
        export_equipment_list(self.store, self.path)

    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()

    def rows(self):
        rows, errors = read_equipment_list(self.path, self.store)
        self.assertEqual(errors, [])
        return rows

    def test_unchanged_export_contains_all_nodes_and_is_noop(self):
        rows = self.rows()
        self.assertEqual([row.ident for row in rows], [self.device, self.unit])
        self.assertEqual(rows[1].values["serial"], "001")
        plan = compare_equipment_list(self.store, rows)
        self.assertEqual(plan.changes, [])
        self.assertEqual(plan.errors, [])
        self.assertEqual(plan.omitted, [])
        self.assertEqual(apply_equipment_changes(self.store, plan, "2026-10-01"), (0, None))
        self.assertEqual(self.store.entries(), [])
        workbook = load_workbook(self.path)
        try:
            sheet = workbook["장비목록"]
            self.assertEqual(tuple(cell.value for cell in sheet[1]), LIST_HEADERS)
            self.assertEqual(sheet["D3"].number_format, "@")
            self.assertEqual(sheet.freeze_panes, "D2")
            self.assertEqual(sheet.auto_filter.ref, "A1:K3")
            self.assertEqual(workbook["SEMP 정보"].sheet_state, "hidden")
        finally:
            workbook.close()

    def test_edit_excel_add_serial_and_apply_dated_history_with_backup(self):
        workbook = load_workbook(self.path)
        sheet = workbook["장비목록"]
        sheet["E3"] = "점검용"
        sheet["F3"] = "정비실"
        sheet["G3"] = "수정 설명"
        sheet["H3"] = "정비 완료"
        sheet.append(["개별 장비", "장비", "GCU", "002", "예비용", "창고", "새 장비", "#예비", "", "", self.device])
        workbook.save(self.path)
        workbook.close()
        plan = compare_equipment_list(self.store, self.rows())
        self.assertEqual(plan.errors, [])
        self.assertEqual(len(plan.changes), 2)
        self.assertEqual(set(plan.changes[0].differences()), {"purpose", "location", "description", "tags"})
        count, backup = apply_equipment_changes(self.store, plan, "2026-10-01")
        self.assertEqual(count, 2)
        self.assertTrue(backup.is_file())
        unit = self.store.node(self.unit)
        self.assertEqual(unit["purpose"], "점검용")
        self.assertEqual(unit["description"], "수정 설명")
        self.assertEqual(unit["tags"], "#정비 #완료")
        self.assertEqual(self.store.current_state(self.unit)[0], "정비실")
        self.assertEqual({node["serial"] for node in self.store.nodes() if node["kind"] == "unit"}, {"001", "002"})
        self.assertTrue(all(entry["day"] == "2026-10-01" for entry in self.store.entries()))
        export_equipment_list(self.store, self.path)
        self.assertEqual(compare_equipment_list(self.store, self.rows()).changes, [])

    def test_missing_rows_do_not_delete_equipment(self):
        plan = compare_equipment_list(self.store, [])
        self.assertEqual(set(plan.omitted), {self.device, self.unit})
        self.assertEqual(plan.changes, [])
        apply_equipment_changes(self.store, plan, "2026-10-01")
        self.assertEqual(len(self.store.nodes()), 3)

    def test_new_device_and_units_can_be_out_of_order(self):
        values = dict(name="PSU", serial="SN-010", purpose="예비용", location="창고", description="", tags="")
        rows = self.rows() + [EquipmentListRow(5, "unit", "", "", values),
            EquipmentListRow(6, "device", "", "", {**values, "serial": "", "purpose": ""})]
        plan = compare_equipment_list(self.store, rows, new_device_parent=self.group)
        self.assertEqual(plan.errors, [])
        apply_equipment_changes(self.store, plan, "2026-10-01")
        device = next(node for node in self.store.nodes() if node["kind"] == "device" and node["name"] == "PSU")
        self.assertEqual(device["parent"], self.group)
        unit = next(node for node in self.store.nodes() if node["serial"] == "SN-010")
        self.assertEqual(unit["parent"], device["id"])

    def test_ambiguous_parent_duplicate_ids_and_duplicate_serials_are_blocked(self):
        rows = self.rows()
        duplicate = compare_equipment_list(self.store, rows + [rows[1]])
        self.assertTrue(duplicate.errors)
        with self.assertRaises(ValueError):
            apply_equipment_changes(self.store, duplicate, "2026-10-01")
        group = self.store.add_node("group", "다른 그룹")
        self.store.add_node("device", "GCU", group)
        new = EquipmentListRow(7, "unit", "", "", {**rows[1].values, "serial": "002"})
        self.assertTrue(compare_equipment_list(self.store, rows + [new]).errors)

    def test_comparison_rejects_changes_after_preview(self):
        rows = self.rows()
        rows[1].values["purpose"] = "점검용"
        plan = compare_equipment_list(self.store, rows)
        self.store.bulk_update_purpose([self.unit], "환경용")
        with self.assertRaisesRegex(ValueError, "비교 이후"):
            apply_equipment_changes(self.store, plan, "2026-10-01")
        self.assertEqual(self.store.node(self.unit)["purpose"], "환경용")

    def test_transaction_failure_rolls_back_all_nodes_and_records(self):
        rows = self.rows()
        rows[1].values["purpose"] = "점검용"
        plan = compare_equipment_list(self.store, rows)
        before = equipment_snapshot(self.store)
        with patch.object(self.store, "save_entry", side_effect=RuntimeError("기록 저장 실패")):
            with self.assertRaises(RuntimeError):
                apply_equipment_changes(self.store, plan, "2026-10-01")
        self.assertEqual(equipment_snapshot(self.store), before)
        self.assertEqual(self.store.entries(), [])

    def test_project_identity_and_formula_validation(self):
        other = Store.create(self.tmp.name, "다른 프로젝트", default_groups=False)
        try:
            with self.assertRaisesRegex(ValueError, "다른 프로젝트"):
                read_equipment_list(self.path, other)
        finally:
            other.close()
        workbook = load_workbook(self.path)
        workbook["장비목록"]["G3"] = "=1+1"
        workbook.save(self.path)
        workbook.close()
        rows, errors = read_equipment_list(self.path, self.store)
        self.assertEqual(len(errors), 1)
        self.assertIn("수식", errors[0])

    def test_serial_swaps_preserve_ids(self):
        second = self.store.add_node("unit", "GCU", self.device, serial="002")
        export_equipment_list(self.store, self.path)
        rows = self.rows()
        rows[1].values["serial"], rows[2].values["serial"] = "002", "001"
        plan = compare_equipment_list(self.store, rows)
        self.assertEqual(plan.errors, [])
        apply_equipment_changes(self.store, plan, "2026-10-01")
        self.assertEqual(self.store.node(self.unit)["serial"], "002")
        self.assertEqual(self.store.node(second)["serial"], "001")

    def test_device_rename_keeps_units_and_their_history(self):
        record = self.store.save_entry(title="기존 기록", category="점검", day="2026-01-01",
            html="점검", plain="점검", node_ids=[self.unit])
        rows = self.rows()
        rows[0].values["name"] = "GCU II"
        plan = compare_equipment_list(self.store, rows)
        self.assertEqual(plan.errors, [])
        self.assertEqual(len(plan.changes), 1)
        apply_equipment_changes(self.store, plan, "2026-10-01")
        self.assertEqual(self.store.node(self.unit)["name"], "GCU II")
        self.assertEqual(self.store.targets(record), [self.unit])
        export_equipment_list(self.store, self.path)
        self.assertEqual(compare_equipment_list(self.store, self.rows()).changes, [])

    def test_existing_root_device_stays_at_root_when_default_group_selected(self):
        root_device = self.store.add_node("device", "루트 장치")
        export_equipment_list(self.store, self.path)
        rows = self.rows()
        next(row for row in rows if row.ident == root_device).values["description"] = "수정"
        plan = compare_equipment_list(self.store, rows, new_device_parent=self.group)
        self.assertEqual(plan.errors, [])
        apply_equipment_changes(self.store, plan, "2026-10-01")
        self.assertIsNone(self.store.node(root_device)["parent"])

    def test_duplicate_serial_against_omitted_equipment_is_rejected(self):
        self.store.add_node("unit", "GCU", self.device, serial="002")
        rows = self.rows()
        rows[1].values["serial"] = "002"
        plan = compare_equipment_list(self.store, rows)
        self.assertTrue(plan.errors)
        with self.assertRaises(ValueError):
            apply_equipment_changes(self.store, plan, "2026-10-01")
        self.assertEqual(self.store.node(self.unit)["serial"], "001")

    def test_literal_formula_text_and_identity_survive_export_and_reopen(self):
        self.store.update_node(self.device, "GCU", "", "", "=literal text", "")
        export_equipment_list(self.store, self.path)
        rows = self.rows()
        self.assertEqual(rows[0].values["description"], "=literal text")
        identity = self.store.project_identity
        self.store.close()
        self.store = Store(self.store.path)
        self.assertEqual(self.store.project_identity, identity)
        self.assertEqual(compare_equipment_list(self.store, self.rows()).changes, [])

    def test_location_can_be_cleared_and_future_location_does_not_replace_today(self):
        self.store.save_entry(title="이동", category="위치", day="2026-01-01", html="창고", plain="창고", location="창고", node_ids=[self.unit])
        export_equipment_list(self.store, self.path)
        rows = self.rows()
        rows[1].values["location"] = ""
        plan = compare_equipment_list(self.store, rows)
        apply_equipment_changes(self.store, plan, "2026-02-01")
        self.assertEqual(self.store.current_state(self.unit)[0], "미지정")
        export_equipment_list(self.store, self.path)
        rows = self.rows()
        rows[1].values["location"] = "미래 현장"
        apply_equipment_changes(self.store, compare_equipment_list(self.store, rows), "2099-01-01")
        self.assertEqual(self.store.current_state(self.unit)[0], "미지정")


class EquipmentSyncGuiTests(unittest.TestCase):
    setUp = EquipmentSyncTests.setUp
    tearDown = EquipmentSyncTests.tearDown
    def test_preview_cancel_and_apply(self):
        app = QApplication.instance() or QApplication([])
        for font in ("malgun.ttf", "malgunbd.ttf", "segoeui.ttf"):
            path = Path("C:/Windows/Fonts") / font
            if path.exists():
                QFontDatabase.addApplicationFont(str(path))
        app.setFont(QFont("Malgun Gothic", 10))
        app.setStyle("Fusion")
        app.setStyleSheet(STYLE)
        workbook = load_workbook(self.path)
        workbook["장비목록"]["E3"] = "점검용"
        workbook.save(self.path)
        workbook.close()
        settings = QSettings(str(Path(self.tmp.name) / "settings.ini"), QSettings.Format.IniFormat)
        window = MainWindow(settings)
        window.load_project(self.store.path)
        dialog = EquipmentCompareDialog(window)
        with patch("SEMP.QFileDialog.getOpenFileName", return_value=(str(self.path), "Excel (*.xlsx)")):
            dialog.load()
        self.assertTrue(dialog.apply_button.isEnabled())
        self.assertEqual(dialog.table.rowCount(), 1)
        self.assertEqual(dialog.table.item(0, 3).text(), "비행용")
        self.assertEqual(dialog.table.item(0, 4).text(), "점검용")
        self.assertEqual(self.store.node(self.unit)["purpose"], "비행용")
        dialog.reject()
        self.assertEqual(self.store.node(self.unit)["purpose"], "비행용")
        dialog.day.setDate(QDate(2026, 10, 1))
        dialog.calendar_button.click()
        app.processEvents()
        self.assertTrue(dialog.mini_calendar.isVisible())
        self.assertEqual(dialog.mini_calendar.selectedDate(), QDate(2026, 10, 1))
        dialog.mini_calendar.hide()
        dialog.show()
        app.processEvents()
        dialog.grab().save(str(Path(__file__).resolve().parent.parent / ".qa" / "equipment-compare.png"))
        dialog.apply()
        self.assertEqual(self.store.node(self.unit)["purpose"], "점검용")
        # A failed replacement file must not re-enable the previous comparison.
        dialog = EquipmentCompareDialog(window)
        with patch("SEMP.QFileDialog.getOpenFileName", return_value=(str(self.path), "Excel (*.xlsx)")):
            dialog.load()
        workbook = load_workbook(self.path)
        workbook["장비목록"]["A1"] = "잘못된 양식"
        workbook.save(self.path)
        workbook.close()
        with patch("SEMP.QFileDialog.getOpenFileName", return_value=(str(self.path), "Excel (*.xlsx)")), patch.object(window, "show_error"):
            dialog.load()
        dialog.destination.setCurrentIndex(1)
        self.assertFalse(dialog.apply_button.isEnabled())
        self.assertIsNone(dialog.comparison)
        dialog.reject()
        window.close()


if __name__ == "__main__":
    unittest.main()
